# -*- coding: utf-8 -*-
r"""ddrescue 虛擬機救援的 Windows 端控制程式（操作手冊：docs/方法二_虛擬機ddrescue.md）

  python rescue_dd.py vm-check          確認虛擬機、共用資料夾、硬碟唯讀、Windows 沒抓到硬碟
  python rescue_dd.py mft               R1：只救 MBR、boot sector、整個 $MFT
  python rescue_dd.py plan              R2：從映像檔的 MFT 比對 G 槽，產生要讀的範圍（不碰硬碟）
  python rescue_dd.py copy [--batch bNN] R3：讀一批（不指定＝照順序下一批），邊讀邊輸出到 G 槽、邊釋放映像檔；讀完這批就結束
  python rescue_dd.py batches           各批次狀態
  python rescue_dd.py retry             R4：所有批次讀完後，一起補讀壞區
  python rescue_dd.py extract [--final] R5：從映像檔輸出到 G 槽；--final 才決定部分損壞/損壞
  python rescue_dd.py report            印目前統計（不碰硬碟、不碰虛擬機）
  python rescue_dd.py --selftest        不需虛擬機、不需硬碟的完整流程測試

rescue2.py 提供 NTFS 解析、build_plan、finalize_ready/deliver、write_xlsx；這裡只負責「讀映像檔」和「指揮虛擬機裡的 ddrescue」。
"""
import argparse, bisect, csv, ctypes, os, re, shutil, sqlite3, subprocess, sys, tempfile, time
import rescue2 as R

# ===== 可調參數 =====
VM_MAC = "00:50:56:2a:0d:e1"
LEASES = r"C:\ProgramData\VMware\vmnetdhcp.leases"
ASKPASS = r"C:\PhotoRescueVM\askpass.cmd"
SHARE_CFG = r"C:\PhotoRescueVM\share-password.txt"
DD_DIR = r"C:\PhotoRescueDD"                         # Windows 端；虛擬機裡掛在 VM_DD
VM_DD = "/mnt/dd"
IMG, MAP = "seagate.img", "seagate.map"
DEV_GLOB = "/dev/disk/by-id/usb-Seagate_Expansion_Desk_*"
CSV_PATH = r"C:\照片救援\壞軌紀錄.csv"
OTHER_TOPS = ["家人A", "家人B", "家人C", "手機備份", "game"]     # 「照片」以外要整個比對的第一層資料夾（software、字型、music、電子書使用者決定不救）
# 分批計畫（以下是這次實際用的 13 批，名稱已去識別化）。每個檔案歸到「前綴最長、最吻合」的那一批；
# ask=True 的批次輪到時要先問使用者。訂批次的原則見 docs/方法二_虛擬機ddrescue.md「分批計畫怎麼訂」。
BATCHES = [
    ("b01", "照片、家人C、手機備份", ["E:\\照片\\", "E:\\家人C\\", "E:\\手機備份\\"], False),
    ("b02", "家人A", ["E:\\家人A\\"], False),
    ("b03", "找不到所屬資料夾的檔案", ["E:\\照片救援_無法判斷資料夾\\"], False),
    ("b04", "家人B-1 個人小資料夾", ["E:\\家人B\\"], False),
    ("b05", "家人B-2 相機照片", ["E:\\家人B\\個人\\照片2020存入\\", "E:\\家人B\\個人\\照片2022存入\\", "E:\\家人B\\個人\\2018相機\\"], False),
    ("b06", "家人B-3 畢冊照片 第47、48屆", ["E:\\家人B\\畢冊照片\\第47屆\\", "E:\\家人B\\畢冊照片\\第48屆\\"], False),
    ("b07", "家人B-4 畢冊照片其餘", ["E:\\家人B\\畢冊照片\\"], False),
    ("b08", "家人B-5 個人其餘", ["E:\\家人B\\個人\\"], False),
    ("b09", "家人B-6 舊電腦資料、202304", ["E:\\家人B\\2019電腦資料\\", "E:\\家人B\\202304\\"], False),
    ("b10", "家人B-7 20240828、20241113、20260120", ["E:\\家人B\\20240828\\", "E:\\家人B\\20241113\\", "E:\\家人B\\20260120\\"], False),
    ("b11", "家人B-8 筆電與手機備份",
     ["E:\\家人B\\20240108\\", "E:\\家人B\\筆電D槽\\", "E:\\家人B\\手機檔案\\", "E:\\家人B\\20250124\\", "E:\\家人B\\20240418\\"], False),
    ("b12", "家人B-9 錄音", ["E:\\家人B\\個人\\music\\"], False),
    ("b13", "game", ["E:\\game\\"], True),
]
BATCH_NAME = {bid: name for bid, name, _, _ in BATCHES}
SECTOR = 4096
DD_BASE = "-d -b 4096 -c 16 --mapfile-interval=60"   # --mapfile-interval：ddrescue 預設依 mapfile 大小自動決定存檔間隔，9/17 b02 實測
                                                     # 6 MB 的 mapfile 開始後 10 分鐘才第一次存，Windows 回報一直顯示 0% → 固定每 60 秒存一次
                                                     # 直接讀、4K 磁區、一次 64 KiB。9/17 虛擬機實測：搭配 --domain-mapfile 時
                                                     # -c 256（1 MiB）會把整段好資料留在「待修剪」不讀，-c 64 也會漏；-c 16/32 正常
COPY_OPTS = "-n -e +50 -T 10m"                       # 第一輪：不刮取；壞區太多或 10 分鐘讀不到就結束，交給這裡判斷
RETRY_OPTS = "-A -r1 -T 10m"                         # 補讀：-A 把「待修剪/待刮取」改回未讀再試（實測沒有 -A 時 ddrescue 不會回頭處理）
READ_MIN, REST_MIN = 120, 60
POLL_SEC = 30
MIN_FREE_GB = 50                                     # C 槽（映像檔）每 5 分鐘輸出並釋放後仍剩不到 50 GB → 停止讀取、找 Opus
G_MIN_FREE_GB, G_RESUME_FREE_GB = 30, 60             # G 槽（Google Drive 暫存在 D 槽）剩不到 30 GB 暫停，等上傳釋出到 60 GB
SPACE_WAIT_SEC = 300
EXTRACT_GROUP = 200                                  # 輸出時每 200 個檔案檢查一次 G 槽空間
R.DB = os.path.join(R.WORK, "dd.sqlite")
R.TMP = os.path.join(R.WORK, "dd_tmp")
# ====================


# ---------- mapfile ----------
def read_map(path, tries=5):
    """ddrescue mapfile → [(pos, size, status)]。ddrescue 會定期整份重寫，讀到寫一半的內容就重讀。"""
    for _ in range(tries):
        try:
            lines = [l.split() for l in open(path, encoding="ascii").read().splitlines() if l.strip() and not l.startswith("#")]
        except FileNotFoundError:
            return []
        try:
            blocks = [(int(p, 0), int(s, 0), st) for p, s, st in lines[1:]]      # 第 1 行是 current_pos 狀態列
            if all(blocks[i][0] + blocks[i][1] == blocks[i + 1][0] for i in range(len(blocks) - 1)):
                return blocks
        except (ValueError, IndexError):
            pass
        time.sleep(1)
    raise RuntimeError(f"mapfile 格式不對或一直在寫入中：{path}")


def map_pos(path):
    """ddrescue 目前讀到硬碟的哪個位置（位元組），以及 mapfile 多久沒更新（秒）。
    看門狗用這兩個數字判斷是不是卡住了；讀不到就回傳 (None, None)。"""
    try:
        stale = time.time() - os.path.getmtime(path)
        for l in open(path, encoding="ascii"):
            l = l.strip()
            if l.startswith("0x"):                    # current_pos 那一行是第一個 0x 開頭的
                return int(l.split()[0], 16), stale
    except (OSError, ValueError, IndexError):
        pass
    return None, None


def write_map(path, blocks, comment="rescue_dd"):
    with open(path, "w", encoding="ascii", newline="\n") as f:
        f.write(f"# Mapfile. Created by {comment}\n# current_pos  current_status  current_pass\n0x00000000     ?               1\n"
                "#      pos        size  status\n")
        for pos, size, st in blocks:
            f.write(f"0x{pos:08X}  0x{size:08X}  {st}\n")


class MapIndex:
    """查某個範圍在 mapfile 裡的狀態；mapfile 沒涵蓋的地方算 '?'（沒讀過）。"""
    def __init__(self, blocks):
        self.blocks, self.starts = blocks, [b[0] for b in blocks]

    def segments(self, off, n):
        out, end, pos = [], off + n, off
        i = max(bisect.bisect_right(self.starts, off) - 1, 0)
        while pos < end:
            if i < len(self.blocks) and self.blocks[i][0] <= pos < self.blocks[i][0] + self.blocks[i][1]:
                b_end = self.blocks[i][0] + self.blocks[i][1]
                st = self.blocks[i][2]
                i += 1
            else:
                nxt = self.starts[i] if i < len(self.starts) and self.starts[i] > pos else end
                b_end, st = nxt, "?"
            seg_end = min(b_end, end)
            if out and out[-1][2] == st:
                out[-1] = (out[-1][0], seg_end, st)
            else:
                out.append((pos, seg_end, st))
            pos = seg_end
        return out


def domain_blocks(ranges, total):
    """[(off, n)] → 對齊 4K、合併後的 domain mapfile 內容（要讀的標 '+'，其他 '?'）與要讀的位元組數。"""
    merged = []
    for off, n in sorted(ranges):
        a, b = off - off % SECTOR, min(-(-(off + n) // SECTOR) * SECTOR, total)
        if merged and a <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], b)
        elif b > a:
            merged.append([a, b])
    blocks, pos = [], 0
    for a, b in merged:
        if a > pos:
            blocks.append((pos, a - pos, "?"))
        blocks.append((a, b - a, "+"))
        pos = b
    if pos < total:
        blocks.append((pos, total - pos, "?"))
    return blocks, sum(b - a for a, b in merged)


def tally(domain_path, map_path):
    """domain 範圍內的位元組數：'+' 讀到、'?' 沒讀、'retry' 第一輪跳過還能補讀（* /）、'-' 補讀後確定壞；'bad' = retry + '-'。"""
    idx = MapIndex(read_map(map_path))
    t = {"+": 0, "?": 0, "retry": 0, "-": 0}
    for pos, size, st in read_map(domain_path):
        if st != "+":
            continue
        for a, b, s in idx.segments(pos, size):
            t[s if s in "+?-" else "retry"] += b - a
    t["bad"] = t["retry"] + t["-"]
    return t


# ---------- 映像檔 ----------
class NotYetRead(Exception):
    def __init__(self, off, n):
        super().__init__(f"位移 {off} 長 {n} 還沒讀過")
        self.off, self.n = off, n


class ImageReader:
    """給 rescue2.Volume 用的讀取器：以 mapfile 為準。沒讀過 → NotYetRead；讀過但壞 → rescue2.ReadError；稀疏檔的洞絕不當成資料。"""
    def __init__(self, img, map_path):
        self.img, self.map_path, self.f = img, map_path, open(img, "rb")
        self.reload()

    def reload(self):
        self.idx = MapIndex(read_map(self.map_path))

    def read(self, off, n):
        sts = {s for _, _, s in self.idx.segments(off, n)}
        if sts == {"+"}:
            self.f.seek(off)
            return self.f.read(n).ljust(n, b"\0")
        if "?" in sts:
            raise NotYetRead(off, n)
        raise R.ReadError(-5, off, n)

    def close(self):
        self.f.close()


def ensure_sparse_image(path, size):
    """建立（或確認）稀疏映像檔。絕不縮小或重建已存在的映像檔。注意：不能用 Python truncate()，它會真的寫零。"""
    k = ctypes.WinDLL("kernel32", use_last_error=True)
    if os.path.exists(path):
        if not k.GetFileAttributesW(path) & 0x200:
            raise RuntimeError(f"{path} 已存在但不是稀疏檔，停止（避免佔滿 C 槽）")
        if os.path.getsize(path) < size:
            raise RuntimeError(f"{path} 大小 {os.path.getsize(path)} 小於硬碟 {size}，停止")
        return
    import ctypes.wintypes as W, msvcrt
    k.DeviceIoControl.argtypes = [ctypes.c_void_p, W.DWORD, ctypes.c_void_p, W.DWORD, ctypes.c_void_p, W.DWORD, ctypes.POINTER(W.DWORD), ctypes.c_void_p]
    k.SetFilePointerEx.argtypes = [ctypes.c_void_p, ctypes.c_longlong, ctypes.c_void_p, W.DWORD]
    k.SetEndOfFile.argtypes = [ctypes.c_void_p]
    with open(path, "xb", buffering=0) as f:
        h, ret = msvcrt.get_osfhandle(f.fileno()), W.DWORD()
        if not k.DeviceIoControl(h, 0x900C4, None, 0, None, 0, ctypes.byref(ret), None):        # FSCTL_SET_SPARSE
            raise OSError(ctypes.get_last_error(), "無法設定稀疏檔")
        if not (k.SetFilePointerEx(h, size, None, 0) and k.SetEndOfFile(h)):
            raise OSError(ctypes.get_last_error(), "無法設定映像檔大小")


def allocated_bytes(path):
    import ctypes.wintypes as W
    k = ctypes.WinDLL("kernel32")
    k.GetCompressedFileSizeW.restype = W.DWORD
    hi = W.DWORD()
    lo = k.GetCompressedFileSizeW(path, ctypes.byref(hi))
    return (hi.value << 32) | lo


def zero_ranges(path, ranges):
    """把稀疏映像檔裡的範圍歸還給 C 槽（FSCTL_SET_ZERO_DATA）。不是刪檔：映像檔還在，mapfile 也不動。
    回傳成功釋放的範圍；打不開映像檔回傳 None（例如共用資料夾正鎖著）。"""
    import ctypes.wintypes as W

    class ZeroData(ctypes.Structure):
        _fields_ = [("FileOffset", ctypes.c_longlong), ("BeyondFinalZero", ctypes.c_longlong)]
    k = ctypes.WinDLL("kernel32", use_last_error=True)
    k.CreateFileW.restype = ctypes.c_void_p
    k.CreateFileW.argtypes = [W.LPCWSTR, W.DWORD, W.DWORD, ctypes.c_void_p, W.DWORD, W.DWORD, ctypes.c_void_p]
    k.DeviceIoControl.argtypes = [ctypes.c_void_p, W.DWORD, ctypes.c_void_p, W.DWORD, ctypes.c_void_p, W.DWORD, ctypes.POINTER(W.DWORD), ctypes.c_void_p]
    k.CloseHandle.argtypes = [ctypes.c_void_p]
    h = k.CreateFileW(path, 0x40000000, 7, None, 3, 0, None)          # GENERIC_WRITE, share read/write/delete, OPEN_EXISTING
    if h in (None, ctypes.c_void_p(-1).value):
        return None
    done, ret = [], W.DWORD()
    try:
        for a, b in ranges:
            zd = ZeroData(a, b)
            if k.DeviceIoControl(h, 0x980C8, ctypes.byref(zd), ctypes.sizeof(zd), None, 0, ctypes.byref(ret), None):
                done.append((a, b))
    finally:
        k.CloseHandle(h)
    return done


def free_bytes(path):
    return shutil.disk_usage(path).free


def batch_of(path, batches=None):
    """檔案屬於哪一批：前綴最長者勝（例如 家人B\\畢冊照片\\第47屆… 歸 b06，其餘畢冊照片歸 b07）。都不符合回傳 None。"""
    pl, best, best_len = path.lower(), None, -1
    for bid, _, prefixes, _ in (batches or BATCHES):
        for p in prefixes:
            if len(p) > best_len and pl.startswith(p.lower()):
                best, best_len = bid, len(p)
    return best


# ---------- 虛擬機 ----------
def vm_ip():
    try:
        text = open(LEASES, encoding="ascii", errors="ignore").read()
    except OSError:
        return None
    ips = [m.group(1) for m in re.finditer(r"lease (\d+\.\d+\.\d+\.\d+) \{([^}]*)\}", text) if VM_MAC in m.group(2).lower()]
    return ips[-1] if ips else None


def ssh(ip, cmd, timeout=60):
    """回傳 (結束代碼, 輸出)。逾時回傳 (None, '')：虛擬機可能卡住。"""
    env = dict(os.environ, SSH_ASKPASS=ASKPASS, SSH_ASKPASS_REQUIRE="force", DISPLAY="x")
    # 虛擬機每次從光碟開機都會換 SSH 主機金鑰；記住舊金鑰會讓 ssh 拒絕密碼登入。每次連線前清空（虛擬機只在本機 NAT 網段）
    known = os.path.join(os.path.dirname(ASKPASS), "known_hosts_vm")
    open(known, "w").close()
    args = ["ssh", "-o", "StrictHostKeyChecking=no", "-o", f"UserKnownHostsFile={known}", "-o", "LogLevel=ERROR",
            "-o", "ConnectTimeout=10", "-o", "ServerAliveInterval=15", f"root@{ip}", cmd]
    try:
        p = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout, env=env,
                           stdin=subprocess.DEVNULL)
        return p.returncode, (p.stdout + p.stderr).strip()
    except subprocess.TimeoutExpired:
        return None, ""


def host_sees_drive():
    """Windows 的儲存驅動抓到這顆硬碟了嗎（應該要 False：它只能在虛擬機裡）。查不到（逾時）也當成有問題。"""
    ps = "if (Get-CimInstance Win32_DiskDrive | Where-Object Model -like '*Seagate Expansion*') { 'YES' } else { 'NO' }"
    out = R.ps_run(ps, 30)
    return out != "NO"


class VmRunner:
    """在虛擬機裡執行 ddrescue。輸入一定是 /dev/disk/by-id 的 Seagate，輸出一定在 SMB 掛載的 /mnt/dd。"""
    def __init__(self):
        self.ip = vm_ip()
        if not self.ip:
            raise SystemExit("找不到虛擬機的 IP（虛擬機沒開？）")
        rc, out = ssh(self.ip, "echo ok", 20)
        if rc != 0:
            raise SystemExit("SSH 連不進虛擬機。請在虛擬機畫面打：dmesg -n 1  以及  echo root:rescue | chpasswd; iptables -P INPUT ACCEPT; iptables -F")
        self.mount()
        self.dev, self.size = self.find_device()

    # cache=none：關掉 cifs 的客戶端寫入快取，讓 ddrescue 的寫入直接穿透到 Windows。
    # 9/20 實測：預設的 cache=strict 會把映像檔的寫入堆在快取裡，硬碟卡住或行程被強制
    # 終止時就 flush 不出去，24 小時累積 11,885 次「延遲寫入失敗」（Event 141），
    # Modified Page List 一度漲到 13.5 GB 把記憶體吃光。這顆碟只有 5~25 MB/s，
    # 遠低於 SMB 的能力，關快取實際上不會變慢。
    MOUNT_OPTS = ("credentials=/root/.smbcred,vers=3.1.1,cache=none,"
                  "uid=0,gid=0,file_mode=0644,dir_mode=0755")

    def mount(self):
        cfg = dict(l.split("=", 1) for l in open(SHARE_CFG, encoding="ascii").read().split())
        cmd = (f"umask 077; printf 'username={cfg['user']}\\npassword={cfg['password']}\\n' > /root/.smbcred; mkdir -p {VM_DD}; "
               f"mountpoint -q {VM_DD} || timeout 60 mount -t cifs //{cfg['host']}/{cfg['share']} {VM_DD} "
               f"-o {self.MOUNT_OPTS}; findmnt -n -o FSTYPE {VM_DD}")
        rc, out = ssh(self.ip, cmd)
        if out.strip().splitlines()[-1:] != ["cifs"]:
            raise SystemExit(f"虛擬機掛載共用資料夾失敗：{out}")

    def find_device(self):
        cmd = (f"for d in {DEV_GLOB}; do case $d in *-part*) ;; *) [ -e \"$d\" ] && echo $d;; esac; done")
        rc, out = ssh(self.ip, cmd)
        devs = [l for l in out.splitlines() if l.startswith("/dev/disk/by-id/usb-Seagate")]
        if len(devs) != 1:
            raise SystemExit("虛擬機裡找不到 Seagate 硬碟（沒接上、或還沒自動接進虛擬機）" if not devs else f"找到不只一顆：{devs}")
        dev = devs[0]
        # 9/17 實測：這顆碟通電後要 7 分鐘才「準備好」，期間回報 Not Ready。Linux 預設 30 秒逾時會一直重置它，改成 180 秒
        rc, out = ssh(self.ip, f"echo 180 > /sys/block/$(basename $(readlink -f {dev}))/device/timeout; "
                               f"blockdev --setro {dev} && blockdev --getro {dev} && blockdev --getss {dev} && blockdev --getsize64 {dev}", 600)
        vals = out.split()
        if rc != 0 or len(vals) != 3 or vals[0] != "1" or vals[1] != str(SECTOR):
            raise SystemExit(f"硬碟唯讀或磁區大小不對，停止：{out}")
        return dev, int(vals[2])

    def _cmdline(self, domain, opts):
        assert self.dev.startswith("/dev/disk/by-id/usb-Seagate") and "-part" not in self.dev
        return (f"ddrescue {DD_BASE} {opts} --domain-mapfile={VM_DD}/{domain} {self.dev} {VM_DD}/{IMG} {VM_DD}/{MAP}")

    def _preflight(self, domain):
        rc, out = ssh(self.ip, f"findmnt -n -o FSTYPE {VM_DD}; blockdev --getro {self.dev}")
        if out.split() != ["cifs", "1"]:              # 沒掛載的話，映像檔會寫進虛擬機記憶體把它塞爆
            raise SystemExit(f"開始前檢查失敗（共用資料夾沒掛載或硬碟不是唯讀）：{out}")
        # 9/17 實測：Windows 剛改寫的 domain 檔，虛擬機經 SMB 可能還讀到舊快取（被截斷），ddrescue 會回報 Bad block line。
        # 確認虛擬機讀到的內容和 Windows 寫的一模一樣才執行。
        import hashlib
        want = hashlib.sha256(open(os.path.join(DD_DIR, domain), "rb").read()).hexdigest()
        for _ in range(10):
            rc, out = ssh(self.ip, f"sha256sum {VM_DD}/{domain} | cut -d' ' -f1", 30)
            if out.strip() == want:
                return
            ssh(self.ip, "sync; echo 1 > /proc/sys/vm/drop_caches; sleep 2", 30)
        raise SystemExit(f"虛擬機讀到的 {domain} 跟 Windows 寫的不一致（共用資料夾快取），停止")

    def run(self, domain, opts, timeout=3600):
        self._stop_leftover()
        self._preflight(domain)
        say(f"虛擬機執行：{self._cmdline(domain, opts)}")
        rc, out = ssh(self.ip, self._cmdline(domain, opts) + f" > {VM_DD}/ddrescue_last.log 2>&1; echo EXIT=$?", timeout)
        return rc

    def _stop_leftover(self):
        """控制程式當掉重跑時，虛擬機裡上一次的 ddrescue 可能還在跑；兩個同時寫同一份 mapfile 會弄壞它 → 先正常停掉（9/17 b01 遇到）。"""
        if self.running():
            say("虛擬機裡還有上一次的 ddrescue 在跑，先正常停止（進度會寫回 mapfile）")
            self.stop()
            if self.running() is not False:
                raise SystemExit("虛擬機裡的 ddrescue 停不下來或虛擬機沒有回應，沒有啟動新的一輪。請切換 Opus 處理")

    def start(self, domain, opts):
        self._stop_leftover()
        self._preflight(domain)
        say(f"虛擬機背景執行：{self._cmdline(domain, opts)}")
        ssh(self.ip, f"setsid nohup {self._cmdline(domain, opts)} > {VM_DD}/ddrescue_run.log 2>&1 < /dev/null &")

    def running(self):
        rc, out = ssh(self.ip, "pgrep -x ddrescue", 30)
        return None if rc is None else rc == 0

    def stop(self, wait=600):
        ssh(self.ip, "pkill -INT -x ddrescue", 30)
        t0 = time.time()
        while self.running() and time.time() - t0 < wait:
            time.sleep(5)

    def probe(self, off):
        rc, _ = ssh(self.ip, f"timeout 30 dd if={self.dev} of=/dev/null bs={SECTOR} count=1 skip={off // SECTOR} iflag=direct status=none", 60)
        return rc == 0

    def drive_present(self):
        """硬碟裝置還在不在（純看裝置節點，不去讀它，以免加重卡死）。
        回傳 True=還在、False=掉了、None=虛擬機沒回應。
        這是分流的依據：裝置還在通常等一等會自己好；裝置掉了（USB error -110）只有斷電重開能救。"""
        rc, out = ssh(self.ip, f"[ -e {self.dev} ] && echo YES || echo NO", 30)
        if rc is None:
            return None
        return out.strip().endswith("YES")

    def status_line(self):
        cmd = (f"echo dev=$(cat /sys/block/$(basename $(readlink -f {self.dev}))/device/state 2>/dev/null || echo 不見了); "
               "dmesg | grep -ciE 'reset (high|super)-speed|device offlined|I/O error'")
        rc, out = ssh(self.ip, cmd, 30)
        return "虛擬機沒有回應" if rc is None else out.replace("\n", "，USB重置/錯誤訊息累計 ")


# ---------- 回報 ----------
S = dict(stage="", state="", note="")
PREV = {"t": time.time(), "good": 0}


def say(msg):
    R.say(msg)


def report_text(c, domain=None):
    rows = dict(c.execute("SELECT status, COUNT(*) FROM files GROUP BY status").fetchall())
    L = [f"【照片搬移回報】{R.now()}   ddrescue 虛擬機救援 {S['stage']}", f"狀態：{S['state'] or '—'}"]
    mp = os.path.join(DD_DIR, MAP)
    if domain and os.path.exists(os.path.join(DD_DIR, domain)):
        t = tally(os.path.join(DD_DIR, domain), mp)
        tot = t["+"] + t["?"] + t["bad"] or 1
        dt = max(time.time() - PREV["t"], 1)
        rate = (t["+"] - PREV["good"]) / dt
        eta = f"，照這 {dt / 60:.0f} 分鐘的速度還要 {t['?'] / rate / 3600:.1f} 小時" if rate > 0 and t["?"] else ""
        L.append(f"範圍：共 {tot / 1e9:.2f} GB，已讀到 {t['+'] / 1e9:.2f} GB（{t['+'] * 100 / tot:.1f}%），"
                 f"壞區 {t['bad'] / 1e6:.1f} MB（其中待補讀 {t['retry'] / 1e6:.1f} MB），還沒讀 {t['?'] / 1e9:.2f} GB；速度 {rate / 1e6:.1f} MB/s{eta}")
        # 9/20：只看「速度 0.0」會誤判——這顆碟讀到資料稀疏的區段時速度本來就是 0，但位置一直在前進。
        # 把「讀到哪」和「多久沒動」一起寫出來，看的人不必再進虛擬機挖也能分辨卡死或正常。
        pos, stale = map_pos(mp)
        if pos is not None:
            L.append(f"進度檔：目前讀到硬碟 {pos / 1e9:.1f} GB 的位置；"
                     + (f"已 {stale / 60:.0f} 分鐘沒更新（正常是每分鐘更新一次，超過 {STALL_MIN} 分鐘要查）"
                        if stale > 90 else f"{stale:.0f} 秒前才更新過，讀取正常進行中"))
        PREV.update(t=time.time(), good=t["+"])
    L.append(f"檔案：待處理 {rows.get('new', 0)}，已完整救回 {rows.get('done', 0)}，部分損壞已存 {rows.get('partial_saved', 0)}，"
             f"損壞未存 {rows.get('damaged', 0)}，不支援 {rows.get('unsupported', 0)}，之前已在 G 槽 {rows.get('done_before', 0)}")
    m = re.match(r"domain_(b\d+)\.map$", domain or "")
    if m:
        n_bad = len(bad_files(c, m.group(1)))
        L.append(f"壞軌：這一批目前有 {n_bad} 個檔案碰到壞區（清單：python rescue_dd.py bad）" if n_bad else "壞軌：這一批目前沒有檔案碰到壞區")
    if S["note"]:
        L.append(f"備註：{S['note']}")
    return "\n".join(L)


def bad_files(c, batch):
    """這一批讀取中已知碰到壞區的檔案 [(路徑, 大小, 壞區位元組)]（最後補讀前的暫時結果）。"""
    return c.execute("""SELECT f.path, f.size, (SELECT SUM(nbytes) FROM chunks WHERE file_id=f.id AND state='bad')
                        FROM files f WHERE f.batch=? AND f.status='new'
                        AND EXISTS(SELECT 1 FROM chunks WHERE file_id=f.id AND state='bad') ORDER BY f.path""", (batch,)).fetchall()


def emit(c, domain=None):
    text = report_text(c, domain)
    print("\n" + "=" * 70 + "\n" + text + "\n" + "=" * 70, flush=True)
    R.log(text)
    try:
        with open(R.REPORT_TXT, "w", encoding="utf-8-sig") as f:
            f.write(text + "\n")
    except OSError:
        pass


# ---------- R1：救 MFT ----------
def run_settled(runner, domain, opts, timeout):
    """跑一次；範圍內還有 ddrescue 沒處理的「待修剪/待刮取」就加 -A 再補一次。只用在 MBR/MFT 這種小而關鍵的範圍。"""
    runner.run(domain, opts, timeout=timeout)
    if tally(os.path.join(DD_DIR, domain), os.path.join(DD_DIR, MAP))["retry"]:
        say(f"{domain} 還有待修剪的區塊，加 -A 補讀一次")
        runner.run(domain, "-A -r1", timeout=timeout)


def attrlist_extents(vol):
    """非常駐 $ATTRIBUTE_LIST 的資料位置 [(裝置位移, 位元組數)]（只取到實際大小）。"""
    out = []
    for r in vol.records.values():
        al = r["attrlist"]
        if not al or al["res"]:
            continue
        left = al["size"]
        for off, n in vol.runs_to_extents(al["runs"]):
            if left <= 0:
                break
            if off is not None:
                out.append((off, min(n, left)))
            left -= n
    return out


def cmd_mft(runner, c):
    S["stage"] = "R1 救 MFT"
    img, mp = os.path.join(DD_DIR, IMG), os.path.join(DD_DIR, MAP)
    ensure_sparse_image(img, runner.size)
    R.meta_set(c, "disk_size", runner.size)
    rd = ImageReader(img, mp)
    for _ in range(20):                                    # MBR → boot sector → MFT 第 0 筆記錄（與屬性清單）：缺什麼補讀什麼
        try:
            vol = R.Volume(rd, R.find_partition(rd, SECTOR), SECTOR)
            runs = vol.find_mft_runs()
            break
        except NotYetRead as e:
            blocks, n = domain_blocks([(e.off, max(e.n, 1 << 20))], runner.size)
            write_map(os.path.join(DD_DIR, "domain_meta.map"), blocks)
            S["state"] = f"讀取目錄的中繼資料（位移 {e.off}）"
            say(S["state"])
            run_settled(runner, "domain_meta.map", "-r1", timeout=1800)
            rd.reload()
            if {s for _, _, s in rd.idx.segments(e.off, e.n)} == {"?"}:
                raise SystemExit(f"ddrescue 沒有讀位移 {e.off}（看 {DD_DIR}\\ddrescue_last.log）")
    else:
        raise SystemExit("MFT 位置找了 20 次還找不到")
    ext = [(off, n) for off, n in vol.runs_to_extents(runs) if off is not None]
    blocks, n = domain_blocks(ext, runner.size)
    write_map(os.path.join(DD_DIR, "domain_mft.map"), blocks)
    if tally(os.path.join(DD_DIR, "domain_mft.map"), mp)["?"]:
        S["state"] = f"讀取整個 MFT（{n / 1e6:.0f} MB）"
        say(S["state"])
        run_settled(runner, "domain_mft.map", "-n", timeout=6 * 3600)
    else:                                                   # 重跑 mft 時不要再花幾小時重試 MFT 壞區（先救好的；壞區留給 R4）
        say(f"MFT（{n / 1e6:.0f} MB）之前已經讀過，略過")
    part = vol.part_off
    for _ in range(3):
        # 9/17 實測：碎片很多的大檔，$ATTRIBUTE_LIST 放不進 MFT 記錄，另外存在 MFT 以外的叢集 → 讀完 MFT 後補讀這些位置再解析
        rd.reload()
        vol = R.Volume(rd, part, SECTOR)                  # 每次重新建：_link_extensions 會改動記錄，不能對同一份重跑
        try:
            vol.load_mft()
            break
        except NotYetRead:
            ext_al = attrlist_extents(vol)
            blocks, n_al = domain_blocks(ext_al, runner.size)
            write_map(os.path.join(DD_DIR, "domain_attrlist.map"), blocks)
            S["state"] = f"補讀 {len(ext_al)} 份存在 MFT 以外的屬性清單（{n_al / 1e6:.1f} MB）"
            say(S["state"])
            run_settled(runner, "domain_attrlist.map", "-r1", timeout=6 * 3600)
    else:
        raise SystemExit("屬性清單補讀 3 次仍有沒讀到的位置（看 ddrescue_last.log）")
    t = tally(os.path.join(DD_DIR, "domain_mft.map"), mp)
    msg = (f"MFT {n / 1e6:.0f} MB：讀到 {t['+'] / 1e6:.0f} MB，壞區 {t['bad'] / 1e6:.1f} MB，沒讀 {t['?'] / 1e6:.1f} MB；"
           f"記錄 {len(vol.records)} 筆，讀不到/損毀 {len(vol.bad_recs)} 筆")
    R.meta_set(c, "mft", msg)
    say(msg)
    rd.close()
    return vol


# ---------- R2：規劃 ----------
def make_keep(csv_path=CSV_PATH, others=OTHER_TOPS):
    """範圍（使用者 9/17 決定）：其他資料夾整個比對；照片只挑壞軌紀錄裡「待救援」的檔案與資料夾；孤兒檔一律列入。
    沒有方法一的壞軌紀錄（換一顆硬碟、直接用方法二）時，照片也整個比對，G 槽已有的照樣會跳過。"""
    files, dirs = set(), []
    if not os.path.exists(csv_path):
        others = ["照片"] + list(others)
    else:
        with open(csv_path, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                if r["狀態"] != "待救援":
                    continue
                p = r["來源完整路徑"].rstrip("\\").lower()
                (files.add(p) if r["類型"] == "檔案" else dirs.append(p + "\\"))
    prefixes = tuple(["e:\\" + t.lower() + "\\" for t in others] + ["e:\\" + R.ORPHAN_DIR.lower() + "\\"] + dirs)

    def keep(path):
        pl = path.lower()
        return pl in files or pl.startswith(prefixes)
    return keep


def ensure_columns(c):
    """files 表加上 batch（屬於哪一批）與 freed（映像檔裡的資料已釋放）；再建立故障追蹤用的表。"""
    for sql in ("ALTER TABLE files ADD COLUMN batch TEXT", "ALTER TABLE files ADD COLUMN freed INT DEFAULT 0"):
        try:
            c.execute(sql)
        except sqlite3.OperationalError:
            pass
    # 9/20：硬碟在某些位置一讀就整顆鎖死，必須記下來跨執行沿用，否則每次重啟都再撞一次
    c.executescript("""
    CREATE TABLE IF NOT EXISTS hostile(start INT, end INT, reason TEXT, at TEXT);
    CREATE TABLE IF NOT EXISTS incidents(at TEXT, kind TEXT, pos INT, note TEXT);
    """)
    c.commit()


# ---------- 危險區與故障事件 ----------
HOSTILE_MARGIN = 5 * 10**9          # 掉線位置前後各留這麼多當緩衝
STALL_MIN = 5                       # mapfile 幾分鐘沒更新就當成停滯，開始查原因
HANG_WAIT_MIN = 45                  # 硬碟還在、只是卡住 → 最多等這麼久（實測常常會自己恢復）
PROGRESS_MIN_BYTES = 10**6          # 看門狗：讀到的新資料超過這麼多才算「有進度」。補讀一次只刮 4 KB，
                                    # 45 分鐘刮不到 1 MB 的位置，繼續刮只是讓硬碟一直卡死（9/24 實測 3 小時只多 230 KB）
WAIT_DRIVE_MIN = 30                 # 硬碟不見了 → 等人來重新上電，最多等這麼久
NEED_HELP = os.path.join(R.WORK, "需要處理.txt")


def hostile_zones(c):
    return [(a, b) for a, b in c.execute("SELECT start, end FROM hostile ORDER BY start")]


def add_hostile(c, start, end, reason):
    """記下一段會害硬碟鎖死的位置；跟既有區段重疊就合併。"""
    start, end = max(0, int(start)), int(end)
    for a, b in hostile_zones(c):
        if a <= end and start <= b:                       # 有重疊 → 併成一段
            start, end = min(a, start), max(b, end)
            c.execute("DELETE FROM hostile WHERE start=? AND end=?", (a, b))
    c.execute("INSERT INTO hostile VALUES(?,?,?,?)", (start, end, reason, R.now()))
    c.commit()
    say(f"已記錄危險區 {start / 1e9:.0f}～{end / 1e9:.0f} GB（{reason}），之後的批次會跳過，留到最後補讀")


def add_incident(c, kind, pos, note=""):
    c.execute("INSERT INTO incidents VALUES(?,?,?,?)", (R.now(), kind, int(pos or 0), note))
    c.commit()


def recent_incidents(c, kind, minutes):
    """最近這段時間內同類型故障的次數（用來判斷同一個位置是不是重複掛掉）。"""
    cutoff = time.strftime("%Y/%m/%d %H:%M:%S", time.localtime(time.time() - minutes * 60))
    return c.execute("SELECT COUNT(*), MAX(pos) FROM incidents WHERE kind=? AND at>=?", (kind, cutoff)).fetchone()


def need_help(c, domain, why):
    """寫一份「要人處理」的旗標檔，讓監控的人／AI 一看就知道發生什麼事、該做什麼。"""
    text = f"{R.now()}\n{why}\n\n{report_text(c, domain)}\n"
    try:
        with open(NEED_HELP, "w", encoding="utf-8-sig") as f:
            f.write(text)
    except OSError:
        pass
    say(why)


def clear_need_help():
    try:
        if os.path.exists(NEED_HELP):
            open(NEED_HELP, "w", encoding="utf-8-sig").close()
    except OSError:
        pass


def cmd_plan(c, keep=None, disk_size=None, batches=None):
    S["stage"] = "R2 規劃"
    ensure_columns(c)
    if c.execute("SELECT COUNT(*) FROM files WHERE status IN ('done','partial_saved')").fetchone()[0]:
        # 已輸出的檔案在映像檔裡的資料可能已經釋放；重新規劃會把它們當成「讀到了」再輸出一次全是 0 的檔案
        raise SystemExit("已經開始輸出檔案，不能重新規劃（映像檔裡已輸出的資料可能已釋放）")
    img, mp = os.path.join(DD_DIR, IMG), os.path.join(DD_DIR, MAP)
    disk_size = disk_size or int(R.meta_get(c, "disk_size", "0"))
    if not disk_size:
        raise SystemExit("還沒跑過 mft")
    rd = ImageReader(img, mp)
    vol = R.Volume(rd, R.find_partition(rd, SECTOR), SECTOR)
    try:
        vol.load_mft()
    except NotYetRead as e:
        raise SystemExit(f"MFT 還有沒讀過的部分（{e}），請先跑 mft")
    st = R.build_plan(vol, c, want_tops=["照片"] + OTHER_TOPS, keep=keep or make_keep())
    rows = c.execute("SELECT id, path FROM files").fetchall()
    c.executemany("UPDATE files SET batch=?, freed=0 WHERE id=?", [(batch_of(p, batches), fid) for fid, p in rows])
    same = []                                                # G 槽已有同名同大小、只是修改時間不同（之前用別的方式搬過）→ 不再讀
    for fid, dst, size in c.execute("SELECT id, dst, size FROM files WHERE status='new'").fetchall():
        try:
            if os.stat(R.lp(dst)).st_size == size:
                same.append((fid,))
        except OSError:
            pass
    c.executemany("UPDATE files SET status='done_before', note='G 槽已有同名同大小（修改時間不同）' WHERE id=?", same)
    c.executemany("DELETE FROM chunks WHERE file_id=?", same)
    # 兩個檔案用到同一段磁區（硬連結等）：任一個輸出後都不能釋放那段，否則另一個會輸出成全是 0 → 兩邊都標 freed=-1 永不釋放
    shared, end, owner = set(), -1, None
    for fid, off, n in c.execute("SELECT file_id, dev_off, nbytes FROM chunks WHERE dev_off>=0 ORDER BY dev_off").fetchall():
        if off < end and fid != owner:
            shared.update((fid, owner))
        if off + n > end:
            end, owner = off + n, fid
    c.executemany("UPDATE files SET freed=-1 WHERE id=?", [(f,) for f in shared])
    c.commit()
    per = {b: (n, s or 0) for b, n, s in c.execute("SELECT batch, COUNT(*), SUM(size) FROM files WHERE status='new' GROUP BY batch")}
    total = sum(s for _, s in per.values())
    lines = [f"===== ddrescue 規劃 {R.now()} =====", R.meta_get(c, "mft", ""),
             f"範圍內檔案 {st['total']}：之前已在 G 槽 {st['done_before'] + len(same)}（其中同大小、時間不同 {len(same)}），"
             f"要救 {st['to_read'] - len(same)}，不支援/記錄異常 {st['unsupported']}，找不到所屬資料夾 {st['orphan']}",
             f"要救合計 {total / 1e9:.1f} GB。分批讀取、邊讀邊搬到 G 槽邊釋放映像檔，C 槽剩 {free_bytes(DD_DIR) / 1e9:.0f} GB",
             "批次 | 內容 | 要救檔案 | 要救 GB"]
    for bid, name, _, ask in (batches or BATCHES):
        n_f, s = per.get(bid, (0, 0))
        lines.append(f"  {bid} | {name}{'（輪到時先問使用者）' if ask else ''} | {n_f} | {s / 1e9:.2f}")
    if None in per:
        lines.append(f"  （不屬於任何批次，不會讀）| | {per[None][0]} | {per[None][1] / 1e9:.2f}")
    if shared:
        lines.append(f"共用磁區的檔案 {len(shared)} 個：輸出後不釋放映像檔空間")
    n = total
    text = "\n".join(lines)
    say(text)
    with open(os.path.join(R.WORK, "ddrescue盤點.txt"), "w", encoding="utf-8-sig") as f:
        f.write(text + "\n")
    R.meta_set(c, "plan", text)
    rd.close()
    return st, n


# ---------- R3/R4：讀取 ----------
DONE_BATCHES = "SELECT substr(key, 12) FROM meta WHERE key LIKE 'batch_done_%'"


def next_batch(c, batches=None):
    """照順序第一個還沒讀完的批次 → (id, 名稱, 要不要先問)；全部讀完回傳 (None, None, False)。"""
    for bid, name, _, ask in (batches or BATCHES):
        if not R.meta_get(c, "batch_done_" + bid):
            return bid, name, ask
    return None, None, False


def mark_unread(map_path, ranges):
    """把 mapfile 裡這些範圍從 '+' 改回 '?'（對齊磁區，其他狀態不動），回傳改了多少位元組。會先備份 mapfile。"""
    marks, prev = [], None
    for off, n in sorted((o - o % SECTOR, -(-(o + n) // SECTOR) * SECTOR) for o, n in ranges):
        if prev and off <= prev[1]:
            prev[1] = max(prev[1], n)
        else:
            prev = [off, n]
            marks.append(prev)
    blocks, out, k, changed = read_map(map_path), [], 0, 0
    for pos, size, st in blocks:
        cur, end = pos, pos + size
        while cur < end:
            while k < len(marks) and marks[k][1] <= cur:
                k += 1
            if k < len(marks) and marks[k][0] < end:
                a, b = max(marks[k][0], cur), min(marks[k][1], end)
                if a > cur:
                    out.append((cur, a - cur, st))
                new = "?" if st == "+" else st
                changed += (b - a) if new != st else 0
                out.append((a, b - a, new))
                cur = b
            else:
                out.append((cur, end - cur, st))
                cur = end
    merged = []
    for pos, size, st in out:
        if merged and merged[-1][2] == st and merged[-1][0] + merged[-1][1] == pos:
            merged[-1] = (merged[-1][0], merged[-1][1] + size, st)
        else:
            merged.append((pos, size, st))
    if changed:
        shutil.copy2(map_path, map_path + ".bak-mark")
        write_map(map_path, merged, comment="rescue_dd mark-unextracted-as-unread")   # mapfile 只吃 ASCII
    return changed


def limit_zones(c, map_path, zones):
    """retry 分兩段（9/24 使用者決定）：先補讀安全的壞區，會讓硬碟整顆鎖死的危險區最後才碰。
    zones='safe' 從讀取範圍扣掉危險區；'danger' 只留危險區；None 不動。回傳扣掉的位元組。"""
    hz = hostile_zones(c)
    if zones == "safe":
        cut = [(a, b - a) for a, b in hz]
    elif zones == "danger":                          # 危險區之間的空隙全部扣掉（hostile 已合併、排序，彼此不重疊）
        edges = [0] + [x for z in hz for x in z] + [int(R.meta_get(c, "disk_size", "0"))]
        cut = [(a, b - a) for a, b in zip(edges[::2], edges[1::2]) if b > a]
    else:
        return 0
    return mark_unread(map_path, cut) if cut else 0


def recheck_unextracted(c, batch, sleep=time.sleep):
    """9/19 事故：硬碟出狀況時 Windows 出現 1169 筆 seagate.img「延遲寫入失敗」，那段資料在 mapfile 是 '+'、映像檔裡卻可能是空的。
    每次開始讀一批前，把「已標記讀到、但還沒輸出到 G 槽」的區段標回未讀重新讀（只有輸出到 G 槽的資料才算數）。"""
    ranges = c.execute("""SELECT ch.dev_off, ch.nbytes FROM chunks ch JOIN files f ON f.id=ch.file_id
                          WHERE ch.state='new' AND ch.dev_off>=0 AND f.status='new' AND f.batch=?""", (batch,)).fetchall()
    if not ranges:
        return 0
    n = mark_unread(os.path.join(DD_DIR, MAP), ranges)
    if n:
        say(f"上次讀到但還沒輸出到 G 槽的 {n / 1e9:.2f} GB，重新讀一次（避免當機/寫入失敗留下空資料）")
    return n


def write_domain(c, name, where, args=()):
    ranges = c.execute(f"SELECT ch.dev_off, ch.nbytes FROM chunks ch JOIN files f ON f.id=ch.file_id WHERE ch.dev_off>=0 AND {where}", args).fetchall()
    blocks, n = domain_blocks(ranges, int(R.meta_get(c, "disk_size", "0")))
    write_map(os.path.join(DD_DIR, name), blocks)
    return n


def cmd_copy(runner, c, retry=False, batch=None, sleep=time.sleep, batches=None, skip=None, folder=None, zones=None):
    """copy：讀一批（沒指定就照順序下一批），讀完就結束等使用者確認下一批。
    folder：改成只讀某個資料夾底下還沒救到的檔案，不受分批限制（硬碟快不行時，用來小範圍搶救）。
    retry：所有批次讀完後，一起補讀待處理檔案的壞區。"""
    ensure_columns(c)
    if not R.meta_get(c, "plan"):                              # 看資料庫，不只看檔案：共用資料夾可能留著演練的 domain 檔
        raise SystemExit("還沒跑過 plan")
    # 9/21 踩到的坑：這個檢查原本只有 cmd_mft 有。映像檔被刪掉後沒再跑過 mft，
    # ddrescue 就自己在 SMB 上建了一個「普通檔案」。非稀疏檔寫到磁碟位置 2744 GB，
    # NTFS 會從 0 開始實際配置並填零，半小時吃掉 511 GB 把 C 槽塞到只剩 69 GB。
    # copy 每次開始前都要確認，不能假設 mft 階段留下的檔案還在、還是稀疏的。
    ensure_sparse_image(os.path.join(DD_DIR, IMG), runner.size)
    if retry:
        label = {"safe": "（危險區以外）", "danger": "（只讀危險區）"}.get(zones, "（所有批次一起）")
        domain, opts, S["stage"] = "domain_retry.map", RETRY_OPTS, "R4 補讀壞區" + label
        # 壞區（bad）＋讀取批次時留下沒讀到的密集壞軌區（new），只限已讀完的批次（例如使用者沒同意的 game 不會被讀）
        write_domain(c, domain, f"ch.state IN ('bad','new') AND f.status='new' AND f.batch IN ({DONE_BATCHES})")
        n_cut = limit_zones(c, os.path.join(DD_DIR, domain), zones)
        if n_cut:
            say(f"這一輪{label}：讀取範圍扣掉 {n_cut / 1e9:.2f} GB")
    elif folder:
        # 9/20：硬碟嚴重衰弱時，整批 100 GB 讀不完。改成一個資料夾一個資料夾搶，
        # 讀完就停，讓人看結果再決定下一個。狀態仍照常寫回資料庫，之後 retry 一樣會補讀。
        like = folder.rstrip("\\") + "\\%"
        n = c.execute("SELECT COUNT(*) FROM files WHERE path LIKE ? AND status='new'", (like,)).fetchone()[0]
        if not n:
            raise SystemExit(f"{folder} 底下沒有還沒救到的檔案")
        domain, opts = "domain_folder.map", COPY_OPTS
        S["stage"] = f"R3 資料夾：{folder}（{n} 個檔案）"
        write_domain(c, domain, "f.path LIKE ? AND f.status NOT IN ('done_before','unsupported')", (like,))
        zones = hostile_zones(c)
        if zones:
            n_skip = mark_unread(os.path.join(DD_DIR, domain), [(a, b - a) for a, b in zones])
            say(f"這一輪跳過 {len(zones)} 個危險區共 {n_skip / 1e9:.2f} GB，留到最後補讀")
    else:
        if batch is None:
            batch, name, ask = next_batch(c, batches)
            if batch is None:
                raise SystemExit("所有批次都讀完了。下一步：retry（補讀壞區）")
            if ask:
                raise SystemExit(f"輪到 {batch}（{name}）：這一批要先問使用者，同意後執行 copy --batch {batch}")
        name = {b[0]: b[1] for b in (batches or BATCHES)}.get(batch)
        if not name:
            raise SystemExit(f"沒有 {batch} 這個批次")
        domain, opts, S["stage"] = f"domain_{batch}.map", COPY_OPTS, f"R3 第 {batch} 批：{name}"
        recheck_unextracted(c, batch)
        write_domain(c, domain, "f.batch=? AND f.status NOT IN ('done_before','unsupported')", (batch,))
        if skip:                     # --skip 是人工指定，一併存進資料庫，之後每一批都自動跳過
            a, b = skip
            add_hostile(c, a, b, "人工指定 --skip")
        # 9/20：危險區＝實測會讓硬碟整顆鎖死的位置。從讀取範圍扣掉，但區塊維持「沒讀過」，
        # 所以最後的 retry 階段（條件是 state IN ('bad','new')）一定會回頭補讀，不會漏掉資料。
        zones = hostile_zones(c)
        if zones:
            n_skip = mark_unread(os.path.join(DD_DIR, domain), [(a, b - a) for a, b in zones])
            say(f"這一輪跳過 {len(zones)} 個危險區共 {n_skip / 1e9:.2f} GB（"
                + "、".join(f"{a / 1e9:.0f}~{b / 1e9:.0f}GB" for a, b in zones) + "），留到最後補讀")
    R.meta_set(c, "cur_domain", domain)
    dp, mp = os.path.join(DD_DIR, domain), os.path.join(DD_DIR, MAP)
    finished = lambda t: t["?"] == 0 and (not retry or t["retry"] == 0)
    last_report, prev, unproductive = 0.0, tally(dp, mp), 0
    clear_need_help()
    # 看門狗：上次讀到新資料是什麼時候。9/24 補讀實測的兩個教訓：
    # (1) 要跨 ddrescue 重啟累計：每輪「讀到卡點 → 硬碟卡死 → ddrescue 結束 → 重啟」都不到 45 分鐘，計時每輪歸零就永遠不會跳過；
    # (2) 要看讀到的資料量，不能看讀取位置：-A 重啟時位置會先跳回範圍開頭再回到卡點，看起來像「有前進」。
    # 兩次加起來在同一個卡點空轉約 6 小時、讓硬碟每半小時卡死一次。讀取位置只拿來記「卡在哪裡」。
    got, moved_at, moved_pos = prev["+"], time.time(), map_pos(mp)[0]
    while not finished(prev):
        runner.start(domain, opts)
        S["state"], started, paused = "執行中", time.time(), False
        while True:
            sleep(POLL_SEC)
            if host_sees_drive():
                runner.stop(60)
                S["state"] = "★★★ 硬碟掉回 Windows 了，請立刻拔掉硬碟的 USB ★★★（已停止讀取）"
                add_incident(c, "掉回Windows", map_pos(mp)[0] or 0)
                need_help(c, domain, "硬碟從虛擬機掉出來、被 Windows 抓到了。請拔掉硬碟 USB，"
                                     "再依序：關電源 → 等 10 秒 → 開電源 → 立刻在 VMware 點 Connect 接回虛擬機")
                emit(c, domain)
                raise SystemExit(S["state"])
            alive = runner.running()
            if alive is None:
                S["state"] = "虛擬機沒有回應（可能卡住）。請檢查 VMware 視窗；映像檔與進度都在 Windows 上，不會遺失"
                need_help(c, domain, S["state"])
                emit(c, domain)
                raise SystemExit(S["state"])
            # ---- 看門狗 ----
            # ddrescue 自己的 -T 10m 在硬碟卡死時不會觸發：整支程式被凍在不可中斷的 I/O 等待（D 狀態），
            # 連自己的計時器都執行不了（9/20 實測卡了 2.5 小時毫無反應）。控制程式在虛擬機外面不會被凍住，由它判斷。
            if alive:
                moved_pos = map_pos(mp)[0] or moved_pos
                now_got = tally(dp, mp)["+"]
                if now_got - got >= PROGRESS_MIN_BYTES:
                    got, moved_at = now_got, time.time()
                stuck = (time.time() - moved_at) / 60
                if stuck >= STALL_MIN:
                    present = runner.drive_present()
                    if present is False:
                        # 硬碟從匯流排上掉了（USB error -110）。只有人工斷電重開能救，繼續探測沒有意義。
                        runner.stop(60)
                        add_incident(c, "硬碟掉線", moved_pos or 0, f"停滯 {stuck:.0f} 分鐘後發現裝置不見")
                        S["state"] = f"硬碟掉線了（讀到 {(moved_pos or 0) / 1e9:.1f} GB 處），已停止讀取"
                        need_help(c, domain, f"硬碟從虛擬機的匯流排上掉了（停在 {(moved_pos or 0) / 1e9:.1f} GB）。"
                                             "這種情況只有斷電重開能救，程式不會再空轉。請：關硬碟電源 → 等 10 秒 → "
                                             "開電源 → 立刻在 VMware 點 VM → Removable Devices → Connect，然後重新啟動自動接續")
                        emit(c, domain)
                        raise SystemExit(S["state"])
                    if stuck >= HANG_WAIT_MIN:
                        # 硬碟還在，只是這個位置讀不過去。等夠久了還是不動 → 記成危險區，跳過它繼續讀別的地方。
                        runner.stop(120)
                        at = moved_pos or 0
                        add_incident(c, "位置卡死", at, f"卡住 {stuck:.0f} 分鐘")
                        add_hostile(c, at - HOSTILE_MARGIN, at + HOSTILE_MARGIN, f"讀到此處卡住超過 {HANG_WAIT_MIN} 分鐘")
                        mark_unread(dp, [(max(0, at - HOSTILE_MARGIN), 2 * HOSTILE_MARGIN)])
                        S["state"] = f"硬碟在 {at / 1e9:.1f} GB 卡住 {stuck:.0f} 分鐘，跳過這一段繼續讀其他位置"
                        emit(c, domain)
                        paused = True
                        break
                    S["state"] = (f"已 {stuck:.0f} 分鐘沒讀到新資料（讀取位置 {(moved_pos or 0) / 1e9:.1f} GB），硬碟還在，"
                                  f"繼續等（這顆碟常常會自己恢復；累計滿 {HANG_WAIT_MIN} 分鐘才會跳過這一段）")
            if time.time() - last_report >= R.REPORT_EVERY_SEC:
                S["note"] = runner.status_line()
                paused = step_extract(c, runner, domain, sleep)       # 邊讀邊搬到 G 槽、邊釋放映像檔
                emit(c, domain)
                last_report = time.time()
            if paused or not alive:
                break
            if time.time() - started >= READ_MIN * 60:
                runner.stop()
                step_extract(c, runner, domain, sleep)
                S["state"] = f"讀滿 {READ_MIN} 分鐘，休息 {REST_MIN} 分鐘（自動繼續）"
                emit(c, domain)
                sleep(REST_MIN * 60)
                paused = True
                break
        t = tally(dp, mp)
        if paused:                                     # 休息、等 G 槽、剛跳過一段：這段時間不算卡住
            prev, got, moved_at = t, t["+"], time.time()
            continue
        if finished(t):
            break
        waited = wait_drive(runner, c, domain, sleep)  # ddrescue 自己結束（壞區太多、10 分鐘讀不到、或硬碟被拔掉）→ 先確認硬碟還能讀
        if not waited and t == prev and retry:        # ddrescue 結束了但什麼都沒改變：絕不無限重跑（9/17 演練曾每 4 秒重啟一次）
            S["state"] = (f"ddrescue 已結束且沒有任何進度，停止這一輪。範圍內還沒讀 {t['?'] / 1e6:.1f} MB、"
                          f"待補讀 {t['retry'] / 1e6:.1f} MB。請 Claude 看 {DD_DIR}\\ddrescue_run.log 判斷原因")
            say(S["state"])
            emit(c, domain)
            return
        # 9/17 b01 實測：最後 307 MB 是密集壞軌，每讀一次卡好幾分鐘還會讓硬碟卡死。第一輪讀取時硬碟正常卻連續兩輪讀到不滿 1 MB
        # → 這一批先收尾，剩下的留給所有批次讀完後的 retry（使用者選的「壞的最後一起補讀」）
        unproductive = 0 if waited or t["+"] - prev["+"] >= 1e6 else unproductive + 1
        prev = t
        if not retry and unproductive >= 2:
            say(f"剩下 {t['?'] / 1e6:.1f} MB 連續兩輪幾乎讀不到（密集壞軌），這一批先收尾，留到全部批次讀完後一起補讀")
            break
        say(f"ddrescue 結束但範圍還沒讀完（沒讀 {t['?'] / 1e6:.1f} MB），硬碟有回應，繼續")
    if retry:                                          # 補讀過的壞區重新依 mapfile 切成好/壞，讀到的部分寫進暫存
        c.execute("UPDATE chunks SET state='new' WHERE state='bad' AND file_id IN (SELECT id FROM files WHERE status='new')")
        c.commit()
        cmd_extract(c, final=False, sleep=sleep)
        S["state"] = "這一輪已完成"
        return batch_report(c, "retry", domain)
    close_batch(c, batch, sleep, domain, folder)


def close_batch(c, batch, sleep=time.sleep, domain=None, folder=None):
    """一批結束：輸出、記錄讀完、寫批次報告。沒讀到的部分留給 retry。
    batch 為 None 代表 --folder 的臨時範圍：照常輸出與寫報告，但不標記 batch_done
    （那是分批計畫的進度，資料夾模式不該動它，否則 next_batch 會誤判）。
    folder 給定時 batch_report 改用路徑篩選，而不是拿 bid 當假批次代號去篩
    （9/20 的 bug：篩不到任何一筆，報告整份顯示 0）。"""
    cmd_extract(c, final=False, sleep=sleep, xlsx=False)
    if batch:
        R.meta_set(c, "batch_done_" + batch, R.now())
    write_bad_xlsx(c)
    S["state"] = "這一輪已完成"
    return batch_report(c, batch or "資料夾", domain or f"domain_{batch}.map", folder=folder)


def wait_drive(runner, c, domain, sleep=time.sleep):
    """探測一個確定讀得到的位置；讀不到就請使用者拔插並等它回來。回傳 True = 硬碟曾經卡死/不見（這輪沒進度不能怪壞軌）。"""
    good = next(((p, s) for p, s, st in read_map(os.path.join(DD_DIR, MAP)) if st == "+"), None)
    waited, t0 = False, time.time()
    while good and not runner.probe(good[0]):
        waited = True
        # 9/20：這裡原本是無限迴圈，硬碟真的死掉時會一直空轉探測（實測空轉 2.5 小時沒人知道）。
        # 現在等滿就停下來叫人，並把原因寫進旗標檔。
        if time.time() - t0 >= WAIT_DRIVE_MIN * 60:
            add_incident(c, "等不到硬碟", good[0], f"等了 {WAIT_DRIVE_MIN} 分鐘")
            need_help(c, domain, f"等了 {WAIT_DRIVE_MIN} 分鐘硬碟還是讀不到，停止等待。"
                                 "請：關硬碟電源 → 等 10 秒 → 開電源 → 立刻在 VMware 點 Connect，然後重新啟動自動接續")
            emit(c, domain)
            raise SystemExit("等不到硬碟，已停止")
        S["state"] = (f"硬碟卡死或沒接上（連確定讀得到的位置都讀不到，已等 {(time.time() - t0) / 60:.0f}/{WAIT_DRIVE_MIN} 分鐘）。"
                      "請：關硬碟電源 → 等 10 秒 → 開電源 → "
                      "立刻在 VMware 選單 VM → Removable Devices → Seagate Expansion Desk → Connect（接到虛擬機，不能留在 Windows）"
                      "。程式每分鐘檢查，恢復後自己繼續")
        if host_sees_drive():
            S["state"] = "★ 硬碟現在接在 Windows 上：請立刻在 VMware 選單 VM → Removable Devices → Seagate Expansion Desk → Connect ★"
        emit(c, domain)
        sleep(60)
    if waited:
        runner.dev, runner.size = runner.find_device()             # 重新接上是一顆「新」裝置：重新設唯讀與 180 秒逾時
        say("硬碟恢復了，已重新設為唯讀，繼續")
    return waited


def batch_report(c, bid, domain, folder=None):
    """bid='retry' 不篩；folder 給定（--folder 模式，bid 不是真正的批次代號）就用路徑篩；
    否則照批次代號篩。9/20 曾經誤把 folder 模式的 bid 當成批次代號傳進 SQL，
    batch='資料夾' 這種字串在資料表裡從不存在，篩出來全是 0，報告整份是假的。"""
    if bid == "retry":
        where, args = "", ()
    elif folder:
        where, args = " AND path LIKE ?", (folder.rstrip("\\") + "\\%",)
    else:
        where, args = " AND batch=?", (bid,)
    rows = dict(c.execute("SELECT status, COUNT(*) FROM files WHERE 1=1" + where + " GROUP BY status", args).fetchall())
    n_bad = c.execute("SELECT COUNT(*) FROM files f WHERE status='new'" + where +
                      " AND EXISTS(SELECT 1 FROM chunks WHERE file_id=f.id AND (state='bad' OR (state='new' AND dev_off>=0)))", args).fetchone()[0]
    t = tally(os.path.join(DD_DIR, domain), os.path.join(DD_DIR, MAP))
    nxt, name, ask = next_batch(c)
    title = f"資料夾：{folder}" if folder else ('補讀壞區' if bid == 'retry' else '第 ' + bid + ' 批：' + BATCH_NAME.get(bid, ''))
    text = "\n".join([
        f"===== {title} 結束 {R.now()} =====",
        f"讀取範圍：讀到 {t['+'] / 1e9:.2f} GB，壞區 {t['bad'] / 1e6:.1f} MB，沒讀 {t['?'] / 1e6:.1f} MB",
        f"檔案：已完整救回 {rows.get('done', 0)}，有壞區或沒讀到、待最後補讀 {n_bad}，其他未完成 {rows.get('new', 0) - n_bad}，"
        f"部分損壞已存 {rows.get('partial_saved', 0)}，救不回來 {rows.get('damaged', 0)}，寫入 G 失敗 {rows.get('dst_error', 0)}，"
        f"不支援 {rows.get('unsupported', 0)}，原檔空白未存 {rows.get('empty', 0)}，之前已在 G 槽 {rows.get('done_before', 0)}",
        f"壞檔清單：{R.XLSX} 的「{BAD_SHEET}」工作表（依批次篩選）",
        f"C 槽剩 {free_bytes(DD_DIR) / 1e9:.0f} GB，G 槽剩 {free_bytes(R.DST_ROOT) / 1e9:.0f} GB，映像檔實際佔用 {allocated_bytes(os.path.join(DD_DIR, IMG)) / 1e9:.1f} GB",
        ("下一步：所有批次讀完 → retry" if nxt is None else f"下一批：{nxt}（{name}）{'——這批要先問使用者' if ask else ''}，使用者確認後執行 copy"),
    ])
    say(text)
    with open(os.path.join(R.WORK, f"ddrescue批次報告_{bid}.txt"), "w", encoding="utf-8-sig") as f:
        f.write(text + "\n")
    return text


# ---------- R5：輸出與釋放 ----------
def wait_g_space(c, on_wait=None, sleep=time.sleep):
    """G 槽（Google Drive 暫存）不夠就先停讀取，等上傳釋出空間。"""
    if free_bytes(R.DST_ROOT) >= G_MIN_FREE_GB * 1e9:
        return
    if on_wait:
        on_wait()
    while free_bytes(R.DST_ROOT) < G_RESUME_FREE_GB * 1e9:
        S["state"] = (f"G 槽只剩 {free_bytes(R.DST_ROOT) / 1e9:.0f} GB（低於 {G_MIN_FREE_GB} GB），暫停讀取與輸出，"
                      f"等 Google Drive 上傳釋出空間到 {G_RESUME_FREE_GB} GB（每 {SPACE_WAIT_SEC // 60} 分鐘檢查，會自己繼續）")
        emit(c)
        sleep(SPACE_WAIT_SEC)
    S["state"] = "G 槽空間恢復，繼續"
    say(S["state"])


def step_extract(c, runner, domain, sleep=time.sleep):
    """讀取中的定期工作：輸出＋釋放；空間不足時停下 ddrescue。回傳 True = ddrescue 被停下了（外層要重新啟動）。"""
    stopped = []

    def stop():
        if not stopped:
            runner.stop()
            stopped.append(1)
    cmd_extract(c, on_wait=stop, sleep=sleep, xlsx=False)
    if free_bytes(DD_DIR) < MIN_FREE_GB * 1e9:
        stop()
        S["state"] = (f"C 槽剩不到 {MIN_FREE_GB} GB，已輸出並釋放仍不夠（多半是很多大檔讀到一半），已停止讀取。"
                      "請切換 Opus 處理；進度都保留")
        emit(c, domain)
        raise SystemExit(S["state"])
    return bool(stopped)


def need_to_free(c):
    """還需要靠打洞省空間嗎？C 槽放得下剩餘待救量就不必打。

    9/20 踩到的坑：打洞是為了原始的 1.8 TB 規模設計的（C 槽只有 931 GB，不邊讀邊
    釋放一定爆），但它原本無條件執行。每打一個洞就在稀疏映像檔裡多切出片段，
    實測累積到 1,929,562 個片段（93 萬個洞）後撞上 NTFS 片段對應表的上限，
    之後所有寫入與打洞全部失敗——一天 11,885 次「延遲寫入失敗」，
    髒頁塞爆 13.5 GB 記憶體。空間夠的時候完全不要打洞。"""
    left = c.execute("SELECT SUM(size) FROM files WHERE status='new'").fetchone()[0] or 0
    return free_bytes(DD_DIR) < left * 1.5 + MIN_FREE_GB * 1e9


def free_delivered(c):
    """已搬到 G 槽的檔案：把映像檔裡它讀到的資料歸還給 C 槽。共用磁區的檔案（freed=-1）不動。
    空間充足時只更新記帳、不實際打洞，避免把映像檔切碎（見 need_to_free）。"""
    rows = c.execute("""SELECT ch.dev_off, ch.nbytes FROM files f JOIN chunks ch ON ch.file_id=f.id
                        WHERE f.status IN ('done','partial_saved') AND f.freed=0 AND ch.state='good' AND ch.dev_off>=0""").fetchall()
    ranges = [(o, o + n) for o, n in rows]
    if ranges and not need_to_free(c):
        c.execute("UPDATE files SET freed=1 WHERE status IN ('done','partial_saved') AND freed=0")
        c.commit()
        return 0
    if ranges:
        ok = zero_ranges(os.path.join(DD_DIR, IMG), ranges)
        if ok is None or len(ok) != len(ranges):
            say(f"釋放映像檔空間沒有全部成功（{0 if ok is None else len(ok)}/{len(ranges)}），下次再試")
            return 0
    c.execute("UPDATE files SET freed=1 WHERE status IN ('done','partial_saved') AND freed=0")
    c.commit()
    return sum(b - a for a, b in ranges)


def cmd_extract(c, final=False, on_wait=None, sleep=time.sleep, xlsx=True):
    S["stage"] = S["stage"] or "R5 輸出" + ("（最終）" if final else "")
    ensure_columns(c)
    rd = ImageReader(os.path.join(DD_DIR, IMG), os.path.join(DD_DIR, MAP))
    rows = c.execute("SELECT id, file_id, foff, dev_off, nbytes FROM chunks WHERE state='new' AND dev_off>=0 ORDER BY file_id, foff").fetchall()
    sizes = {}
    for cid, fid, foff, dev_off, nbytes in rows:
        segs = rd.idx.segments(dev_off, nbytes)
        if any(s == "?" for _, _, s in segs):
            continue                                        # 還沒讀到，下次再輸出
        if fid not in sizes:
            sizes[fid] = c.execute("SELECT size FROM files WHERE id=?", (fid,)).fetchone()[0]
        c.execute("DELETE FROM chunks WHERE id=?", (cid,))
        for a, b, s in segs:
            if s == "+":
                rd.f.seek(a)
                R.write_tmp(fid, foff + (a - dev_off), rd.f.read(b - a), sizes[fid])
            c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(?,?,?,?,?)",
                      (fid, foff + (a - dev_off), a, b - a, "good" if s == "+" else "bad"))
        c.commit()
    rd.close()
    # 全部待處理的檔案都要檢查：內容就在 MFT 裡的小檔（resident）和空檔沒有磁區要讀
    # 有大小但硬碟上完全沒有內容（整個稀疏/未初始化＝全 0，9/17 實測多為沒下載完的 mp3）：不存到雲端，列進壞檔清單
    c.execute("""UPDATE files SET status='empty', note='原檔在硬碟上沒有內容（整個是 0，多半是當初沒下載完或沒存完），未存到雲端'
                 WHERE status='new' AND size>0 AND EXISTS(SELECT 1 FROM chunks WHERE file_id=files.id)
                 AND NOT EXISTS(SELECT 1 FROM chunks WHERE file_id=files.id AND state<>'zero')""")
    c.commit()
    fids = [r[0] for r in c.execute("SELECT id FROM files WHERE status='new'").fetchall()]
    for i in range(0, len(fids), EXTRACT_GROUP):
        wait_g_space(c, on_wait, sleep)
        R.finalize_ready(c, fids[i:i + EXTRACT_GROUP], final=final)
    freed = free_delivered(c)
    if xlsx:
        write_bad_xlsx(c)
    S["note"] = f"這次釋放映像檔 {freed / 1e9:.2f} GB；C 槽剩 {free_bytes(DD_DIR) / 1e9:.0f} GB，G 槽剩 {free_bytes(R.DST_ROOT) / 1e9:.0f} GB"
    if xlsx:
        S["state"] = "輸出完成"
        emit(c)


BAD_SHEET = "壞軌檔案清單"
BAD_NAMES = {"new": "待補讀", "partial_saved": "部分損壞已存", "damaged": "救不回來", "unsupported": "不支援", "dst_error": "寫入G失敗",
             "empty": "原檔空白未存"}


def bad_rows(c):
    return c.execute(f"""SELECT f.batch, f.path, f.dst, f.size, f.status, f.note,
                          CASE WHEN f.status='new' THEN (SELECT SUM(nbytes) FROM chunks WHERE file_id=f.id AND
                            (state IN ('bad','trimmed') OR (state='new' AND dev_off>=0))) ELSE f.bad_bytes END
                        FROM files f WHERE f.status IN ('damaged','partial_saved','unsupported','dst_error','empty')
                          OR (f.status='new' AND EXISTS(SELECT 1 FROM chunks WHERE file_id=f.id AND (state='bad'
                              OR (state='new' AND dev_off>=0 AND f.batch IN ({DONE_BATCHES})))))
                        ORDER BY f.batch, f.path""").fetchall()


def write_bad_xlsx(c):
    """每批結束更新：哪些路徑的檔案有壞區、結果如何。取代 rescue2 的「第二階段結果」（那張會列出所有未完成檔案，太大）。"""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return say("沒有安裝 openpyxl，壞檔清單無法寫入 Excel")
    fills = {"new": "FFF2CC", "partial_saved": "FCE4D6", "damaged": "FFC7CE", "unsupported": "FFC7CE", "dst_error": "FFF2CC", "empty": "D9D9D9"}
    try:
        wb = load_workbook(R.XLSX) if os.path.exists(R.XLSX) else Workbook()
        if not os.path.exists(R.XLSX):
            wb.remove(wb.active)
        if BAD_SHEET in wb.sheetnames:
            del wb[BAD_SHEET]
        ws = wb.create_sheet(BAD_SHEET, 0)
        ws.append(["批次", "來源路徑", "G槽路徑", "大小(MB)", "壞區(KB)", "壞區%", "結果", "說明"])
        for bid, path, dst, size, status, note, bb in bad_rows(c):
            bb = bb or 0
            ws.append([f"{bid} {BATCH_NAME.get(bid, '')}" if bid else "", path, dst if status == "partial_saved" else "",
                       round(size / 1048576, 2), round(bb / 1024, 1), round(bb * 100.0 / size, 2) if size else 0,
                       BAD_NAMES.get(status, status), note if status != "new" else ""])
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fills.get(status, "FFFFFF"))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, w in zip("ABCDEFGH", (22, 70, 70, 10, 10, 8, 14, 40)):
            ws.column_dimensions[col].width = w
        wb.save(R.XLSX)
    except Exception as e:
        S["note"] = f"壞檔清單 Excel 更新失敗（檔案是不是開著？）：{e}"
        say(S["note"])


def batches_text(c):
    per = {}
    for bid, st, n, s in c.execute("SELECT batch, status, COUNT(*), SUM(size) FROM files GROUP BY batch, status"):
        per.setdefault(bid, {})[st] = (n, s or 0)
    L = ["批次 | 內容 | 狀態 | 待處理 | 已救回 | 部分損壞已存 | 救不回來 | 之前已在G槽"]
    for bid, name, _, ask in BATCHES:
        d = per.get(bid, {})
        g = lambda k: d.get(k, (0, 0))[0]
        done = R.meta_get(c, "batch_done_" + bid)
        L.append(f"  {bid} | {name} | {'已讀完 ' + done if done else ('要先問使用者' if ask else '未讀')} | "
                 f"{g('new')}（{d.get('new', (0, 0))[1] / 1e9:.1f} GB）| {g('done')} | {g('partial_saved')} | {g('damaged')} | {g('done_before')}")
    return "\n".join(L)


# ---------- 自我測試 ----------
class FakeRunner:
    """用 Python 模擬 ddrescue：只讀 domain 裡的 4K 磁區，壞磁區第一輪標 '/'、補讀時標 '-'；-A 把 * / 改回 ?。
    stuck=True 模擬 ddrescue 結束卻什麼都沒做（9/17 真實演練遇到的情況），用來測試不會無限重跑。"""
    def __init__(self, src, bad, dd_dir):
        self.src, self.bad, self.dir, self.size, self.calls, self.stuck, self.never = src, bad, dd_dir, len(src), [], False, set()

    def _do(self, domain, opts):
        self.calls.append((domain, opts))
        if self.stuck:
            return 0
        idx = MapIndex(read_map(os.path.join(self.dir, MAP)))
        status = bytearray(b"?" * (self.size // SECTOR))
        for a, b, s in idx.segments(0, self.size):
            status[a // SECTOR:b // SECTOR] = s.encode() * ((b - a) // SECTOR)
        if "-A" in opts.split():
            status = bytearray(ord("?") if ch in b"*/" else ch for ch in status)
        with open(os.path.join(self.dir, IMG), "r+b") as img:
            for pos, size, st in read_map(os.path.join(self.dir, domain)):
                if st != "+":
                    continue
                for off in range(pos, pos + size, SECTOR):
                    if off in self.never:                       # 模擬密集壞軌：ddrescue 一直讀不到、也沒標壞就結束
                        continue
                    cur = chr(status[off // SECTOR])
                    if cur == "+" or (cur != "?" and "-n" in opts):
                        continue
                    if off in self.bad:
                        status[off // SECTOR] = ord("/" if "-n" in opts else "-")
                    else:
                        img.seek(off)
                        img.write(self.src[off:off + SECTOR])
                        status[off // SECTOR] = ord("+")
        blocks, i = [], 0
        while i < len(status):
            j = i
            while j < len(status) and status[j] == status[i]:
                j += 1
            blocks.append((i * SECTOR, (j - i) * SECTOR, chr(status[i])))
            i = j
        write_map(os.path.join(self.dir, MAP), blocks)
        return 0

    run = lambda self, domain, opts, timeout=0: self._do(domain, opts)
    start = lambda self, domain, opts: self._do(domain, opts)
    running = lambda self: False
    stop = lambda self, wait=0: None
    probe = lambda self, off: True
    status_line = lambda self: "（測試）"


def selftest():
    global DD_DIR, host_sees_drive, BATCHES, BATCH_NAME, free_bytes, NEED_HELP, need_to_free
    real_batches, real_free, real_need = BATCHES, free_bytes, need_to_free
    base = tempfile.mkdtemp(prefix="rescue_dd_")
    DD_DIR = os.path.join(base, "dd")
    os.makedirs(DD_DIR)
    R.WORK, R.TMP = os.path.join(base, "work"), os.path.join(base, "tmp")
    NEED_HELP = os.path.join(R.WORK, "需要處理.txt")     # 9/23：原本漏改，自我測試會清空正在救援的旗標檔
    R.DB, R.LOG = os.path.join(R.WORK, "dd.sqlite"), os.path.join(R.WORK, "log.txt")
    R.REPORT_TXT, R.XLSX, R.DST_RETRY_SLEEP = os.path.join(base, "回報.txt"), os.path.join(base, "清單.xlsx"), 0
    os.makedirs(R.WORK)
    dst_root = os.path.join(base, "G")
    real_dst = R.DST_ROOT
    R.DST_ROOT = dst_root
    host_sees_drive = lambda: False

    # 1. mapfile / domain / MapIndex 基本行為
    blocks, n = domain_blocks([(5000, 100), (4096 * 3, 10), (4096 * 2, 4096)], 4096 * 10)
    assert blocks == [(0, 4096, "?"), (4096, 4096 * 3, "+"), (4096 * 4, 4096 * 6, "?")] and n == 4096 * 3, blocks
    write_map(os.path.join(DD_DIR, "t.map"), blocks)
    assert read_map(os.path.join(DD_DIR, "t.map")) == blocks
    idx = MapIndex([(0, 100, "+"), (100, 50, "-")])
    assert idx.segments(90, 100) == [(90, 100, "+"), (100, 150, "-"), (150, 190, "?")], idx.segments(90, 100)

    class _V:                                               # 非常駐屬性清單的位置：只取實際大小、略過稀疏段與常駐清單
        records = {1: {"attrlist": dict(res=False, runs=[(10, 2), (None, 1), (20, 5)], size=4096 * 4 + 100)},
                   2: {"attrlist": dict(res=True, value=b"x")}, 3: {"attrlist": None}}
        runs_to_extents = staticmethod(lambda runs: [(None if l is None else l * 4096, n * 4096) for l, n in runs])
    assert attrlist_extents(_V()) == [(40960, 8192), (81920, 4096 + 100)], attrlist_extents(_V())

    # 2. 稀疏映像檔：1 TB 的檔案不能真的佔空間；已存在的絕不重建
    big = os.path.join(base, "sparse.img")
    ensure_sparse_image(big, 10 ** 12)
    assert os.path.getsize(big) == 10 ** 12 and allocated_bytes(big) < 10 ** 6, allocated_bytes(big)
    ensure_sparse_image(big, 10 ** 12)
    os.remove(big)

    # 3. 迷你 NTFS 當假硬碟，整條流程
    src, expect, bad, part = R._make_image()
    size = -(-len(src) // SECTOR) * SECTOR
    src = src.ljust(size, b"\0")
    ensure_sparse_image(os.path.join(DD_DIR, IMG), size)
    runner = FakeRunner(src, bad, DD_DIR)
    runner.size = size
    c = R.db_open()
    vol = cmd_mft(runner, c)
    assert any(d == "domain_meta.map" for d, _ in runner.calls)
    assert tally(os.path.join(DD_DIR, "domain_mft.map"), os.path.join(DD_DIR, MAP))["?"] == 0   # 迷你硬碟：找中繼資料時就讀完 MFT，MFT 那輪會被略過
    n_calls = len(runner.calls)
    cmd_mft(runner, c)                                          # 重跑 mft：已讀過的一律不再叫 ddrescue
    assert len(runner.calls) == n_calls, runner.calls[n_calls:]
    assert 40 in vol.bad_recs and 35 in vol.records                   # MFT 第 40 筆在壞磁區 → 孤兒
    # G 槽已有「已存在.jpg」
    ex = expect["已存在.jpg"]
    p = os.path.join(dst_root, "照片", "2015旅遊", "已存在.jpg")
    os.makedirs(os.path.dirname(p))
    open(p, "wb").write(ex["data"])
    os.utime(p, (ex["mtime"], ex["mtime"]))
    # 壞軌紀錄：照片只列這幾個檔；其他資料夾（家人A）整個比對
    csv_p = os.path.join(base, "壞軌紀錄.csv")
    with open(csv_p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["類型", "來源完整路徑", "狀態"])
        for name in ("a.jpg", "e.jpg", "d.nef", "big.mov", "已存在.jpg", "r.txt"):
            w.writerow(["檔案", f"E:\\照片\\2015旅遊\\{name}", "待救援"])
        w.writerow(["檔案", "E:\\照片\\2015旅遊\\b.txt", "已可讀取（已重新搬移）"])
    keep = make_keep(csv_p, others=["家人A"])
    assert keep("E:\\照片\\2015旅遊\\a.jpg") and not keep("E:\\照片\\2015旅遊\\b.txt")
    no_csv = make_keep(os.path.join(base, "沒有這個.csv"), others=["家人A"])      # 沒有方法一的紀錄：照片整個比對
    assert no_csv("E:\\照片\\2015旅遊\\b.txt") and no_csv("E:\\家人A\\x.jpg") and not no_csv("E:\\software\\x")
    BATCHES = [("t1", "照片", ["E:\\照片\\"], False), ("t2", "家人A與孤兒", ["E:\\家人A\\", "E:\\" + R.ORPHAN_DIR + "\\"], False),
               ("t3", "要先問", ["E:\\不存在\\"], True)]
    BATCH_NAME = {b[0]: b[1] for b in BATCHES}
    # G 槽已有同名同大小、修改時間不同的 c.jpg → 不再讀
    p = os.path.join(dst_root, "家人A", "c.jpg")
    os.makedirs(os.path.dirname(p))
    open(p, "wb").write(b"x" * expect["c.jpg"]["size"])
    st, n = cmd_plan(c, keep=keep, disk_size=size)
    planned = {r[0].rsplit("\\", 1)[1]: (r[1], r[2]) for r in c.execute("SELECT path, status, batch FROM files")}
    assert set(planned) == {"a.jpg", "e.jpg", "d.nef", "big.mov", "已存在.jpg", "r.txt", "c.jpg", "orphan.jpg"}, planned
    assert planned["已存在.jpg"][0] == "done_before" and planned["c.jpg"][0] == "done_before", planned
    assert planned["a.jpg"][1] == "t1" and planned["orphan.jpg"][1] == "t2", planned
    assert batch_of("E:\\家人B\\畢冊照片\\第47屆\\x.jpg", real_batches) == "b06"
    assert batch_of("e:\\家人B\\畢冊照片\\其他\\x.jpg", real_batches) == "b07" and batch_of("E:\\software\\x", real_batches) is None
    # 整個檔案都是稀疏/未初始化（全 0，沒有磁區要讀、不會有暫存檔）也要能輸出（9/17 b01 真實遇到 100 個 wav 讓程式當掉）
    zero_dst = os.path.join(dst_root, "家人A", "全0.wav")
    c.execute("INSERT INTO files(id,mft,path,size,mtime,atime,dst,status,note,bad_bytes,orphan,batch,freed) VALUES(9999,0,'E:\\家人A\\全0.wav',10000,?,?,?,'new','',0,0,'t2',0)",
              (ex["mtime"], ex["mtime"], zero_dst))
    c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(9999,0,-1,10000,'zero')")
    c.commit()

    # 第 1 批：讀取中 G 槽空間不足 → 停下 ddrescue 等上傳，恢復後繼續；讀完輸出並釋放映像檔
    write_map(os.path.join(DD_DIR, MAP), [(p, s, "?" if st == "+" else st) for p, s, st in read_map(os.path.join(DD_DIR, MAP))])
    g_low, stops = [2], []                          # ↑ 迷你硬碟讀 MFT 時就全讀完了：標回沒讀，讓讀取迴圈真的跑
    free_bytes = lambda path: (g_low.__setitem__(0, g_low[0] - 1) or 10e9) if path == dst_root and g_low[0] > 0 else 999e9
    # 9/21 起空間夠就不打洞，這裡強制要打，才測得到「已輸出的資料有被釋放」；
    # 「空間夠不打」的方向由 test_watchdog.py 的 test_need_to_free 測。9/23 才發現這個自我測試從 9/21 起一直失敗。
    need_to_free = lambda c: True
    runner.stop = lambda wait=0: stops.append(1)
    img = os.path.join(DD_DIR, IMG)
    before, alloc0 = len(runner.calls), allocated_bytes(img)
    cmd_copy(runner, c, sleep=lambda s: None)
    assert runner.calls[before] == ("domain_t1.map", COPY_OPTS), runner.calls[before:]
    assert stops and "等 Google Drive 上傳" in open(R.LOG, encoding="utf-8").read()
    status = {r[0].rsplit("\\", 1)[1]: r[1] for r in c.execute("SELECT path, status FROM files")}
    for name in ("a.jpg", "big.mov", "r.txt"):
        assert status[name] == "done", (name, status)
        row = c.execute("SELECT dst FROM files WHERE path LIKE ?", ("%\\" + name,)).fetchone()
        assert open(row[0], "rb").read() == expect[name]["data"], name
        assert abs(os.path.getmtime(row[0]) - expect[name]["mtime"]) <= 2
    assert status["e.jpg"] == "new" and status["d.nef"] == "new" and status["orphan.jpg"] == "new"   # 有壞區先不決定；第 2 批還沒讀
    assert status["全0.wav"] == "empty" and not os.path.exists(zero_dst), status
    assert not os.path.exists(os.path.join(dst_root, "照片", "2015旅遊", "b.txt")), "範圍外的檔案不能輸出"
    t = tally(os.path.join(DD_DIR, "domain_t1.map"), os.path.join(DD_DIR, MAP))
    assert t["?"] == 0 and t["bad"] > 0, t
    # 釋放：已輸出檔案在映像檔裡的資料變 0、實際佔用變小；還沒輸出的 e.jpg 不能動
    fid_big = c.execute("SELECT id FROM files WHERE path LIKE '%big.mov'").fetchone()[0]
    assert c.execute("SELECT freed FROM files WHERE id=?", (fid_big,)).fetchone()[0] == 1
    with open(img, "rb") as f:
        for off, nb in c.execute("SELECT dev_off, nbytes FROM chunks WHERE file_id=? AND state='good'", (fid_big,)).fetchall():
            f.seek(off)
            assert f.read(nb) == bytes(nb), "已輸出的檔案應該被釋放"
    assert allocated_bytes(img) < alloc0 or alloc0 == 0, (alloc0, allocated_bytes(img))
    assert R.meta_get(c, "batch_done_t1") and os.path.exists(os.path.join(R.WORK, "ddrescue批次報告_t1.txt"))
    names = {r[1].rsplit("\\", 1)[1]: BAD_NAMES[r[4]] for r in bad_rows(c)}
    assert names == {"e.jpg": "待補讀", "d.nef": "待補讀", "全0.wav": "原檔空白未存"}, names
    try:
        import openpyxl
        ws = openpyxl.load_workbook(R.XLSX)[BAD_SHEET]
        assert {r[1].value.rsplit("\\", 1)[1] for r in ws.iter_rows(min_row=2)} == {"e.jpg", "d.nef", "全0.wav"}
    except ImportError:
        print("（沒有 openpyxl，略過 Excel 檢查）")

    # 第 2 批、第 3 批（要先問）、全部讀完
    # 當機/寫入失敗後：已標記讀到但還沒輸出的區段要標回未讀重新讀（9/19 事故）
    fid_e = c.execute("SELECT id FROM files WHERE path LIKE '%e.jpg'").fetchone()[0]
    c.execute("UPDATE chunks SET state='new' WHERE file_id=? AND state='good'", (fid_e,))     # 假裝還沒輸出
    c.commit()
    rng = c.execute("SELECT dev_off, nbytes FROM chunks WHERE file_id=? AND state='new' AND dev_off>=0", (fid_e,)).fetchall()
    assert rng and recheck_unextracted(c, "t1") > 0
    idx2 = MapIndex(read_map(os.path.join(DD_DIR, MAP)))
    assert all(s != "+" for off, n in rng for _, _, s in idx2.segments(off, n)), "已讀到但沒輸出的區段要標回未讀"
    assert os.path.exists(os.path.join(DD_DIR, MAP + ".bak-mark"))
    runner.run("domain_t1.map", "-r1")                                                        # 重新讀一次就補回來
    cmd_extract(c, sleep=lambda s: None)
    assert c.execute("SELECT COUNT(*) FROM chunks WHERE file_id=? AND state='new'", (fid_e,)).fetchone()[0] == 0

    # 第 2 批：孤兒檔在密集壞軌裡一直讀不到，中間硬碟還卡死一次 → 等硬碟恢復（重設唯讀）後，連續兩輪讀不到就收尾，留給 retry
    fid_or = c.execute("SELECT id FROM files WHERE path LIKE '%orphan.jpg'").fetchone()[0]
    runner.never = {o for o, n in c.execute("SELECT dev_off, nbytes FROM chunks WHERE file_id=? AND dev_off>=0", (fid_or,))}
    probes, found = [False], []
    runner.probe = lambda off: probes.pop(0) if probes else True
    runner.find_device = lambda: found.append(1) or ("dev", size)
    n_calls = len(runner.calls)
    cmd_copy(runner, c, sleep=lambda s: None)
    assert len(runner.calls) - n_calls == 3 and found, (runner.calls[n_calls:], found)     # 卡死那輪不算、之後兩輪沒進度
    assert R.meta_get(c, "batch_done_t2") and c.execute("SELECT status FROM files WHERE id=?", (fid_or,)).fetchone()[0] == "new"
    assert "orphan.jpg" in {r[1].rsplit("\\", 1)[1] for r in bad_rows(c)}
    assert "沒讀到" in open(os.path.join(R.WORK, "ddrescue批次報告_t2.txt"), encoding="utf-8-sig").read()
    runner.never = set()
    for want, kw in (("先問使用者", {}), (None, dict(batch="t3")), ("所有批次都讀完了", {})):
        try:
            cmd_copy(runner, c, sleep=lambda s: None, **kw)
            assert want is None, want
        except SystemExit as e:
            assert want and want in str(e), (want, e)

    # C 槽輸出並釋放後仍不夠 → 停止
    free_bytes = lambda path: 1e9 if path == DD_DIR else 999e9
    try:
        step_extract(c, runner, "domain_t1.map")
        raise AssertionError("C 槽不夠應該停止")
    except SystemExit as e:
        assert "C 槽" in str(e)
    free_bytes = lambda path: 999e9

    # 補讀：ddrescue 結束卻沒進度 → 要自己停下來，不能無限重跑，壞區維持原狀
    runner.stuck, n_calls = True, len(runner.calls)
    blocks = read_map(os.path.join(DD_DIR, MAP))
    write_map(os.path.join(DD_DIR, MAP), [(p, s, "*" if st == "-" else st) for p, s, st in blocks])   # 做出「待修剪」讓補讀有事可做
    cmd_copy(runner, c, retry=True, sleep=lambda s: None)
    assert len(runner.calls) - n_calls == 1 and "沒有任何進度" in open(R.LOG, encoding="utf-8").read(), runner.calls[n_calls:]
    assert runner.calls[-1] == ("domain_retry.map", RETRY_OPTS)
    runner.stuck = False
    # 補讀（壞磁區仍然壞）→ 最終輸出
    cmd_copy(runner, c, retry=True, sleep=lambda s: None)
    assert runner.calls[-1] == ("domain_retry.map", RETRY_OPTS)
    t = tally(os.path.join(DD_DIR, "domain_retry.map"), os.path.join(DD_DIR, MAP))
    assert t["retry"] == 0 and t["-"] > 0 and t["?"] == 0, t
    cmd_extract(c, final=True)
    status = {r[0].rsplit("\\", 1)[1]: r[1] for r in c.execute("SELECT path, status FROM files")}
    assert status["e.jpg"] == "partial_saved" and status["d.nef"] == "damaged" and status["orphan.jpg"] == "done", status
    got = open(os.path.join(dst_root, "照片", "2015旅遊", "e_部分損壞.jpg"), "rb").read()
    exp = bytearray(expect["e.jpg"]["data"])
    exp[20 * SECTOR:21 * SECTOR] = bytes(SECTOR)
    assert got == bytes(exp), "e.jpg 只有壞的那個磁區補 0"
    names = {r[1].rsplit("\\", 1)[1]: BAD_NAMES[r[4]] for r in bad_rows(c)}
    assert names == {"e.jpg": "部分損壞已存", "d.nef": "救不回來", "全0.wav": "原檔空白未存"}, names
    try:
        cmd_plan(c, keep=keep, disk_size=size)
        raise AssertionError("已輸出後不能重新規劃")
    except SystemExit as e:
        assert "不能重新規劃" in str(e)
    # 沒讀過的位置不能當資料
    rd = ImageReader(os.path.join(DD_DIR, IMG), os.path.join(DD_DIR, MAP))
    try:
        rd.read(size - SECTOR, SECTOR)
        raise AssertionError("沒讀過的位置應該丟 NotYetRead")
    except NotYetRead:
        pass
    rd.close()
    BATCHES, BATCH_NAME, free_bytes, need_to_free = real_batches, {b[0]: b[1] for b in real_batches}, real_free, real_need
    assert vm_ip_parse_ok()
    R.DST_ROOT = real_dst
    print("SELFTEST OK")
    shutil.rmtree(base, ignore_errors=True)


def vm_ip_parse_ok():
    global LEASES
    real, tmp = LEASES, os.path.join(tempfile.mkdtemp(), "leases")
    open(tmp, "w").write("lease 192.168.131.127 {\n hardware ethernet 00:50:56:aa:bb:cc;\n}\n"
                         "lease 192.168.131.128 {\n hardware ethernet 00:50:56:2A:0D:E1;\n}\n")
    LEASES = tmp
    try:
        return vm_ip() == "192.168.131.128"
    finally:
        LEASES = real


# ---------- 主程式 ----------
def main():
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="ddrescue 虛擬機救援（Windows 端）")
    ap.add_argument("mode", nargs="?", choices=["vm-check", "mft", "plan", "copy", "retry", "extract", "report", "batches", "close", "bad"])
    ap.add_argument("--final", action="store_true", help="extract 時決定部分損壞/損壞（不再補讀之後才用）")
    ap.add_argument("--batch", help="copy 指定批次（例如 b13）；不指定就照順序讀下一批")
    ap.add_argument("--skip", help="copy 這一輪跳過的硬碟位置，單位 GB，例如 --skip 2450-2460（留到最後補讀）")
    ap.add_argument("--folder", help=r'copy 只讀這個資料夾底下還沒救到的檔案，例如 --folder "E:\家人B\20250124\Facebook"')
    ap.add_argument("--zones", choices=["safe", "danger"], help="retry 分段：safe 先補讀危險區以外的壞區，danger 最後只讀危險區；不指定＝全部一起")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.zones and a.mode != "retry":
        raise SystemExit("--zones 只能搭配 retry")
    if a.selftest:
        return selftest()
    if not a.mode:
        return ap.print_help()
    c = R.db_open()
    ensure_columns(c)
    if a.mode == "report":
        return print(report_text(c, R.meta_get(c, "cur_domain", "")))
    if a.mode == "bad":                    # 不碰硬碟：列出這一批（或 --batch 指定）目前碰到壞區的檔案
        m = re.match(r"domain_(b\d+)\.map$", R.meta_get(c, "cur_domain", ""))
        batch = a.batch or (m and m.group(1))
        rows = bad_files(c, batch) if batch else []
        print(f"第 {batch} 批目前有 {len(rows)} 個檔案碰到壞區（缺的比例是暫時的，全部批次讀完後會再補讀）")
        return [print(f"  {s / 1e6:.1f} MB，缺 {b * 100 / s:.0f}%  {p}") for p, s, b in rows]
    if a.mode == "batches":
        return print(batches_text(c))
    if a.mode == "plan":
        return cmd_plan(c)
    if a.mode == "extract":
        return cmd_extract(c, final=a.final)
    if a.mode == "close":                  # 不碰硬碟：把讀到一半的批次收尾，沒讀到的留給 retry（9/17 b01 最後 307 MB 密集壞軌）
        if not a.batch or a.batch not in BATCH_NAME:
            raise SystemExit("用法：close --batch bNN")
        S["stage"] = f"收尾第 {a.batch} 批：{BATCH_NAME[a.batch]}"
        return close_batch(c, a.batch)
    if host_sees_drive():
        raise SystemExit("Windows 看得到 Seagate 硬碟（它應該只在虛擬機裡）。請拔掉硬碟 USB，先開虛擬機再接上。")
    R.write_lock()
    try:
        ctypes.windll.kernel32.SetThreadExecutionState(0x80000001)
    except Exception:
        pass
    try:
        runner = VmRunner()
        if a.mode == "vm-check":
            print(f"虛擬機 {runner.ip}，共用資料夾已掛載，硬碟 {runner.dev}（{runner.size / 1e9:.0f} GB，唯讀），Windows 沒有抓到硬碟。{runner.status_line()}")
        elif a.mode == "mft":
            cmd_mft(runner, c)
        else:
            skip = tuple(int(float(x) * 1e9) for x in a.skip.split("-")) if a.skip else None
            cmd_copy(runner, c, retry=(a.mode == "retry"), batch=a.batch, skip=skip, folder=a.folder, zones=a.zones)
    finally:
        try:
            os.remove(R.LOCK)
        except OSError:
            pass


if __name__ == "__main__":
    main()
