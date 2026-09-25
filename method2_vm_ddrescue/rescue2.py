# -*- coding: utf-8 -*-
r"""第二階段：直接讀硬碟磁區的照片救援（現在由 rescue_dd.py 當 NTFS 解析函式庫使用；單獨執行的模式沒有在真實硬碟上跑過，來龍去脈見 docs/演進過程.md）

  python rescue2.py step0        階段 0：只讀 MFT 做盤點 + 抽樣驗證（需要系統管理員；由排程工作執行）
  python rescue2.py run          階段 A/A2/B + 輸出（需要系統管理員；由排程工作執行；可隨時中斷、重跑接續）
  python rescue2.py report       只印出目前的統計（不需要系統管理員，不碰硬碟）
  python rescue2.py --selftest   用自製的迷你 NTFS 映像做完整測試，不碰任何真的硬碟

原則：對來源硬碟只讀不寫；硬碟設為離線後 Windows 也碰不到它；G 槽已有的完整檔案絕不覆蓋。
"""
import argparse, ctypes, ctypes.wintypes as W, datetime, json, os, random, shutil, sqlite3, struct, subprocess, sys, threading, time

# ===== 可調參數 =====
DISK_MODEL = "Seagate Expansion"                   # 用型號找實體磁碟（磁碟編號可能會變）
WANT_TOPS = ["照片", "家人A", "家人B", "家人C", "手機備份", "電子書", "字型", "music", "software", "game"]
DST_ROOT = r"G:\我的雲端硬碟\!家裡硬碟"              # E:\照片\x → G:\...\!家裡硬碟\照片\x
WORK = r"C:\照片救援"
DB = os.path.join(WORK, "phase2.sqlite")
TMP = os.path.join(WORK, "p2_tmp")
LOG = os.path.join(WORK, "進度紀錄.log")
LOCK = os.path.join(WORK, "rescue.lock")
REPORT_TXT = r"G:\我的雲端硬碟\!家裡硬碟\進度回報.txt"
XLSX = r"G:\我的雲端硬碟\!家裡硬碟\照片搬移異常與壞軌清單.xlsx"
ORPHAN_DIR = "照片救援_無法判斷資料夾"               # 所屬資料夾的 MFT 記錄讀不到的檔案放這裡
CHUNK = 1 << 20                                    # 階段 A 一次讀 1 MiB；讀失敗就整塊標壞、跳過
TRIM_STEP = 64 << 10                               # 階段 B 從壞塊兩端往內縮的步長
SKIP_MIN, SKIP_MAX = 1 << 20, 64 << 20             # 連續壞區的倍增跳躍
SLOW_SEC = 10                                      # 一次讀取超過這麼久算「慢」（硬碟在重試）
HUNG_STREAK = 3                                    # 連續幾次慢/失敗 → 讀探測區確認硬碟還活著
PROBE_SEC = 30
READ_MIN, REST_MIN = 120, 60                       # 讀 2 小時、休息 1 小時，自動循環
POLL_SEC = 30                                      # 斷線/卡死時多久檢查一次
GONE_STREAK = 6                                    # 連續幾次「裝置不見」（× POLL_SEC）就當成卡死處理
STUCK_SEC = 180                                    # 一次讀取卡了超過這麼久才「斷線」（使用者拔掉、USB 自己重置）：這塊直接標壞，不重讀
STALL_KILL_SEC = 600                               # 一次讀取卡超過這麼久：把這個位置記進黑名單（重啟後直接當壞區）、強制結束程式、用排程工作自動重啟接續
# v2.1：不做 pnputil 軟體重置 USB。9/16 06:22 第一階段在讀取卡住時重啟裝置造成藍屏 0x50，卡死一律等使用者拔電源。
MFT_BAD_STREAK = 8                                 # 讀 MFT 時同一個 1 MiB 裡連續幾個磁區壞掉，就放棄這 1 MiB 剩下的部分
REPORT_EVERY_SEC = 300
DST_RETRY_SLEEP = 60                               # 寫入 G 槽失敗多久後重試
KEEP_PARTIAL_EXT = {".jpg", ".jpeg"}               # 只有 JPG 的部分損壞檔會存（壞區 ≤ KEEP_PARTIAL_PCT%）
KEEP_PARTIAL_PCT = 5.0
VERIFY_N = 200
SKIP_NAMES = {"thumbs.db", "desktop.ini", ".ds_store", "@eadir", ".@__thumb"}
# ====================

S = dict(phase="第二階段", state="啟動中", cur="", stage="", done=0, partial=0, damaged=0, bad_chunks=0,
         read_bytes=0, reads=0, fails=0, chunks_total=0, chunks_left=0, prog_prev=0, resting_until=0.0, hung_since=0.0,
         active_sec=0.0, last_good="", note="",
         read_start=0.0, read_off=-1,       # 目前正在進行的那一次讀取（給卡住看門狗用）；0 = 沒有讀取在進行
         relaunched=0)                      # 這次執行是第幾次「卡住 → 強制結束 → 自動重啟」之後的接續（0 = 使用者自己啟動的）
STUCK = set()                               # 之前卡超過 STALL_KILL_SEC 被強制結束的讀取位置（存在 meta.stuck_offs）：直接當壞區
watchdog = {"fired": None}
PREV = {"t": time.time(), "bytes": 0}
LOGLOCK = threading.Lock()


def now():
    return datetime.datetime.now().strftime("%Y/%m/%d %H:%M:%S")


def log(text):
    with LOGLOCK:
        try:
            with open(LOG, "a", encoding="utf-8") as f:
                f.write(text + "\n")
        except OSError:
            pass


def say(msg):
    line = f"[{now()[11:]}] {msg}"
    print(line, flush=True)
    log(line)


# ---------- 讀取層 ----------
class ReadError(Exception):
    def __init__(self, code, off=0, n=0):
        super().__init__(f"讀取失敗 code={code} off={off} n={n}")
        self.code, self.off, self.n = code, off, n


class DiskGone(ReadError):
    """裝置不見了（拔掉、卡死到 Windows 移除它）。"""


GONE_CODES = {2, 3, 6, 21, 55, 1167, 1117 + 100000}   # 檔案/路徑不存在、無效控制代碼、裝置未就緒、網路名稱、裝置未連接


class RawDisk:
    """以唯讀 + 不經快取開啟實體磁碟。read() 的 off/n 必須是磁區大小的倍數。"""
    def __init__(self, path, sector):
        k = ctypes.WinDLL("kernel32", use_last_error=True)
        k.CreateFileW.restype = ctypes.c_void_p
        k.CreateFileW.argtypes = [W.LPCWSTR, W.DWORD, W.DWORD, ctypes.c_void_p, W.DWORD, W.DWORD, ctypes.c_void_p]
        k.ReadFile.argtypes = [ctypes.c_void_p, ctypes.c_void_p, W.DWORD, ctypes.POINTER(W.DWORD), ctypes.c_void_p]
        k.SetFilePointerEx.argtypes = [ctypes.c_void_p, ctypes.c_longlong, ctypes.POINTER(ctypes.c_longlong), W.DWORD]
        k.CloseHandle.argtypes = [ctypes.c_void_p]
        k.VirtualAlloc.restype = ctypes.c_void_p
        k.VirtualAlloc.argtypes = [ctypes.c_void_p, ctypes.c_size_t, W.DWORD, W.DWORD]
        k.VirtualFree.argtypes = [ctypes.c_void_p, ctypes.c_size_t, W.DWORD]
        self.k, self.sector, self.path = k, sector, path
        self.h = k.CreateFileW(path, 0x80000000, 3, None, 3, 0x20000000, None)   # GENERIC_READ, share R/W, NO_BUFFERING
        if self.h == ctypes.c_void_p(-1).value:
            raise DiskGone(ctypes.get_last_error(), 0, 0)
        self.cap = CHUNK + sector
        self.buf = k.VirtualAlloc(None, self.cap, 0x3000, 0x04)

    def read(self, off, n):
        assert off % self.sector == 0 and n % self.sector == 0 and n <= self.cap
        k = self.k
        if not k.SetFilePointerEx(self.h, off, None, 0):
            raise ReadError(ctypes.get_last_error(), off, n)
        got = W.DWORD(0)
        if not k.ReadFile(self.h, self.buf, n, ctypes.byref(got), None):
            code = ctypes.get_last_error()
            raise (DiskGone if code in GONE_CODES else ReadError)(code, off, n)
        if got.value != n:
            raise ReadError(-1, off, n)
        return ctypes.string_at(self.buf, n)

    def close(self):
        try:
            self.k.CloseHandle(self.h)
            self.k.VirtualFree(self.buf, 0, 0x8000)
        except Exception:
            pass


class FileDisk:
    """測試用：磁碟內容在記憶體裡，bad 是會讀失敗的 4K 區塊起始位移集合；gone 可以模擬拔線。"""
    def __init__(self, data, sector=4096, bad=(), slow=()):
        self.data, self.sector, self.bad, self.slow, self.gone, self.reads = bytearray(data), sector, set(bad), set(slow), False, []

    def read(self, off, n):
        assert off % self.sector == 0 and n % self.sector == 0
        self.reads.append((off, n))
        if self.gone:
            raise DiskGone(1167, off, n)
        for b in range(off, off + n, self.sector):
            if b in self.bad:
                raise ReadError(23, off, n)
        return bytes(self.data[off:off + n])

    def close(self):
        pass


def find_disk():
    """用型號找實體磁碟：回傳 (路徑, 磁區大小, 總大小)。"""
    ps = ("Get-CimInstance Win32_DiskDrive | Where-Object Model -like '*" + DISK_MODEL + "*' | "
          "Select-Object -First 1 DeviceID, BytesPerSector, Size | ConvertTo-Json -Compress")
    out = ps_run(ps, 300)
    if not out:
        return None
    d = json.loads(out)
    return d["DeviceID"], int(d["BytesPerSector"]), int(d["Size"])


def ps_run(ps, timeout):
    """跑 PowerShell，回傳 stdout（逾時或失敗回傳 ''；硬碟卡死時 Get-Disk/CIM 查詢會卡住，不能讓它把程式弄當）。"""
    try:
        return subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True,
                              errors="ignore", timeout=timeout).stdout.strip()
    except (subprocess.TimeoutExpired, OSError):
        return ""


def set_disk_offline(offline=True):
    ps_run(f"Get-Disk | Where-Object FriendlyName -like '*{DISK_MODEL}*' | Set-Disk -IsOffline ${'true' if offline else 'false'}", 300)


def is_admin():
    try:
        return bool(ctypes.windll.shell32.IsUserAnAdmin())
    except Exception:
        return False


class Reader:
    """在 RawDisk 之上加：斷線等待、卡死偵測（探測區）、讀 N 分鐘休息 M 分鐘、統計。"""
    def __init__(self, open_fn, probe_off, sector):
        self.open_fn, self.dev, self.probe_off, self.sector = open_fn, None, probe_off, sector
        self.streak, self.gone_streak, self.active, self.session_start = 0, 0, 0.0, time.time()
        self.next_rest = time.time() + READ_MIN * 60

    def _ensure(self):
        while self.dev is None:
            try:
                self.dev = self.open_fn()
            except (DiskGone, ReadError, OSError) as e:
                S["state"] = f"硬碟不在（{e}），每 {POLL_SEC} 秒檢查一次"
                time.sleep(POLL_SEC)
        return self.dev

    def _close(self):
        if self.dev:
            self.dev.close()
        self.dev = None

    def _probe(self):
        t0 = time.time()
        try:
            self._ensure().read(self.probe_off, self.sector)
            return time.time() - t0 < PROBE_SEC
        except Exception:
            return False

    def _reopen_and_probe(self):
        """重新開啟裝置並讀探測區。卡死的硬碟不可能在 PROBE_SEC 內讀到探測區；讀得到就是恢復了。"""
        try:
            self.dev = self.open_fn()
        except Exception:
            return False
        if self._probe():
            return True
        self._close()
        return False

    def wait_power_cycle(self):
        """硬碟卡死：停止讀取，等使用者拔掉電源重插（不做任何軟體重置，避免在讀取卡住時拆裝置造成藍屏）。"""
        S["hung_since"] = S["hung_since"] or time.time()
        self._close()
        S["state"] = ("硬碟卡死（連確定完好的區域都讀不到）。已停止讀取，請拔掉硬碟電源與 USB → 等 1 分鐘 → 先插電源再插 USB；"
                      f"程式每 {POLL_SEC} 秒自動檢查，恢復後自己繼續")
        say(S["state"])
        while True:
            time.sleep(POLL_SEC)
            if self._reopen_and_probe():
                break
        self._recovered()

    def _recovered(self):
        S["hung_since"] = 0.0
        self.streak = self.gone_streak = 0
        S["state"] = "執行中"
        say("硬碟已恢復，繼續")

    def maybe_rest(self):
        if time.time() < self.next_rest:
            return
        S["resting_until"] = time.time() + REST_MIN * 60
        S["state"] = f"讀滿 {READ_MIN} 分鐘，休息 {REST_MIN} 分鐘（自動繼續）"
        say(S["state"])
        self._close()
        time.sleep(REST_MIN * 60)
        S["resting_until"] = 0.0
        self.next_rest = time.time() + READ_MIN * 60
        S["state"] = "執行中"

    def read(self, off, n):
        """回傳 bytes；壞區丟 ReadError。斷線與卡死在這裡處理完才回來。"""
        self.maybe_rest()
        while True:
            if off in STUCK:                                    # 上次在這裡卡了 10 分鐘以上被強制結束：直接當壞區，不再讀
                S["reads"] += 1
                S["fails"] += 1
                raise ReadError(-4, off, n)
            dev = self._ensure()
            t0 = time.time()
            try:
                S.update(read_start=t0, read_off=off)
                try:
                    data = dev.read(off, n)
                finally:
                    S["read_start"] = 0.0                       # 只有真的在核心裡讀的那段時間算「卡住」
                dt = time.time() - t0
                self.active += dt
                S["reads"] += 1
                S["read_bytes"] += n
                self.gone_streak = 0
                self.streak = self.streak + 1 if dt > SLOW_SEC else 0
                if self.streak >= HUNG_STREAK and not self._probe():
                    self.wait_power_cycle()
                return data
            except DiskGone as e:
                dt = time.time() - t0
                self._close()
                if dt > STUCK_SEC:                              # 卡了很久才斷線（拔掉/USB 重置）：這塊就是壞的，不重讀，免得再卡一次
                    self.active += dt
                    S["reads"] += 1
                    S["fails"] += 1
                    say(f"！讀取卡住 {int(dt)} 秒後硬碟斷線（code {e.code}）：這塊標為壞區，不再重讀")
                    raise ReadError(e.code, off, n) from None
                self.gone_streak += 1
                if self.gone_streak >= GONE_STREAK:            # 一直「不見」= 卡死到 Windows 把它移掉了，走重置流程
                    self.wait_power_cycle()
                    continue
                say(f"！硬碟斷線（code {e.code}），等待重新連接")
                S["state"] = "硬碟斷線，等待重新連接"
                time.sleep(POLL_SEC)
                continue
            except ReadError as e:
                dt = time.time() - t0
                self.active += dt
                S["reads"] += 1
                S["fails"] += 1
                self.streak += 1
                if self.streak >= HUNG_STREAK and not self._probe():
                    self.wait_power_cycle()
                    continue                                    # 剛才的失敗可能是卡死造成的：恢復後同一塊再讀一次
                raise


# ---------- NTFS 解析 ----------
FILETIME_EPOCH = 116444736000000000


def ft2epoch(ft):
    return (ft - FILETIME_EPOCH) / 1e7 if ft else 0.0


def fixup(rec, stride=512):
    """套用 update sequence（每 512 bytes 的最後 2 bytes 被換成 USN，原值存在 USA）。回傳 bytearray 或 None（損毀）。"""
    if rec[:4] not in (b"FILE", b"INDX"):
        return None
    usa_off, usa_cnt = struct.unpack_from("<HH", rec, 4)
    if usa_cnt < 2 or usa_off + usa_cnt * 2 > len(rec) or (usa_cnt - 1) * stride > len(rec):
        return None
    out = bytearray(rec)
    usn = out[usa_off:usa_off + 2]
    for i in range(1, usa_cnt):
        pos = i * stride - 2
        if out[pos:pos + 2] != usn:
            return None
        out[pos:pos + 2] = out[usa_off + i * 2:usa_off + i * 2 + 2]
    return out


def decode_runs(buf, off):
    """資料段：[(lcn 或 None=稀疏, 長度_clusters), ...]"""
    runs, lcn = [], 0
    while off < len(buf):
        h = buf[off]
        if h == 0:
            break
        ls, os_ = h & 0xF, h >> 4
        off += 1
        length = int.from_bytes(buf[off:off + ls], "little")
        off += ls
        if os_ == 0:
            runs.append((None, length))
        else:
            delta = int.from_bytes(buf[off:off + os_], "little", signed=True)
            off += os_
            lcn += delta
            runs.append((lcn, length))
    return runs


def parse_record(raw, recno):
    """解析一筆 MFT 記錄。回傳 dict 或 None（損毀/非 FILE）。損毀記錄裡的任何欄位都可能是垃圾，解析炸掉一律當損毀。"""
    try:
        return _parse_record(raw, recno)
    except (struct.error, IndexError, ValueError, UnicodeDecodeError):
        return None


def _parse_record(raw, recno):
    rec = fixup(raw)
    if rec is None:
        return None
    seq, links, first_attr, flags = struct.unpack_from("<HHHH", rec, 0x10)
    base = struct.unpack_from("<Q", rec, 0x20)[0] & 0xFFFFFFFFFFFF
    r = dict(no=recno, seq=seq, inuse=bool(flags & 1), isdir=bool(flags & 2), base=base,
             names=[], data=None, named_data=False, attrlist=None, mtime=0.0, atime=0.0, unsupported="")
    off = first_attr
    while off + 8 <= len(rec):
        atype, alen = struct.unpack_from("<II", rec, off)
        if atype == 0xFFFFFFFF or alen == 0 or off + alen > len(rec):
            break
        nonres, nlen, noff, aflags = rec[off + 8], rec[off + 9], struct.unpack_from("<H", rec, off + 10)[0], struct.unpack_from("<H", rec, off + 12)[0]
        name = rec[off + noff:off + noff + nlen * 2].decode("utf-16-le", "replace") if nlen else ""
        if nonres:
            svcn, evcn, runoff = struct.unpack_from("<QQH", rec, off + 0x10)
            alloc, real, init = struct.unpack_from("<QQQ", rec, off + 0x28)
            body = dict(res=False, svcn=svcn, evcn=evcn, runs=decode_runs(rec, off + runoff), size=real, init=init, flags=aflags)
        else:
            vlen, voff = struct.unpack_from("<IH", rec, off + 0x10)
            body = dict(res=True, value=bytes(rec[off + voff:off + voff + vlen]), size=vlen, flags=aflags)
        if atype == 0x10 and body["res"] and len(body["value"]) >= 0x20:
            r["mtime"], r["atime"] = ft2epoch(struct.unpack_from("<Q", body["value"], 8)[0]), ft2epoch(struct.unpack_from("<Q", body["value"], 0x18)[0])
        elif atype == 0x30 and body["res"]:
            v = body["value"]
            parent = struct.unpack_from("<Q", v, 0)[0] & 0xFFFFFFFFFFFF
            fnlen, ns = v[0x40], v[0x41]
            r["names"].append((parent, v[0x42:0x42 + fnlen * 2].decode("utf-16-le", "replace"), ns))
        elif atype == 0x80:
            if name:
                r["named_data"] = True
            elif r["data"] is None:
                r["data"] = body
                if aflags & 0x4001:               # 壓縮 / 加密：不支援
                    r["unsupported"] = "壓縮" if aflags & 1 else "加密"
        elif atype == 0x20:
            r["attrlist"] = body
        off += alen
    return r


def parse_attrlist(value):
    """$ATTRIBUTE_LIST 的項目：[(type, svcn, mftref, namelen)]"""
    out, off = [], 0
    while off + 0x1A <= len(value):
        atype, elen = struct.unpack_from("<IH", value, off)
        if elen == 0:
            break
        nlen = value[off + 6]
        svcn, ref = struct.unpack_from("<QQ", value, off + 8)
        out.append((atype, svcn, ref & 0xFFFFFFFFFFFF, nlen))
        off += elen
    return out


class Volume:
    def __init__(self, reader, part_off, sector):
        self.rd, self.part_off, self.sector = reader, part_off, sector
        boot = reader.read(part_off, sector)
        if boot[3:11] != b"NTFS    ":
            raise RuntimeError(f"分割區起點 {part_off} 不是 NTFS（{boot[3:11]!r}）")
        self.bps = struct.unpack_from("<H", boot, 0x0B)[0]
        spc = boot[0x0D]
        self.spc = spc if spc < 0x80 else 1 << (256 - spc)
        self.cluster = self.bps * self.spc
        self.total_sectors, self.mft_lcn, self.mftmirr_lcn = struct.unpack_from("<QQQ", boot, 0x28)
        cpr = struct.unpack_from("<b", boot, 0x40)[0]
        self.rec_size = cpr * self.cluster if cpr > 0 else 1 << (-cpr)
        self.mft_runs = None
        self.records = {}          # recno -> parsed
        self.bad_recs = set()      # 讀不到 / 損毀的記錄編號
        self.mft_bytes = 0

    def lcn_off(self, lcn):
        return self.part_off + lcn * self.cluster

    def runs_to_extents(self, runs):
        """[(lcn, len)] → [(device_off 或 None, nbytes)]"""
        return [(None if lcn is None else self.lcn_off(lcn), ln * self.cluster) for lcn, ln in runs]

    def read_extents(self, extents, size, on_chunk=None):
        """依序讀 extents（測試/MFT 用；正式流程走 chunk 表）。壞區 → 丟 ReadError。"""
        out, left = bytearray(), size
        for off, n in extents:
            n = min(n, left)
            if n <= 0:
                break
            if off is None:
                out += bytes(n)
            else:
                pos = 0
                while pos < n:
                    m = min(CHUNK, n - pos)
                    mm = -(-m // self.sector) * self.sector
                    out += self.rd.read(off + pos, mm)[:m]
                    pos += m
            left -= n
        return bytes(out[:size])

    def find_mft_runs(self):
        """$MFT 本身的資料段 [(lcn, len)]（MFT 碎片多時，擴充記錄在 $ATTRIBUTE_LIST 裡）。ddrescue 版先用它決定要救哪些位置。"""
        rec0 = self.rd.read(self.lcn_off(self.mft_lcn), max(self.rec_size, self.sector))[:self.rec_size]
        r0 = parse_record(rec0, 0)
        if not r0 or not r0["data"] or r0["data"]["res"]:
            raise RuntimeError("MFT 第 0 筆記錄讀不到或損毀，無法解析這個分割區")
        runs = list(r0["data"]["runs"])
        if r0["attrlist"]:
            for atype, svcn, ref, nlen in parse_attrlist(self._attr_value(r0["attrlist"], runs)):
                if atype == 0x80 and ref != 0 and nlen == 0:
                    ext = parse_record(self._read_rec_from_runs(ref, runs), ref)
                    if ext and ext["data"] and not ext["data"]["res"]:
                        runs += ext["data"]["runs"]
        return runs

    def load_mft(self, progress=None):
        """讀整個 MFT。讀不到的區塊裡的記錄記入 bad_recs。"""
        runs = self.find_mft_runs()
        self.mft_runs = runs
        total = sum(ln for _, ln in runs) * self.cluster
        self.mft_bytes = total
        recno, done = 0, 0
        for lcn, ln in runs:
            nbytes = ln * self.cluster
            pos = 0
            while pos < nbytes:
                m = min(CHUNK, nbytes - pos)
                m = -(-m // self.sector) * self.sector
                off = self.lcn_off(lcn) + pos
                try:
                    blob = self.rd.read(off, m)
                    self._ingest(blob, recno)
                except ReadError:
                    streak = 0
                    for sub in range(0, m, self.sector):    # 這 1 MiB 有壞區：改成一個磁區一個磁區讀
                        try:
                            if streak >= MFT_BAD_STREAK:    # 連續壞太多：剩下的直接放棄（每個壞磁區都要付一次 1～7 分鐘的逾時）
                                raise ReadError(-3, off + sub, self.sector)
                            blob = self.rd.read(off + sub, self.sector)
                            self._ingest(blob, recno + sub // self.rec_size)
                            streak = 0
                        except ReadError:
                            streak += 1
                            for k in range(max(1, self.sector // self.rec_size)):
                                self.bad_recs.add(recno + (sub // self.rec_size) + k)
                recno += m // self.rec_size
                pos += m
                done += m
                if progress:
                    progress(done, total)
        self._link_extensions()

    def _ingest(self, blob, first_recno):
        for i in range(0, len(blob) - self.rec_size + 1, self.rec_size):
            raw = blob[i:i + self.rec_size]
            no = first_recno + i // self.rec_size
            if raw[:4] == b"\x00\x00\x00\x00":
                continue                                    # 沒用過的記錄
            r = parse_record(raw, no)
            if r is None:
                self.bad_recs.add(no)
            else:
                self.records[no] = r

    def _read_rec_from_runs(self, recno, runs):
        byte_off, vcn_off = recno * self.rec_size, 0
        for lcn, ln in runs:
            span = ln * self.cluster
            if byte_off < vcn_off + span:
                rel = byte_off - vcn_off
                off = self.lcn_off(lcn) + rel
                aligned = off - off % self.sector
                blob = self.rd.read(aligned, -(-(off - aligned + self.rec_size) // self.sector) * self.sector)
                return blob[off - aligned:off - aligned + self.rec_size]
            vcn_off += span
        raise ReadError(-2, byte_off, self.rec_size)

    def _attr_value(self, body, runs=None):
        if body["res"]:
            return body["value"]
        return self.read_extents(self.runs_to_extents(body["runs"]), body["size"])

    def _link_extensions(self):
        """擴充記錄（base != 0）的 $DATA 併回主記錄；主記錄有 $ATTRIBUTE_LIST 時依 VCN 排序接起來。"""
        ext = {}
        for no, r in self.records.items():
            if r["base"] and r["base"] != no:
                ext.setdefault(r["base"], []).append(r)
        for no, r in list(self.records.items()):
            if r["base"] and r["base"] != no:
                continue
            if not r["attrlist"]:
                continue
            try:
                entries = parse_attrlist(self._attr_value(r["attrlist"]))
            except ReadError:
                r["unsupported"] = "屬性清單讀不到"
                continue
            pieces = []
            for atype, svcn, ref, nlen in entries:
                if atype != 0x80 or nlen:
                    continue
                src = r if ref == no else self.records.get(ref)
                if src is None:
                    r["unsupported"] = f"擴充記錄 {ref} 讀不到"
                    break
                if src["data"] and not src["data"]["res"] and src["data"]["svcn"] == svcn:
                    pieces.append((svcn, src["data"]))
            if r["unsupported"]:
                continue
            if pieces:
                pieces.sort(key=lambda p: p[0])
                main = r["data"] if r["data"] and not r["data"]["res"] else pieces[0][1]
                runs = []
                for _, body in pieces:
                    runs += body["runs"]
                r["data"] = dict(main, runs=runs, size=max(b["size"] for _, b in pieces), init=max(b["init"] for _, b in pieces))

    def path_of(self, no, _seen=None):
        """回傳 (路徑片段 list 由根開始, 是否孤兒)。根目錄 = []"""
        _seen = _seen or set()
        if no == 5:
            return [], False
        r = self.records.get(no)
        if r is None or not r["names"] or no in _seen:
            return [f"{ORPHAN_DIR}", f"parent#{no}"], True
        _seen.add(no)
        parent, name, ns = self.best_name(r)
        head, orphan = self.path_of(parent, _seen)
        return head + [name], orphan

    @staticmethod
    def best_name(r):
        names = [n for n in r["names"] if n[2] != 2] or r["names"]    # 不要 DOS 8.3 短檔名
        return names[0]


# ---------- 計畫與資料庫 ----------
def db_open():
    os.makedirs(WORK, exist_ok=True)
    c = sqlite3.connect(DB)
    c.executescript("""
    PRAGMA journal_mode=WAL; PRAGMA synchronous=NORMAL;
    CREATE TABLE IF NOT EXISTS files(id INTEGER PRIMARY KEY, mft INT, path TEXT, size INT, mtime REAL, atime REAL, dst TEXT,
        status TEXT, note TEXT DEFAULT '', bad_bytes INT DEFAULT 0, orphan INT DEFAULT 0);
    CREATE TABLE IF NOT EXISTS chunks(id INTEGER PRIMARY KEY, file_id INT, foff INT, dev_off INT, nbytes INT, state TEXT);
    CREATE INDEX IF NOT EXISTS ix_chunks ON chunks(state, dev_off);
    CREATE INDEX IF NOT EXISTS ix_chunks_file ON chunks(file_id);
    CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT);
    """)
    return c


def meta_get(c, key, default=None):
    row = c.execute("SELECT value FROM meta WHERE key=?", (key,)).fetchone()
    return row[0] if row else default


def meta_set(c, key, value):
    c.execute("INSERT OR REPLACE INTO meta VALUES(?,?)", (key, str(value)))
    c.commit()


def is_done(dst, size, mtime):
    try:
        st = os.stat("\\\\?\\" + dst if not dst.startswith("\\\\?\\") else dst)
    except OSError:
        return False
    return st.st_size == size and abs(st.st_mtime - mtime) <= 2


def lp(p):
    return p if p.startswith("\\\\?\\") else "\\\\?\\" + p


def build_plan(vol, c, want_tops=None, dst_root=None, keep=None):
    """從 MFT 選出要救的檔案，寫進 files/chunks。回傳統計 dict。keep(path) 回傳 False 的檔案不列入（也不去 G 槽比對）。"""
    want_tops = want_tops if want_tops is not None else WANT_TOPS
    dst_root = dst_root or DST_ROOT
    st = dict(total=0, done_before=0, to_read=0, to_read_bytes=0, resident=0, unsupported=0, orphan=0, skipped_name=0, bad_recs=len(vol.bad_recs))
    c.execute("DELETE FROM files")
    c.execute("DELETE FROM chunks")
    fid = 0
    disk_bytes = vol.total_sectors * vol.bps
    for no, r in vol.records.items():
        if not r["inuse"] or r["isdir"] or (r["base"] and r["base"] != no) or not r["names"] or no < 16:
            continue
        parts, orphan = vol.path_of(no)
        if not parts:
            continue
        if not orphan and parts[0] not in want_tops:
            continue
        if any(p.lower() in SKIP_NAMES for p in parts):
            st["skipped_name"] += 1
            continue
        path = "E:\\" + "\\".join(parts)
        if keep and not keep(path):
            continue
        st["total"] += 1
        dst = os.path.join(dst_root, *parts)
        data = r["data"]
        if data is None:
            size = 0
            data = dict(res=True, value=b"", size=0, flags=0)
        size = data["size"]
        status, note = "new", ""
        if not r["unsupported"] and not data["res"]:          # 損毀記錄的欄位可能是垃圾：大小超過配置、或位置超出硬碟，都不能照著讀
            alloc = sum(ln for _, ln in data["runs"]) * vol.cluster
            if size > alloc or any(lcn is not None and (lcn < 0 or (lcn + ln) * vol.cluster > disk_bytes) for lcn, ln in data["runs"]):
                r["unsupported"] = "MFT 記錄異常（大小或位置不合理）"
        if r["unsupported"]:
            status, note = "unsupported", r["unsupported"]
            st["unsupported"] += 1
        elif is_done(dst, size, r["mtime"]):
            status = "done_before"
            st["done_before"] += 1
        fid += 1
        c.execute("INSERT INTO files(id,mft,path,size,mtime,atime,dst,status,note,bad_bytes,orphan) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                  (fid, no, path, size, r["mtime"], r["atime"], dst, status, note, 0, int(orphan)))
        if orphan:
            st["orphan"] += 1
        if status != "new":
            continue
        st["to_read"] += 1
        if data["res"]:
            c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(?,?,?,?,?)", (fid, 0, -1, size, "resident"))
            c.execute("UPDATE files SET note=? WHERE id=?", ("hex:" + data["value"].hex(), fid))   # 內容就在 MFT 裡，直接帶著走
            st["resident"] += 1
            continue
        foff, init = 0, data.get("init", size)
        for dev_off, nbytes in vol.runs_to_extents(data["runs"]):
            if foff >= size:
                break
            nbytes = min(nbytes, size - foff)
            if dev_off is None or foff >= init:              # 稀疏 / 尚未初始化：全是 0，不用讀
                c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(?,?,?,?,?)", (fid, foff, -1, nbytes, "zero"))
            else:
                pos = 0
                while pos < nbytes:
                    m = min(CHUNK, nbytes - pos, init - (foff + pos))
                    if m <= 0:
                        break
                    c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(?,?,?,?,?)",
                              (fid, foff + pos, dev_off + pos, m, "new"))
                    st["to_read_bytes"] += m
                    pos += m
                if foff + pos < foff + nbytes:
                    c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(?,?,?,?,?)",
                              (fid, foff + pos, -1, nbytes - pos, "zero"))
            foff += nbytes
    c.commit()
    return st


# ---------- 讀取階段 ----------
def tmp_path(fid):
    return os.path.join(TMP, f"{fid}.bin")


def write_tmp(fid, foff, data, size):
    os.makedirs(TMP, exist_ok=True)
    p = tmp_path(fid)
    with open(p, "r+b" if os.path.exists(p) else "w+b") as f:
        if os.path.getsize(p) < size:
            f.truncate(size)
        f.seek(foff)
        f.write(data)


def pass_read(vol, c, states=("new",), use_skip=True, sector=4096):
    """依磁區位置順序讀指定狀態的 chunk。成功→good（寫暫存），失敗→bad；use_skip 時失敗後跳過後面一段（標 skipped）。"""
    rows = c.execute(f"SELECT id,file_id,foff,dev_off,nbytes FROM chunks WHERE state IN ({','.join('?' * len(states))}) ORDER BY dev_off",
                     states).fetchall()
    S["chunks_total"], S["chunks_left"], S["prog_prev"] = len(rows), len(rows), 0
    skip_until, skip = -1, SKIP_MIN
    sizes = {}
    touched = set()
    for cid, fid, foff, dev_off, nbytes in rows:
        S["chunks_left"] -= 1
        if use_skip and dev_off < skip_until:
            c.execute("UPDATE chunks SET state='skipped' WHERE id=?", (cid,))
            continue
        if fid not in sizes:
            sizes[fid] = c.execute("SELECT size FROM files WHERE id=?", (fid,)).fetchone()[0]
        S["cur"] = c.execute("SELECT path FROM files WHERE id=?", (fid,)).fetchone()[0]
        n_al = -(-nbytes // sector) * sector
        try:
            data = vol.rd.read(dev_off, n_al)[:nbytes]
            write_tmp(fid, foff, data, sizes[fid])
            c.execute("UPDATE chunks SET state='good' WHERE id=?", (cid,))
            skip_until, skip = -1, SKIP_MIN
            S["last_good"] = S["cur"]
        except ReadError as e:
            c.execute("UPDATE chunks SET state='bad' WHERE id=?", (cid,))
            S["bad_chunks"] += 1
            say(f"[壞區] {S['cur']} 位移 {foff} 長 {nbytes}（code {e.code}）")
            if use_skip:
                skip_until = dev_off + nbytes + skip
                skip = min(skip * 2, SKIP_MAX)
        touched.add(fid)
        c.commit()
        finalize_ready(c, touched)
        touched.clear()
    c.commit()


def pass_trim(vol, c, sector=4096):
    """階段 B：對「只差一點就能存」的 JPG，把壞塊兩端往內縮，找出真正壞的範圍。"""
    cand = c.execute("""SELECT f.id, f.size, f.path FROM files f WHERE f.status='new' AND f.size>0
                        AND EXISTS(SELECT 1 FROM chunks WHERE file_id=f.id AND state='bad')
                        AND NOT EXISTS(SELECT 1 FROM chunks WHERE file_id=f.id AND state IN ('new','skipped'))""").fetchall()
    for fid, size, path in cand:
        if os.path.splitext(path)[1].lower() not in KEEP_PARTIAL_EXT:
            continue
        bad = c.execute("SELECT id,foff,dev_off,nbytes FROM chunks WHERE file_id=? AND state='bad'", (fid,)).fetchall()
        if (sum(b[3] for b in bad) - len(bad) * (CHUNK - 2 * TRIM_STEP)) * 100.0 / size > KEEP_PARTIAL_PCT and sum(b[3] for b in bad) > CHUNK:
            continue                                        # 就算修剪也救不到 5% 以內：不值得再讀壞區
        S["cur"] = path
        for cid, foff, dev_off, nbytes in bad:
            lo, hi = 0, nbytes                              # [lo,hi) 是還沒確定的範圍
            step = max(TRIM_STEP, sector)
            while True:                                     # 步長 64K → 16K → 4K（永遠是磁區的倍數，讀取位移才會對齊）
                while lo < hi:                              # 從前面往後
                    m = min(step, hi - lo)
                    m_al = -(-m // sector) * sector
                    try:
                        data = vol.rd.read(dev_off + lo, m_al)[:m]
                        write_tmp(fid, foff + lo, data, size)
                        lo += m
                    except ReadError:
                        break
                while hi > lo:                              # 從後面往前
                    m = min(step, hi - lo)
                    st = hi - m
                    st_al = st - st % sector
                    try:
                        data = vol.rd.read(dev_off + st_al, -(-(hi - st_al) // sector) * sector)[st - st_al:st - st_al + m]
                        write_tmp(fid, foff + st, data, size)
                        hi -= m
                    except ReadError:
                        break
                if step <= sector:
                    break
                step = max(step // 4, sector)
            # [lo,hi) 是壞的（狀態 trimmed：已經修剪過，重新啟動不會再讀一次）；其餘已寫入暫存
            c.execute("DELETE FROM chunks WHERE id=?", (cid,))
            if lo > 0:
                c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(?,?,?,?,'good')", (fid, foff, dev_off, lo))
            if hi > lo:
                c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(?,?,?,?,'trimmed')", (fid, foff + lo, dev_off + lo, hi - lo))
            if hi < nbytes:
                c.execute("INSERT INTO chunks(file_id,foff,dev_off,nbytes,state) VALUES(?,?,?,?,'good')", (fid, foff + hi, dev_off + hi, nbytes - hi))
            c.commit()
        finalize_ready(c, {fid})


def finalize_ready(c, fids, final=False):
    """chunk 都有結果的檔案：全部 good → 搬到 G；有 bad → final 時決定存部分損壞或放棄。"""
    for fid in fids:
        row = c.execute("SELECT path,size,mtime,atime,dst,status,note FROM files WHERE id=?", (fid,)).fetchone()
        if not row or row[5] != "new":
            continue
        path, size, mtime, atime, dst, _, note = row
        pending = c.execute("SELECT COUNT(*) FROM chunks WHERE file_id=? AND state IN ('new','skipped')", (fid,)).fetchone()[0]
        if pending:
            continue
        bad = c.execute("SELECT foff,nbytes FROM chunks WHERE file_id=? AND state IN ('bad','trimmed')", (fid,)).fetchall()
        res = c.execute("SELECT nbytes FROM chunks WHERE file_id=? AND state='resident'", (fid,)).fetchone()
        if res:
            write_tmp(fid, 0, bytes.fromhex(note[4:]) if note.startswith("hex:") else b"", size)
        if not bad:
            deliver(c, fid, path, size, mtime, atime, dst, partial=False)
            continue
        if not final:
            continue
        bad_bytes = sum(n for _, n in bad)
        pct = bad_bytes * 100.0 / size if size else 100.0
        if os.path.splitext(path)[1].lower() in KEEP_PARTIAL_EXT and pct <= KEEP_PARTIAL_PCT:
            stem, ext = os.path.splitext(dst)
            deliver(c, fid, path, size, mtime, atime, stem + "_部分損壞" + ext, partial=True, bad_bytes=bad_bytes)
        else:
            c.execute("UPDATE files SET status='damaged', bad_bytes=?, note=? WHERE id=?",
                      (bad_bytes, f"壞區 {bad_bytes / 1024:.0f} KB（{pct:.1f}%），未存到雲端", fid))
            c.commit()
            S["damaged"] += 1
            try:
                os.remove(tmp_path(fid))
            except OSError:
                pass


def deliver(c, fid, path, size, mtime, atime, dst, partial, bad_bytes=0):
    """暫存檔 → G 槽（先 .part 再改名），設定時間，更新狀態。"""
    src = tmp_path(fid)
    if not os.path.exists(src) and not c.execute("SELECT 1 FROM chunks WHERE file_id=? AND state='good' LIMIT 1", (fid,)).fetchone():
        # 沒有任何要讀的資料（空檔、整個稀疏/未初始化＝全 0）→ 本來就不會有暫存檔。有讀到資料卻沒暫存檔則照樣報錯，不能輸出成全 0
        os.makedirs(TMP, exist_ok=True)
        open(src, "wb").close()
    if os.path.getsize(src) != size:
        with open(src, "r+b") as f:
            f.truncate(size)
    for attempt in range(6):
        try:
            os.makedirs(lp(os.path.dirname(dst)), exist_ok=True)
            shutil.copyfile(src, lp(dst + ".part"))
            os.utime(lp(dst + ".part"), (atime or mtime or time.time(), mtime or time.time()))
            os.replace(lp(dst + ".part"), lp(dst))
            break
        except OSError as e:
            if attempt == 5:
                c.execute("UPDATE files SET status='dst_error', note=? WHERE id=?", (f"寫入 G 槽失敗：{e}", fid))
                c.commit()
                say(f"[寫入失敗] {path}：{e}")
                return
            S["state"] = f"寫入 G 槽失敗（{e}），{DST_RETRY_SLEEP} 秒後重試"
            say(S["state"])
            time.sleep(DST_RETRY_SLEEP)
            S["state"] = "執行中"
    try:
        os.remove(src)
    except OSError:
        pass
    if partial:
        c.execute("UPDATE files SET status='partial_saved', bad_bytes=?, dst=?, note=? WHERE id=?",
                  (bad_bytes, dst, f"壞區 {bad_bytes / 1024:.0f} KB（{bad_bytes * 100.0 / size:.1f}%）已補 0，存為 _部分損壞", fid))
        S["partial"] += 1
    else:
        c.execute("UPDATE files SET status='done' WHERE id=?", (fid,))
        S["done"] += 1
    c.commit()


def finalize_all(c):
    fids = [r[0] for r in c.execute("SELECT id FROM files WHERE status='new'").fetchall()]
    finalize_ready(c, fids, final=True)


# ---------- 開啟真實硬碟 ----------
def open_real():
    """回傳 (Reader, Volume, sector)。需要系統管理員。"""
    found = find_disk()
    if not found:
        raise RuntimeError(f"找不到型號含「{DISK_MODEL}」的實體磁碟（硬碟沒接上？）")
    dev_path, sector, _ = found

    def opener():
        f2 = find_disk()
        if not f2:
            raise DiskGone(1167, 0, 0)
        return RawDisk(f2[0], sector)
    rd = Reader(opener, probe_off=0, sector=sector)
    vol = Volume(rd, find_partition(rd, sector), sector)
    rd.probe_off = vol.lcn_off(vol.mft_lcn)                  # 探測區：MFT 開頭（step0 讀過、確定讀得到）
    return rd, vol, sector


def find_partition(rd, sector):
    """讀 MBR（或 GPT）回傳第一個 NTFS 分割區的位元組位移。"""
    mbr = rd.read(0, sector)
    if mbr[510:512] != b"\x55\xAA":
        raise RuntimeError("讀不到 MBR 簽章")
    for i in range(4):
        e = mbr[0x1BE + i * 16:0x1BE + i * 16 + 16]
        ptype, start = e[4], struct.unpack_from("<I", e, 8)[0]
        if ptype in (0x07, 0x17) and start:
            return start * sector
        if ptype == 0xEE:                                   # GPT：讀第一個分割區
            return_lba = struct.unpack_from("<Q", rd.read(sector, sector), 0x48)[0]
            entries = rd.read(return_lba * sector, sector)
            return struct.unpack_from("<Q", entries, 0x20)[0] * sector
    raise RuntimeError("MBR 裡沒有 NTFS 分割區")


# ---------- 回報 ----------
def report(c=None):
    t = time.time()
    dt = max(t - PREV["t"], 1)
    db_ = ""
    if c is not None:
        rows = dict(c.execute("SELECT status, COUNT(*) FROM files GROUP BY status").fetchall())
        ch = dict(c.execute("SELECT state, COUNT(*) FROM chunks GROUP BY state").fetchall())
        db_ = (f"檔案：待處理 {rows.get('new', 0)}，已完整救回 {rows.get('done', 0)}，部分損壞已存 {rows.get('partial_saved', 0)}，"
               f"損壞未存 {rows.get('damaged', 0)}，不支援 {rows.get('unsupported', 0)}，之前已在 G 槽 {rows.get('done_before', 0)}\n"
               f"區塊：未讀 {ch.get('new', 0)}，跳過待重試 {ch.get('skipped', 0)}，已讀 {ch.get('good', 0)}，壞區 {ch.get('bad', 0) + ch.get('trimmed', 0)}")
    st = S["state"]
    if S["hung_since"]:
        st += f"（已等 {int((t - S['hung_since']) // 60)} 分鐘）"
    if S["read_start"] and t - S["read_start"] > SLOW_SEC * 6:
        st += f"（目前這一次讀取已經卡 {int((t - S['read_start']) // 60)} 分 {int((t - S['read_start']) % 60)} 秒，硬碟正在重試；超過 {STALL_KILL_SEC // 60} 分鐘會自動強制跳過）"
    if S["relaunched"]:
        st += f"（卡住強制跳過後自動重啟的第 {S['relaunched']} 次接續，黑名單 {len(STUCK)} 處）"
    prog = ""
    if S["chunks_total"]:
        done_n = S["chunks_total"] - S["chunks_left"]
        rate = (done_n - S["prog_prev"]) / dt                     # 這一輪每秒處理幾個區塊（含跳過的）
        eta = f"，照這 {dt / 60:.0f} 分鐘的速度預估還要 {S['chunks_left'] / rate / 60:.0f} 分鐘" if rate > 0 else ""
        prog = f"這一階段的區塊進度：{done_n:,} / {S['chunks_total']:,}（{done_n * 100.0 / S['chunks_total']:.1f}%）{eta}"
        S["prog_prev"] = done_n
    L = [f"【照片搬移回報】{now()}   第二階段（磁區讀取）{S['stage']}",
         f"狀態：{st}",
         f"目前檔案：{S['cur'] or '—'}",
         prog,
         db_,
         f"本次累計：讀取 {S['read_bytes'] / 1e9:.2f} GB，讀取次數 {S['reads']}，失敗 {S['fails']}，"
         f"過去 {dt / 60:.0f} 分鐘 {(S['read_bytes'] - PREV['bytes']) / 1e6 / dt:.1f} MB/s",
         f"最近一次讀到的檔案：{S['last_good'] or '—'}"]
    if S["note"]:
        L.append(f"備註：{S['note']}")
    PREV.update(t=t, bytes=S["read_bytes"])
    return "\n".join(x for x in L if x)


def emit(c=None):
    text = report(c)
    print("\n" + "=" * 70 + "\n" + text + "\n" + "=" * 70 + "\n", flush=True)
    log(text)
    try:
        with open(REPORT_TXT, "w", encoding="utf-8-sig") as f:
            f.write(text + "\n")
    except OSError:
        pass


def stall_watchdog(c_factory, mode):
    """給回報執行緒用：一次讀取卡超過 STALL_KILL_SEC → 位置記進黑名單 → 強制結束本程式 → 小幫手用排程工作重啟，新程式在這裡直接當壞區。
    自動重啟後如果連一次都還沒讀成功就又卡住，多半是硬碟卡死而不是這個位置壞了：不再強制跳過，等使用者拔電源。"""
    t0, off = S["read_start"], S["read_off"]
    if not t0 or time.time() - t0 < STALL_KILL_SEC or watchdog["fired"] == off:
        return False
    watchdog["fired"] = off
    mins = int((time.time() - t0) // 60)
    if S["relaunched"] and S["reads"] - S["fails"] == 0:
        say(f"！位移 {off} 的讀取已經卡 {mins} 分鐘，而且自動重啟後還沒有任何一次讀取成功：疑似硬碟卡死，不再自動強制跳過。"
            f"請拔掉硬碟電源休息後再重新啟動")
        return False
    c = c_factory()
    stuck = json.loads(meta_get(c, "stuck_offs", "[]"))
    stuck.append(off)
    meta_set(c, "stuck_offs", json.dumps(stuck))
    meta_set(c, "relaunched", S["relaunched"] + 1)
    say(f"！位移 {off} 的讀取（{S['cur']}）已經卡 {mins} 分鐘：這個位置記進黑名單（重啟後直接當壞區），強制結束程式並自動重新啟動接續")
    emit(c)
    c.close()
    relaunch_self(mode)
    return True


def relaunch_self(mode):
    """交給一個獨立的小幫手：強制結束本程式（主執行緒卡在核心的讀取裡，自己結束不了）、等它真的消失、再用排程工作重啟同一個模式。"""
    cmd = [sys.executable, os.path.abspath(__file__), "--relaunch-after", str(os.getpid()), "--relaunch-mode", mode]
    subprocess.Popen(cmd, creationflags=subprocess.DETACHED_PROCESS | subprocess.CREATE_NEW_PROCESS_GROUP,
                     stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, close_fds=True)


def pid_alive(pid):
    out = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"], capture_output=True, text=True, errors="ignore").stdout
    return f'"{pid}"' in out and "python" in out.lower()


def relaunch_helper(pid, mode):
    """小幫手：結束 pid、等它真的消失（卡在核心的讀取要等 Windows 放棄，可能十幾分鐘）、再 schtasks /run 重啟。期間每分鐘更新進度回報.txt。"""
    subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)
    t0 = time.time()
    log(f"[{now()[11:]}] 小幫手：已送出強制結束 PID {pid}，等它結束後用排程工作重啟 {mode}")
    while pid_alive(pid):
        time.sleep(5)
        if int(time.time() - t0) % 60 < 5:
            try:
                with open(REPORT_TXT, "w", encoding="utf-8-sig") as f:
                    f.write(f"【照片搬移回報】{now()}   第二階段（磁區讀取）\n狀態：卡住的位置已記進黑名單，程式強制結束中：卡在核心裡的那次讀取要等 Windows 放棄才會真的結束"
                            f"（已等 {int((time.time() - t0) // 60)} 分鐘）。結束後會自動重新啟動並跳過那個位置，不用做任何事\n")
            except OSError:
                pass
    out = subprocess.run(["schtasks", "/run", "/tn", "PhotoRescue2-" + mode.capitalize()], capture_output=True, text=True, errors="ignore")
    log(f"[{now()[11:]}] 小幫手：PID {pid} 已結束（等了 {int((time.time() - t0) // 60)} 分鐘），重啟 {mode}：{(out.stdout + out.stderr).strip()}")


def reporter(stop, c_factory, mode=""):
    last = 0.0
    while not stop.is_set():
        if time.time() - last >= REPORT_EVERY_SEC:
            last = time.time()
            try:
                c = c_factory()
                emit(c)
                c.close()
            except Exception as e:
                print(f"（回報失敗：{e}）")
        try:
            stall_watchdog(c_factory, mode)
        except Exception as e:
            print(f"（卡住偵測失敗：{e}）")
        try:
            ctypes.windll.kernel32.SetConsoleTitleW(f"[第二階段] {S['state'][:16]} | {os.path.basename(S['cur'])} | 救回{S['done']} 壞區{S['bad_chunks']}")
        except Exception:
            pass
        stop.wait(2)


def write_lock():
    os.makedirs(WORK, exist_ok=True)
    try:
        pid = int(open(LOCK).read().strip())
        out = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"], capture_output=True, text=True, errors="ignore").stdout
        if f'"{pid}"' in out and "python" in out.lower() and pid != os.getpid():
            raise SystemExit(f"已經有一份搬移/救援程式在執行（PID {pid}）。第二階段要等它結束。")
    except (OSError, ValueError):
        pass
    with open(LOCK, "w") as f:
        f.write(str(os.getpid()))


# ---------- 階段 0 / 驗證 / 執行 ----------
def step0(c):
    if c.execute("SELECT COUNT(*) FROM chunks WHERE state IN ('good','bad','trimmed')").fetchone()[0]:
        raise SystemExit("phase2.sqlite 裡已經有 run 的讀取進度，重做 step0 會把它清掉。真的要從頭來，請先把 C:\\照片救援\\phase2.sqlite 改名或刪除。")
    S["stage"] = "階段 0：盤點"
    S["state"] = "讀取 MFT"
    rd, vol, sector = open_real()
    say(f"磁區 {sector} bytes，叢集 {vol.cluster} bytes，MFT 記錄 {vol.rec_size} bytes，分割區位移 {vol.part_off}")
    vol.load_mft(progress=lambda d, t: S.update(cur=f"MFT {d / 1e6:.0f} / {t / 1e6:.0f} MB"))
    say(f"MFT 共 {vol.mft_bytes / 1e6:.0f} MB，記錄 {len(vol.records)} 筆，讀不到/損毀 {len(vol.bad_recs)} 筆")
    S["state"] = "建立計畫、比對 G 槽"
    st = build_plan(vol, c)
    meta_set(c, "plan_time", now())
    meta_set(c, "plan_stats", json.dumps(st, ensure_ascii=False))
    # 各第一層資料夾統計
    per = c.execute("""SELECT substr(path, 4, instr(substr(path, 4), '\\') - 1) AS top,
                       COUNT(*), SUM(status='done_before'), SUM(status='new'), SUM(CASE WHEN status='new' THEN size ELSE 0 END)
                       FROM files GROUP BY top ORDER BY top""").fetchall()
    lines = [f"===== 階段 0 盤點結果 {now()} =====",
             f"MFT 記錄 {len(vol.records)} 筆，讀不到/損毀 {len(vol.bad_recs)} 筆（這些檔案連名字都找不回來，屬於階段 D）",
             f"範圍內檔案 {st['total']}：之前已在 G 槽 {st['done_before']}，要讀 {st['to_read']}（{st['to_read_bytes'] / 1e9:.1f} GB），"
             f"不支援（壓縮/加密/屬性清單損毀）{st['unsupported']}，找不到所屬資料夾 {st['orphan']}",
             "第一層資料夾 | 檔案數 | 已在G槽 | 要讀 | 要讀GB"]
    for top, n, done_b, new, nb in per:
        lines.append(f"  {top} | {n} | {done_b} | {new} | {(nb or 0) / 1e9:.2f}")
    text = "\n".join(lines)
    say(text)
    with open(os.path.join(WORK, "第二階段盤點.txt"), "w", encoding="utf-8-sig") as f:
        f.write(text + "\n")
    # 抽樣驗證：從「之前已在 G 槽」的檔案抽 VERIFY_N 個，從磁區重建、逐位元組比對
    S["stage"] = "階段 0：抽樣驗證"
    rows = c.execute("SELECT mft, path, size, dst FROM files WHERE status='done_before' AND size>0 ORDER BY RANDOM() LIMIT ?", (VERIFY_N,)).fetchall()
    ok = mismatch = unreadable = 0
    for mft, path, size, dst in rows:
        S["cur"] = path
        r = vol.records[mft]
        try:
            data = r["data"]["value"] if r["data"]["res"] else vol.read_extents(vol.runs_to_extents(r["data"]["runs"]), size)
            with open(lp(dst), "rb") as f:
                ref = f.read()
            if data == ref:
                ok += 1
            else:
                mismatch += 1
                say(f"[驗證不符] {path}")
        except ReadError:
            unreadable += 1
    verdict = f"抽樣驗證：{len(rows)} 個檔案，一致 {ok}，不一致 {mismatch}，讀不到 {unreadable}"
    say(verdict)
    meta_set(c, "verify", verdict)
    meta_set(c, "verify_ok", "1" if mismatch == 0 and ok >= min(50, len(rows)) else "0")
    S["state"] = "階段 0 完成，等使用者確認後再執行 run"
    rd._close()
    return mismatch == 0


def run(c):
    if meta_get(c, "verify_ok") != "1":
        raise SystemExit("階段 0 的抽樣驗證還沒通過，不能開始正式讀取。請先執行 step0。")
    rd, vol, sector = open_real()
    S["state"] = "執行中"
    # 需要 Volume 的 records 來讀 attribute？不用：chunk 表已經有裝置位移。
    S["stage"] = "階段 A：順序讀取"
    pass_read(vol, c, ("new",), use_skip=True, sector=sector)
    write_xlsx(c)                                             # 每個階段結束就更新 Excel，跑好幾天也看得到中途結果
    S["stage"] = "階段 A2：重試跳過的區塊"
    pass_read(vol, c, ("skipped",), use_skip=False, sector=sector)
    write_xlsx(c)
    S["stage"] = "階段 B：修剪壞區邊界"
    pass_trim(vol, c, sector=sector)
    S["stage"] = "輸出"
    finalize_all(c)
    write_xlsx(c)
    S["state"] = "已完成"
    rd._close()


def write_xlsx(c):
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return
    rows = c.execute("SELECT path, size, status, bad_bytes, note, dst FROM files WHERE status NOT IN ('done','done_before') ORDER BY path").fetchall()
    try:
        wb = load_workbook(XLSX) if os.path.exists(XLSX) else Workbook()
        if "第二階段結果" in wb.sheetnames:
            del wb["第二階段結果"]
        if not os.path.exists(XLSX):
            wb.remove(wb.active)
        ws = wb.create_sheet("第二階段結果", 0)
        ws.append(["來源路徑", "大小(MB)", "結果", "壞區(KB)", "說明", "存放位置"])
        names = {"damaged": "損壞（未存）", "partial_saved": "部分損壞（已存）", "unsupported": "不支援", "dst_error": "寫入G失敗", "new": "未完成"}
        fills = {"damaged": "FFC7CE", "partial_saved": "FCE4D6", "unsupported": "FFC7CE", "dst_error": "FFF2CC", "new": "FFF2CC"}
        for path, size, status, bb, note, dst in rows:
            ws.append([path, round(size / 1048576, 2), names.get(status, status), round(bb / 1024, 1), note, dst if status == "partial_saved" else ""])
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fills.get(status, "FFFFFF"))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, w in zip("ABCDEF", (70, 10, 16, 10, 40, 70)):
            ws.column_dimensions[col].width = w
        wb.save(XLSX)
    except Exception as e:
        S["note"] = f"xlsx 更新失敗：{e}"


def start_stop(mode, what):
    """不需要系統管理員：透過已登錄的排程工作啟動/停止（給執行 AI 用，只要跑 python 指令）。"""
    task = "PhotoRescue2-" + what.capitalize()
    if mode == "stop":
        out = subprocess.run(["schtasks", "/end", "/tn", task], capture_output=True, text=True, errors="ignore")
        return print((out.stdout + out.stderr).strip() or f"已送出停止 {task}")
    try:
        pid = int(open(LOCK).read().strip())
        tl = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"], capture_output=True, text=True, errors="ignore").stdout
        if f'"{pid}"' in tl and "python" in tl.lower():
            return print(f"已經有搬移/救援程式在執行（PID {pid}），沒有啟動新的。")
    except (OSError, ValueError):
        pass
    out = subprocess.run(["schtasks", "/run", "/tn", task], capture_output=True, text=True, errors="ignore")
    if out.returncode != 0:
        return print(f"啟動失敗：{(out.stdout + out.stderr).strip()}\n請確認已用系統管理員執行過「登錄第二階段排程.bat」。")
    for _ in range(12):
        time.sleep(5)
        try:
            pid = int(open(LOCK).read().strip())
            tl = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"], capture_output=True, text=True, errors="ignore").stdout
            if f'"{pid}"' in tl and "python" in tl.lower():
                return print(f"{task} 已啟動（PID {pid}）。之後用 rescue.py --wait-report 看每 5 分鐘的回報。")
        except (OSError, ValueError):
            pass
    print(f"{task} 已送出啟動，但 60 秒內沒看到程式在跑。看看桌面有沒有新視窗、或執行 python rescue2.py report。")


def print_report(c):
    print(report(c))
    print("盤點：", meta_get(c, "plan_stats", "（還沒做階段 0）"))
    print("驗證：", meta_get(c, "verify", "（還沒做）"))
    if STUCK:
        print(f"卡住黑名單：{len(STUCK)} 處（讀取卡超過 {STALL_KILL_SEC // 60} 分鐘被強制跳過的位置）")


# ---------- 自我測試：自製迷你 NTFS 映像 ----------
def _make_image(sector=4096):
    """回傳 (image bytes, 期望資料 dict, bad_offsets set)。叢集 = 1 磁區 = 4096；MFT 記錄 4096；分割區從磁區 8 開始。"""
    part = 8 * sector
    mft_lcn, nrec = 16, 64
    data_lcn = 100
    total = 400
    img = bytearray(part + total * sector)
    img[0x1BE + 4] = 0x07
    struct.pack_into("<II", img, 0x1BE + 8, 8, total)
    img[510:512] = b"\x55\xAA"
    boot = bytearray(sector)
    boot[3:11] = b"NTFS    "
    struct.pack_into("<H", boot, 0x0B, sector)
    boot[0x0D] = 1
    struct.pack_into("<QQQ", boot, 0x28, total, mft_lcn, mft_lcn + 60)
    struct.pack_into("<b", boot, 0x40, 1)                   # 1 cluster per MFT record = 4096
    img[part:part + sector] = boot
    files = {}
    expect = {}
    bad = set()
    rng = random.Random(7)
    next_lcn = [data_lcn]

    def alloc(n):
        l = next_lcn[0]
        next_lcn[0] += n
        return l

    def attr(atype, body, name=b"", flags=0):
        nlen = len(name) // 2
        hdr_len = 0x18 if nlen == 0 else 0x18 + len(name)
        hdr_len = (hdr_len + 7) & ~7
        return atype, body, name, flags, hdr_len

    def resident(atype, value, flags=0):
        hdr = 0x18
        total_len = (hdr + len(value) + 7) & ~7
        b = bytearray(total_len)
        struct.pack_into("<IIBBHHH", b, 0, atype, total_len, 0, 0, hdr, flags, 0)
        struct.pack_into("<IH", b, 0x10, len(value), hdr)
        b[hdr:hdr + len(value)] = value
        return bytes(b)

    def encode_runs(runs):
        out, prev = bytearray(), 0
        for lcn, ln in runs:
            lb = (ln.bit_length() + 7) // 8 or 1
            if lcn is None:
                out.append(lb)
                out += ln.to_bytes(lb, "little")
            else:
                d = lcn - prev
                ob = 1
                while not (-(1 << (8 * ob - 1)) <= d < (1 << (8 * ob - 1))):
                    ob += 1
                out.append(lb | (ob << 4))
                out += ln.to_bytes(lb, "little") + d.to_bytes(ob, "little", signed=True)
                prev = lcn
        out.append(0)
        return bytes(out)

    def nonresident(atype, runs, size, svcn=0, init=None, flags=0):
        rb = encode_runs(runs)
        hdr = 0x40
        total_len = (hdr + len(rb) + 7) & ~7
        b = bytearray(total_len)
        struct.pack_into("<IIBBHHH", b, 0, atype, total_len, 1, 0, hdr, flags, 0)
        nclus = sum(ln for _, ln in runs)
        struct.pack_into("<QQHH", b, 0x10, svcn, svcn + nclus - 1, hdr, 0)
        struct.pack_into("<QQQ", b, 0x28, nclus * sector, size, size if init is None else init)
        b[hdr:hdr + len(rb)] = rb
        return bytes(b)

    def fname(parent, name, ns=1, isdir=False):
        v = bytearray(0x42 + len(name) * 2)
        struct.pack_into("<Q", v, 0, parent | (1 << 48))
        struct.pack_into("<I", v, 0x38, 0x10000000 if isdir else 0x20)
        v[0x40], v[0x41] = len(name), ns
        v[0x42:] = name.encode("utf-16-le")
        return resident(0x30, bytes(v))

    def stdinfo(mtime):
        v = bytearray(0x48)
        ft = int(mtime * 1e7) + FILETIME_EPOCH
        struct.pack_into("<QQQQ", v, 0, ft, ft, ft, ft)
        return resident(0x10, bytes(v))

    def record(no, attrs, inuse=True, isdir=False, base=0):
        rec = bytearray(sector)
        rec[:4] = b"FILE"
        usa_cnt = sector // 512 + 1
        struct.pack_into("<HH", rec, 4, 0x30, usa_cnt)
        struct.pack_into("<HHHH", rec, 0x10, 1, 1, 0x30 + usa_cnt * 2 + 6 & ~7, (1 if inuse else 0) | (2 if isdir else 0))
        first = (0x30 + usa_cnt * 2 + 7) & ~7
        struct.pack_into("<H", rec, 0x14, first)
        struct.pack_into("<Q", rec, 0x20, base)
        off = first
        for a in attrs:
            rec[off:off + len(a)] = a
            off += len(a)
        rec[off:off + 4] = b"\xFF\xFF\xFF\xFF"
        struct.pack_into("<II", rec, 0x18, off + 8, sector)
        usn = b"\x11\x22"                                    # 套 fixup：每 512 的最後 2 bytes 存進 USA，換成 USN
        rec[0x30:0x32] = usn
        for i in range(1, usa_cnt):
            pos = i * 512 - 2
            rec[0x30 + i * 2:0x30 + i * 2 + 2] = rec[pos:pos + 2]
            rec[pos:pos + 2] = usn
        img[part + (mft_lcn * sector) + no * sector:part + (mft_lcn * sector) + (no + 1) * sector] = rec

    def put_data(runs, data):
        pos = 0
        for lcn, ln in runs:
            if lcn is not None:
                img[part + lcn * sector:part + lcn * sector + ln * sector] = data[pos:pos + ln * sector].ljust(ln * sector, b"\0")
            pos += ln * sector

    T = 1_600_000_000.0
    record(0, [stdinfo(T), fname(5, "$MFT", 3), nonresident(0x80, [(mft_lcn, nrec)], nrec * sector)])
    for i in range(1, 5):
        record(i, [stdinfo(T), fname(5, f"$Sys{i}", 3), resident(0x80, b"")])
    record(5, [stdinfo(T), fname(5, ".", 3, True), resident(0x90, b"\0" * 16)], isdir=True)
    for i in range(6, 16):
        record(i, [stdinfo(T), fname(5, f"$Sys{i}", 3), resident(0x80, b"")])
    record(16, [stdinfo(T), fname(5, "照片", 1, True), resident(0x90, b"\0" * 16)], isdir=True)
    record(17, [stdinfo(T), fname(16, "2015旅遊", 1, True), resident(0x90, b"\0" * 16)], isdir=True)
    record(23, [stdinfo(T), fname(5, "家人A", 1, True), resident(0x90, b"\0" * 16)], isdir=True)
    record(30, [stdinfo(T), fname(5, "不要的", 1, True), resident(0x90, b"\0" * 16)], isdir=True)

    def add_file(no, parent, name, size, runs=None, mtime=None, extra_names=(), bad_sectors=(), init=None, ns=1):
        mtime = T + no if mtime is None else mtime
        d = bytes(rng.getrandbits(8) for _ in range(size))
        if runs is None:
            attrs = [stdinfo(mtime)] + [fname(parent, name, ns)] + [fname(p, n, s) for p, n, s in extra_names] + [resident(0x80, d)]
        else:
            put_data(runs, d)
            attrs = [stdinfo(mtime)] + [fname(parent, name, ns)] + [fname(p, n, s) for p, n, s in extra_names] + [nonresident(0x80, runs, size, init=init)]
        record(no, attrs)
        for lcn, ln in (runs or []):
            if lcn is not None:
                for s in bad_sectors:
                    if lcn <= s < lcn + ln:
                        bad.add(part + s * sector)
        expect[name] = dict(no=no, data=d, size=size, mtime=mtime, parent=parent)
        return d

    a_runs = [(alloc(3), 3)]
    add_file(18, 17, "a.jpg", 3 * sector - 100, a_runs)                          # 多叢集、非整數大小
    add_file(19, 17, "b.txt", 5)                                                  # resident
    s1 = alloc(1); alloc(1); s2 = alloc(1)
    add_file(24, 23, "c.jpg", 3 * sector, [(s1, 1), (None, 1), (s2, 1)])         # 稀疏：中間那叢集讀出來是 0
    expect["c.jpg"]["data"] = expect["c.jpg"]["data"][:sector] + bytes(sector) + expect["c.jpg"]["data"][2 * sector:]
    dl = alloc(2)
    add_file(25, 17, "d.nef", 2 * sector, [(dl, 2)], bad_sectors=[dl + 1])       # 壞磁區 → 損壞不存
    el = alloc(40)
    add_file(26, 17, "e.jpg", 40 * sector, [(el, 40)], bad_sectors=[el + 20])     # 1 壞磁區 = 2.5% → 修剪後存部分損壞
    fl = alloc(20)
    add_file(27, 17, "f.jpg", 20 * sector, [(fl, 20)], bad_sectors=list(range(fl + 5, fl + 15)))   # 50% 壞 → 不存
    add_file(28, 17, "LONGNAME.JPG", 300, extra_names=[(17, "很長的檔名 照片.jpg", 1)], ns=2)  # DOS + Win32 名：用長檔名
    expect["很長的檔名 照片.jpg"] = expect.pop("LONGNAME.JPG")
    gl = alloc(2)
    add_file(29, 17, "已存在.jpg", 2 * sector, [(gl, 2)])                           # G 槽已有 → 不讀
    add_file(31, 30, "x.jpg", 100)                                                # 範圍外
    record(32, [stdinfo(T), fname(17, "deleted.jpg", 1), resident(0x80, b"gone")], inuse=False)   # 已刪除
    add_file(33, 17, "Thumbs.db", 10)                                             # 略過的名字
    hl = alloc(3)
    add_file(34, 17, "h.jpg", 3 * sector, [(hl, 3)], init=1 * sector)              # 只初始化 1 叢集，其餘視為 0
    expect["h.jpg"]["data"] = expect["h.jpg"]["data"][:sector] + bytes(2 * sector)
    # 屬性清單：主記錄 35 有 $ATTRIBUTE_LIST，$DATA 前 4 叢集在 35、後 2 叢集在擴充記錄 36
    b1, b2 = alloc(4), alloc(2)
    big = bytes(rng.getrandbits(8) for _ in range(6 * sector - 7))
    put_data([(b1, 4)], big[:4 * sector])
    put_data([(b2, 2)], big[4 * sector:])
    al = bytearray()
    for atype, svcn, ref in ((0x10, 0, 35), (0x30, 0, 35), (0x80, 0, 35), (0x80, 4, 36)):
        e = bytearray(0x20)
        struct.pack_into("<IHBB", e, 0, atype, 0x20, 0, 0x1A)
        struct.pack_into("<QQH", e, 8, svcn, ref | (1 << 48), 0)
        al += e
    d35 = nonresident(0x80, [(b1, 4)], len(big))
    record(35, [stdinfo(T), fname(17, "big.mov", 1), resident(0x20, bytes(al)), d35])
    record(36, [nonresident(0x80, [(b2, 2)], len(big), svcn=4)], base=35)
    expect["big.mov"] = dict(no=35, data=big, size=len(big), mtime=T, parent=17)
    # 孤兒：父資料夾 40 的 MFT 記錄在壞磁區
    record(40, [stdinfo(T), fname(16, "壞掉的資料夾", 1, True), resident(0x90, b"\0" * 16)], isdir=True)
    bad.add(part + (mft_lcn + 40) * sector)
    ol = alloc(1)
    add_file(41, 40, "orphan.jpg", sector, [(ol, 1)])
    # 連續壞區（測跳過）：3 個相鄰檔案各 1 叢集全壞
    for i, no in enumerate((42, 43, 44)):
        l = alloc(1)
        add_file(no, 17, f"bad{i}.jpg", sector, [(l, 1)], bad_sectors=[l])
    add_file(45, 17, "r.txt", 2000)                                               # resident 但超過 1 KB
    record(46, [stdinfo(T), fname(17, "corrupt.jpg", 1), nonresident(0x80, [(alloc(1), 1)], 10 ** 12)])   # 損毀記錄：大小 1 TB 但只配置 1 叢集
    return bytes(img), expect, bad, part


def selftest():
    import tempfile, faulthandler
    global WORK, DB, TMP, LOG, LOCK, REPORT_TXT, XLSX, CHUNK, SKIP_MIN, SKIP_MAX, TRIM_STEP, READ_MIN, POLL_SEC, DST_RETRY_SLEEP, STUCK_SEC
    faulthandler.dump_traceback_later(90, exit=True)         # 測試卡住就印出卡在哪
    DST_RETRY_SLEEP = 0
    base = tempfile.mkdtemp(prefix="rescue2_")
    WORK, TMP = os.path.join(base, "work"), os.path.join(base, "tmp")
    DB, LOG, LOCK = os.path.join(WORK, "p2.sqlite"), os.path.join(WORK, "log.txt"), os.path.join(WORK, "lock")
    REPORT_TXT, XLSX = os.path.join(base, "回報.txt"), os.path.join(base, "清單.xlsx")
    CHUNK, SKIP_MIN, SKIP_MAX, TRIM_STEP = 8192, 4096, 65536, 4096   # 小尺寸才測得到分塊/跳過/修剪
    dst_root = os.path.join(base, "G")
    img, expect, bad, part = _make_image()
    sector = 4096
    disk = FileDisk(img, sector, bad)
    rd = Reader(lambda: disk, probe_off=part + 16 * sector, sector=sector)
    vol = Volume(rd, part, sector)
    assert vol.bps == sector and vol.cluster == sector and vol.rec_size == sector and vol.mft_lcn == 16
    vol.load_mft()
    assert 40 in vol.bad_recs and 35 in vol.records and 36 in vol.records, vol.bad_recs
    assert vol.records[35]["data"]["size"] == expect["big.mov"]["size"] and len(vol.records[35]["data"]["runs"]) == 2
    assert vol.path_of(18)[0] == ["照片", "2015旅遊", "a.jpg"]
    assert vol.path_of(28)[0][-1] == "很長的檔名 照片.jpg"
    assert vol.path_of(41)[1] is True and vol.path_of(41)[0][0] == ORPHAN_DIR
    # G 槽已有「已存在.jpg」
    ex = expect["已存在.jpg"]
    p = os.path.join(dst_root, "照片", "2015旅遊", "已存在.jpg")
    os.makedirs(os.path.dirname(p))
    open(p, "wb").write(ex["data"])
    os.utime(p, (ex["mtime"], ex["mtime"]))
    c = db_open()
    st = build_plan(vol, c, want_tops=["照片", "家人A"], dst_root=dst_root)
    assert st["done_before"] == 1 and st["orphan"] == 1 and st["skipped_name"] == 1 and st["bad_recs"] == 1 and st["unsupported"] == 1, st
    names = {r[0] for r in c.execute("SELECT path FROM files").fetchall()}
    assert not any("x.jpg" in n or "deleted" in n or "Thumbs" in n for n in names), names
    assert c.execute("SELECT COUNT(*) FROM chunks WHERE state='zero'").fetchone()[0] == 2      # c.jpg 稀疏 + h.jpg 未初始化
    # 階段 A：跳過機制
    disk.reads.clear()
    pass_read(vol, c, ("new",), use_skip=True, sector=sector)
    sk = c.execute("SELECT COUNT(*) FROM chunks WHERE state='skipped'").fetchone()[0]
    assert sk >= 1, "連續壞區後面應該有被跳過的區塊"
    pass_read(vol, c, ("skipped",), use_skip=False, sector=sector)
    assert c.execute("SELECT COUNT(*) FROM chunks WHERE state IN ('new','skipped')").fetchone()[0] == 0
    pass_trim(vol, c, sector=sector)
    assert c.execute("SELECT COUNT(*) FROM chunks WHERE state='trimmed'").fetchone()[0] >= 1
    disk.reads.clear()
    pass_trim(vol, c, sector=sector)                        # 中斷後重跑：修剪過的壞區不能再讀一次
    assert not disk.reads, disk.reads
    finalize_all(c)
    status = dict(c.execute("SELECT path, status FROM files").fetchall())
    for name in ("a.jpg", "b.txt", "c.jpg", "很長的檔名 照片.jpg", "h.jpg", "big.mov", "orphan.jpg", "r.txt"):
        row = [(pth, s) for pth, s in status.items() if pth.endswith("\\" + name)][0]
        assert row[1] == "done", row
        pth = row[0]
        dst = os.path.join(dst_root, *pth[3:].split("\\"))
        with open(dst, "rb") as f:
            got = f.read()
        assert got == expect[name]["data"], f"{name} 內容不符 ({len(got)} vs {len(expect[name]['data'])})"
        assert abs(os.path.getmtime(dst) - expect[name]["mtime"]) <= 2
    sfx = {pth.rsplit("\\", 1)[1]: s for pth, s in status.items()}
    assert sfx["d.nef"] == "damaged" and sfx["f.jpg"] == "damaged" and sfx["e.jpg"] == "partial_saved", sfx
    assert all(sfx[f"bad{i}.jpg"] == "damaged" for i in range(3))
    assert sfx["corrupt.jpg"] == "unsupported", sfx
    pe = os.path.join(dst_root, "照片", "2015旅遊", "e_部分損壞.jpg")
    with open(pe, "rb") as f:
        got = f.read()
    exp = bytearray(expect["e.jpg"]["data"])
    exp[20 * sector:21 * sector] = bytes(sector)
    assert got == bytes(exp), "e.jpg 應該只有第 20 個磁區補 0"
    assert not os.path.exists(os.path.join(dst_root, "照片", "2015旅遊", "d.nef"))
    assert not os.listdir(TMP) if os.path.exists(TMP) else True
    assert sfx["已存在.jpg"] == "done_before"
    write_xlsx(c)
    from openpyxl import load_workbook
    ws = load_workbook(XLSX)["第二階段結果"]
    assert ws.max_row == 1 + 7, ws.max_row                  # d, f, e(部分損壞), bad0-2, corrupt = 7 列
    # 斷線後自動接續
    disk.gone = True
    threading.Timer(0.5, lambda: setattr(disk, "gone", False)).start()
    POLL_SEC = 1
    assert rd.read(part, sector)[3:11] == b"NTFS    "
    # 卡死：連探測區都讀不到 → 停下來等（不做任何 USB 重置）→ 探測區讀得到就恢復 → 同一塊再讀一次
    probe = rd.probe_off
    disk.bad.add(probe)
    threading.Timer(2.5, lambda: disk.bad.discard(probe)).start()
    rd.streak = HUNG_STREAK
    bad_off = next(iter(sorted(b for b in disk.bad if b != probe)))
    try:
        rd.read(bad_off, sector)
        raise AssertionError("壞區應該還是要丟 ReadError")
    except ReadError:
        pass
    assert S["hung_since"] == 0.0 and S["state"] == "執行中", S["state"]
    assert disk.reads[-1] == (bad_off, sector) and disk.reads.count((bad_off, sector)) >= 2, "恢復後應該重讀同一塊"
    # 卡了很久才斷線：這塊直接標壞（丟 ReadError 而不是 DiskGone），不重讀
    STUCK_SEC = -1
    disk.gone = True
    try:
        rd.read(part, sector)
        raise AssertionError("應該丟 ReadError")
    except DiskGone:
        raise AssertionError("卡住後斷線不應該當成一般斷線重讀")
    except ReadError:
        pass
    STUCK_SEC = 180
    disk.gone = False
    assert rd.read(part, sector)[3:11] == b"NTFS    "
    # 卡住看門狗：一次讀取卡超過 STALL_KILL_SEC → 位置進黑名單（meta）、啟動強制結束+重啟；同一位置只做一次；
    # 自動重啟後一次都沒讀成功就又卡住 → 不做；黑名單裡的位置之後直接丟 ReadError 不讀
    relaunches = []
    real_relaunch = globals()["relaunch_self"]
    globals()["relaunch_self"] = lambda mode: relaunches.append(mode)
    try:
        S.update(read_start=time.time() - STALL_KILL_SEC - 1, read_off=part + 300 * sector, relaunched=0)
        assert stall_watchdog(lambda: sqlite3.connect(DB), "run") is True and relaunches == ["run"]
        assert json.loads(meta_get(c, "stuck_offs")) == [part + 300 * sector] and meta_get(c, "relaunched") == "1"
        assert stall_watchdog(lambda: sqlite3.connect(DB), "run") is False and len(relaunches) == 1
        S.update(read_start=time.time() - STALL_KILL_SEC - 1, read_off=part + 301 * sector, relaunched=1, reads=5, fails=5)
        assert stall_watchdog(lambda: sqlite3.connect(DB), "run") is False and len(relaunches) == 1
        S.update(read_off=part + 302 * sector, reads=6)
        assert stall_watchdog(lambda: sqlite3.connect(DB), "run") is True and len(relaunches) == 2
        S.update(read_start=0.0, read_off=-1, relaunched=0)
        STUCK.add(part + 300 * sector)
        disk.reads.clear()
        try:
            rd.read(part + 300 * sector, sector)
            raise AssertionError("黑名單位置應該直接丟 ReadError")
        except ReadError as e:
            assert e.code == -4 and not disk.reads, "黑名單位置不可以真的去讀硬碟"
        STUCK.clear()
    finally:
        globals()["relaunch_self"] = real_relaunch
        watchdog["fired"] = None
    rep = report(c)
    assert "已完整救回" in rep
    print("盤點:", st)
    print("SELFTEST OK")
    shutil.rmtree(base, ignore_errors=True)


# ---------- 主程式 ----------
def main():
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="第二階段：直接讀硬碟磁區的照片救援")
    ap.add_argument("mode", nargs="?", choices=["step0", "run", "report", "start", "stop"])
    ap.add_argument("what", nargs="?", choices=["step0", "run"], help="start/stop 用：要啟動/停止哪個排程工作")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--relaunch-after", type=int, metavar="PID", help=argparse.SUPPRESS)       # 程式自己用：小幫手模式
    ap.add_argument("--relaunch-mode", choices=["step0", "run"], help=argparse.SUPPRESS)
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if a.relaunch_after:
        return relaunch_helper(a.relaunch_after, a.relaunch_mode or "run")
    if not a.mode:
        return ap.print_help()
    if a.mode in ("start", "stop"):
        return start_stop(a.mode, a.what or "run")
    c = db_open()
    STUCK.update(json.loads(meta_get(c, "stuck_offs", "[]")))
    S["relaunched"] = int(meta_get(c, "relaunched", "0"))
    if a.mode == "report":
        return print_report(c)
    if not is_admin():
        raise SystemExit("需要系統管理員權限才能直接讀取實體磁碟。請用系統管理員身分登錄的排程工作執行（單獨執行模式未在真實硬碟驗證，建議改用 rescue_dd.py）。")
    write_lock()
    try:
        ctypes.windll.kernel32.SetThreadExecutionState(0x80000001)
    except Exception:
        pass
    say(f"===== 第二階段（磁區讀取）{a.mode} 開始（PID {os.getpid()}）=====")
    set_disk_offline(True)
    say("硬碟已設為離線（Windows 不會再碰它）")
    stop = threading.Event()
    th = threading.Thread(target=reporter, args=(stop, lambda: sqlite3.connect(DB), a.mode), daemon=True)
    th.start()
    try:
        if a.mode == "step0":
            step0(c)
        else:
            run(c)
    except KeyboardInterrupt:
        S["state"] = "已手動中斷（重新執行會從進度接續）"
    except BaseException as e:
        S["state"] = f"程式發生錯誤而停止：{e!r}"
        raise
    finally:
        stop.set()
        th.join(5)
        if watchdog["fired"] is None:
            try:
                meta_set(c, "relaunched", 0)    # 正常結束（含手動中斷）：下次是使用者自己啟動的，不算接續
            except Exception:
                pass
        emit(c)
        try:
            os.remove(LOCK)
        except OSError:
            pass
    say(f"===== {a.mode} 結束：{S['state']} =====")


if __name__ == "__main__":
    main()
