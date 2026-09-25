# -*- coding: utf-8 -*-
r"""壞軌外接硬碟照片搬移程式（操作手冊：docs/方法一_Windows逐檔複製.md）

  python rescue.py phase1                  第一階段：搬健康檔，壞檔跳過並記錄（可重跑，自動續傳）
  python rescue.py phase1 --only 資料夾名   只跑 E:\照片 底下某一個資料夾（試跑用）
  python rescue.py phase2                  第二階段：慢慢救壞檔
  python rescue.py --selftest              用假資料自我測試，不碰 E 槽
"""
import argparse, csv, ctypes, datetime, os, shutil, subprocess, sys, threading, time

# ===== 可調參數 =====
SRC_ROOT = r"E:\照片"
DST_ROOT = r"G:\我的雲端硬碟\!家裡硬碟\照片"
WORK = r"C:\照片救援"
REPORT_TXT = r"G:\我的雲端硬碟\!家裡硬碟\進度回報.txt"
XLSX = r"G:\我的雲端硬碟\!家裡硬碟\照片搬移異常與壞軌清單.xlsx"
REPORT_EVERY_SEC = 300        # 每 5 分鐘回報一次
STALL_SEC = 180               # 超過這麼久「完全沒讀到新資料」才算卡住（這顆碟一般壞軌約 50~60 秒就回報錯誤）
                              # 卡住後斷線（拔掉重插或 USB 自己重置）的檔案 → 接回後不重讀，直接記成壞軌
COPY_CHUNK = 1024 * 1024      # 第一階段一次複製 1 MiB，用來追蹤有沒有進度
# 舊清單記錄過、會讓硬碟卡很久的壞檔：第一階段不讀，直接留給第二階段（第一層資料夾, 檔名）
KNOWN_BAD = set()             # 例：{("2016-10-30家族合照", "DSC_0001.JPG"), ("2017-12-20出國旅遊", "DSC_0002.JPG")}
CONSEC_FAIL_LIMIT = 5         # 第一階段同一個資料夾連續失敗幾張 → 這個資料夾剩下的照片整批留給第二階段
                              # （原本是「休息 5 分鐘」，9/14 實測 6 次休息後照樣失敗：是實體壞區，休息沒用）
HUNG_DIR_LIMIT = 3            # 連續幾個資料夾（任何層）都讀不到清單 → 探測硬碟；卡死就撤銷這段期間的紀錄並自動停止
                              # （9/15 02:05 硬碟卡死時，每個資料夾都被誤記成「讀不到」）
HUNG_DEFER_LIMIT = 2          # 連續幾個資料夾都因「連續 5 張失敗」而延後（中間沒有任何一張成功）→ 一樣探測硬碟
DST_RETRY_SLEEP = 60          # 寫入 G 槽失敗（Google Drive 鎖檔、暫存區滿）多久後重試
DST_RETRY_LIMIT = 5           # 重試幾次還是寫不進去 → 記成「目的地寫入失敗」留給第二階段（不是壞軌，不進 Excel）
PRESETS = {"others": ["家人A", "家人B", "家人C", "手機備份", "電子書", "字型", "music", "software", "game"]}
                              # --tops 可以用這裡的名稱（bat 檔用 cmd 的編碼讀，直接寫中文會亂碼）
OTHERS_SRC, OTHERS_DST = "E:\\", r"G:\我的雲端硬碟\!家裡硬碟"   # --tops others 沒指定 --src/--dst 時用的來源與目的地根目錄
FILE_STALL_KILL_SEC = 600     # 第一階段單一檔案完全沒進度超過這麼久 → 記成「卡住(壞軌)」、強制結束程式、自動重新啟動接續（跳過這個檔）
                              # v11：不再用 pnputil 重啟 USB 裝置。9/16 06:22 在讀取卡住時重啟裝置造成藍屏 0x50。
                              # 強制結束後，卡在核心裡的那次讀取要等 Windows 放棄才會真的結束（可能十幾分鐘），小幫手會等它結束再重啟
OFFLINE_POLL_SEC = 30         # 斷線時多久檢查一次
MIN_FREE_GB_ON_G = 20
P2_CHUNK = 1024 * 1024        # 第二階段一次讀 1 MiB
P2_BLOCK = 4096               # 讀不到時改成 4 KiB 慢慢讀
P2_RETRY = 1                  # 每個壞區塊的重試次數（這顆碟每次讀失敗要 60~400 秒，重試太多次只是在傷硬碟）
# 硬碟還有沒有回應的探測檔：一個確定完好、已經搬好的檔案。第二階段有東西失敗時，不經快取讀它開頭 4 KB，
# 30 秒內讀不到 → 判定硬碟卡死，自動停止（把剛才那一項還原成待救援），不會把後面的全部誤判成壞掉
PROBE_FILE = r"E:\照片\2007-06-30活動\DSC00001.JPG"   # 換硬碟時改成一個確定完好、已經搬好的檔案
P2_RETRY_SLEEP = 2
P2_BAD_TIME_CAP_SEC = 1200    # 單一檔案花在壞區的時間上限，超過後壞區直接補 0（好的部分照讀）
P2_FILE_TIME_CAP_SEC = 3600   # 單一檔案總時間上限，超過後剩下全部補 0
P2_REST_BETWEEN_FILES = 10    # 第二階段每個檔案之間休息，讓硬碟降溫
P2_DIR_RETRY = 3
P2_DIR_RETRY_SLEEP = 60
SKIP_NAMES = {"thumbs.db", "desktop.ini", ".ds_store", "@eadir", ".@__thumb"}
# ====================

FIELDS = ["項次", "類型", "第一層資料夾", "來源完整路徑", "檔名", "大小(MB)", "來源修改時間",
          "錯誤代碼", "錯誤類型", "錯誤訊息", "第一階段時間", "狀態", "救回比例%", "目的地檔案", "第二階段完成時間"]
PENDING = "待救援"
KINDS = {23: "壞軌", 1117: "壞軌", 483: "壞軌", 13: "壞軌", 22: "壞軌",   # 13、22：這顆碟退化後常見的讀取失敗
         1392: "結構損壞", 121: "逾時(可能壞軌)", 32: "暫時性(檔案被佔用)"}
TRANSIENT = (32, 121)
FILLS = {PENDING: "FFF2CC", "已完整救回": "C6EFCE", "已可讀取（已重新搬移）": "C6EFCE",
         "無法讀取": "FFC7CE", "資料夾結構損壞，需專業工具": "FFC7CE"}   # 其他（部分救回）= 橘

S = dict(phase="", state="啟動中", top="", top_i=0, top_n=0, idx_in_top=0, cur="", cur_start=0.0, last_progress=0.0,
         copied=0, copied_bytes=0, skipped=0, bad=0, known=0, consec=0, last_bad="", offline_since=0.0,
         relaunched=0,              # 這次執行是第幾次「卡住 → 強制結束 → 自動重啟」之後的接續（0 = 使用者自己啟動的）
         p2_i=0, p2_n=0, p2_pos=0.0, p2_badblk=0, p2_full=0, p2_part=0, p2_fail=0, p2_dirok=0, p2_dirfail=0)
PREV = {"t": time.time(), "copied": 0, "bytes": 0, "bad": 0}
RL = threading.RLock()
records = {}          # 來源完整路徑 -> 一列紀錄；壞軌紀錄.csv 是唯一的資料來源
dirty = [True]        # xlsx 需要重建
notes = []


def lp(p):            # 長路徑前綴，避免超過 260 字元出錯
    return p if p.startswith("\\\\?\\") else "\\\\?\\" + p


def now():
    return datetime.datetime.now().strftime("%Y/%m/%d %H:%M:%S")


def err_code(e):
    return getattr(e, "winerror", None) or e.errno or 0


def dst_of(src):
    rel = src[len(SRC_ROOT):].lstrip("\\")
    return os.path.join(DST_ROOT, rel) if rel else DST_ROOT


def top_of(src):
    return src[len(SRC_ROOT):].lstrip("\\").split("\\")[0]


def partial_name(dst):
    stem, ext = os.path.splitext(dst)
    return stem + "_部分損壞" + ext


def code_of(r):
    try:
        return int(str(r["錯誤代碼"]).split()[0])
    except (ValueError, IndexError):
        return 0


def log(text):
    with RL:
        try:
            with open(os.path.join(WORK, "進度紀錄.log"), "a", encoding="utf-8") as f:
                f.write(text + "\n")
        except OSError:
            pass


def say(msg):
    line = f"[{now()[11:]}] {msg}"
    print(line, flush=True)
    log(line)


def pause(msg, secs):
    S["state"] = msg
    say(msg)
    time.sleep(secs)
    S["state"] = "執行中"


# ---------- 斷線、空間 ----------
def online():
    return os.path.exists(lp(SRC_ROOT)) and os.path.exists(lp(os.path.dirname(DST_ROOT)))


def wait_online():
    prev = S["state"]
    S["offline_since"] = S["offline_since"] or time.time()
    say("！偵測到斷線，暫停等待重新連接（斷線不會記成壞軌）")
    while not online():
        who = "E 槽" if not os.path.exists(lp(SRC_ROOT)) else "G 槽（Google Drive）"
        S["state"] = f"{who}斷線，等待重新連接"
        time.sleep(OFFLINE_POLL_SEC)
    S["state"] = "已重新連接，等待硬碟穩定"
    time.sleep(OFFLINE_POLL_SEC)
    S["offline_since"] = 0.0
    S["state"] = prev
    if S["cur_start"]:
        S["cur_start"] = S["last_progress"] = time.time()
    say("已重新連接")


def free_gb():
    return shutil.disk_usage(os.path.dirname(DST_ROOT)).free / 1e9


def ensure_space():
    try:
        while free_gb() < MIN_FREE_GB_ON_G:
            pause(f"G 槽剩餘空間低於 {MIN_FREE_GB_ON_G} GB，等 Google Drive 上傳釋出暫存空間後自動繼續", 60)
    except OSError:
        pass


def mkdir(d):
    fails = 0
    while True:
        try:
            return os.makedirs(lp(d), exist_ok=True)
        except OSError as e:
            if not online():
                wait_online()
                continue
            fails += 1
            if fails > DST_RETRY_LIMIT:
                raise
            pause(f"G 槽建立資料夾失敗（{e.strerror}），{DST_RETRY_SLEEP} 秒後重試（第 {fails}/{DST_RETRY_LIMIT} 次）", DST_RETRY_SLEEP)


# ---------- 硬碟卡死偵測 ----------
class HungDrive(Exception):
    """連續失敗到門檻、而且連確定完好的檔案都讀不到：是硬碟卡死，不是那些檔案壞了。"""


recent = []                   # 上一次成功讀到資料之後新增的紀錄（硬碟卡死時整批撤銷）
streak = {"dir": 0, "defer": 0}
watchdog = {"fired_for": 0.0}


def note_ok():
    """真的從硬碟讀到資料了（複製成功或探測成功）：之前的失敗紀錄確定是真的壞，不再列入撤銷範圍。"""
    recent.clear()
    streak["dir"] = streak["defer"] = 0


def relaunch_args(argv, n):
    """自動重啟時要用的參數：原本的參數 + --relaunched n（拿掉舊的 --relaunched / --relaunch-after）。"""
    out, skip = [], False
    for a in argv:
        if skip:
            skip = False
        elif a in ("--relaunched", "--relaunch-after"):
            skip = True
        else:
            out.append(a)
    return out + ["--relaunched", str(n)]


def relaunch_self():
    """交給一個獨立的小幫手（rescue.py --relaunch-after 本PID）：強制結束本程式、等它真的消失、用同樣的參數開新視窗重跑。
    要用小幫手是因為本程式的主執行緒卡在核心的讀取裡，自己結束不了。"""
    cmd = [sys.executable, os.path.abspath(__file__), "--relaunch-after", str(os.getpid())] + relaunch_args(sys.argv[1:], S["relaunched"] + 1)
    subprocess.Popen(cmd, creationflags=subprocess.DETACHED_PROCESS | subprocess.CREATE_NEW_PROCESS_GROUP,
                     stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, close_fds=True)


def stall_watchdog():
    """給回報執行緒用：第一階段單一檔案太久完全沒進度 → 記成「卡住(壞軌)」寫進 CSV → 強制結束本程式 → 小幫手自動重啟，新程式會跳過它。
    自動重啟後如果一個檔案都還沒複製成功就又卡住，多半是硬碟卡死而不是這個檔壞了：不再強制跳過，等使用者拔電源。"""
    if S["phase"] != "第一階段" or S["state"] != "執行中" or not S["cur_start"] or S["cur_start"] == watchdog["fired_for"]:
        return False
    stalled = time.time() - max(S["last_progress"], S["cur_start"])
    if stalled < FILE_STALL_KILL_SEC:
        return False
    watchdog["fired_for"] = S["cur_start"]
    src = S["cur"]
    if S["relaunched"] and S["copied"] == 0:
        say(f"！{os.path.basename(src)} 已經 {int(stalled // 60)} 分鐘完全沒有進度，而且自動重啟後還沒複製成功任何檔案：疑似硬碟卡死，"
            f"不再自動強制跳過。請拔掉 E 槽電源休息後再重新啟動（這個檔會被記成卡住(壞軌)跳過）")
        return False
    record("檔案", src, OSError(1117, "讀取卡住，強制結束程式"), ekind="卡住(壞軌)",
           note=f"卡住 {int(stalled // 60)} 分鐘沒有進度，強制結束程式後跳過")
    say(f"！{os.path.basename(src)} 已經 {int(stalled // 60)} 分鐘完全沒有進度：已記成卡住(壞軌)，強制結束程式並自動重新啟動（跳過這個檔）")
    emit()
    relaunch_self()
    return True


def relaunch_helper(pid, argv):
    """小幫手：結束 pid、等它真的消失（卡在核心的讀取要等 Windows 放棄）、再用 argv 開新視窗重跑。期間每分鐘更新進度回報.txt。"""
    subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)
    t0 = time.time()
    log(f"[{now()[11:]}] 小幫手：已送出強制結束 PID {pid}，等它結束後用同樣參數重新啟動")
    while pid_alive(pid):
        time.sleep(5)
        if int(time.time() - t0) % 60 < 5:
            try:
                with open(REPORT_TXT, "w", encoding="utf-8-sig") as f:
                    f.write(f"【照片搬移回報】{now()}   第一階段\n狀態：卡住的檔案已記成壞軌，程式強制結束中：卡在核心裡的那次讀取要等 Windows 放棄才會真的結束"
                            f"（已等 {int((time.time() - t0) // 60)} 分鐘）。結束後會自動重新啟動並跳過那個檔，不用做任何事\n")
            except OSError:
                pass
    log(f"[{now()[11:]}] 小幫手：PID {pid} 已結束（等了 {int((time.time() - t0) // 60)} 分鐘），重新啟動：{' '.join(argv)}")
    subprocess.Popen([sys.executable, os.path.abspath(__file__)] + argv, creationflags=subprocess.CREATE_NEW_CONSOLE,
                     cwd=os.path.dirname(os.path.abspath(__file__)), close_fds=True)


def check_hung(reason):
    if os.path.splitdrive(PROBE_FILE)[0].lower() != os.path.splitdrive(SRC_ROOT)[0].lower():
        return                                   # 探測檔不在來源那顆碟上（測試環境）：不探測
    if drive_alive():
        note_ok()
        return
    with RL:
        for k in recent:
            records.pop(k, None)
    n = len(recent)
    recent.clear()
    streak["dir"] = streak["defer"] = 0
    save_records()
    raise HungDrive(f"{reason}，而且連確定完好的檔案都讀不到；這段期間的 {n} 筆紀錄已撤銷")


# ---------- 紀錄 ----------
def load_records():
    records.clear()
    p = os.path.join(WORK, "壞軌紀錄.csv")
    if os.path.exists(p):
        with open(p, encoding="utf-8-sig", newline="") as f:
            for r in csv.DictReader(f):
                if r["錯誤類型"] == "其他" and code_of(r) in KINDS:    # 舊版歸成「其他」的，依新分類更正
                    r["錯誤類型"] = KINDS[code_of(r)]
                records[r["來源完整路徑"]] = r


def save_records():
    p = os.path.join(WORK, "壞軌紀錄.csv")
    with RL:
        rows = list(records.values())
        for i, r in enumerate(rows, 1):
            r["項次"] = i
        try:
            with open(p + ".tmp", "w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, FIELDS, extrasaction="ignore")
                w.writeheader()
                w.writerows(rows)
            os.replace(p + ".tmp", p)
        except OSError as e:
            print(f"（壞軌紀錄.csv 寫入失敗，下次再寫：{e}）")
        dirty[0] = True


def is_pending(src):
    with RL:
        r = records.get(src)
        return bool(r) and r["狀態"] == PENDING


def record(kind, src, e, size=None, mtime=None, ekind=None, note=""):
    c = err_code(e)
    ekind = ekind or KINDS.get(c, "其他")
    with RL:
        records[src] = {
            "類型": kind, "第一層資料夾": top_of(src), "來源完整路徑": src,
            "檔名": os.path.basename(src) if kind == "檔案" else "",
            "大小(MB)": f"{size / 1048576:.2f}" if size is not None else "",
            "來源修改時間": datetime.datetime.fromtimestamp(mtime).strftime("%Y/%m/%d %H:%M:%S") if mtime else "",
            "錯誤代碼": f"{c} (0x{c:X})", "錯誤類型": ekind,
            "錯誤訊息": (str(e.strerror or e) + (f"（{note}）" if note else ""))[:200],
            "第一階段時間": now(), "狀態": PENDING,
            "救回比例%": "", "目的地檔案": "", "第二階段完成時間": ""}
        S["bad"] += 1
        S["last_bad"] = f"{src}  錯誤 {c} (0x{c:X}) {ekind}{f'（{note}）' if note else ''}  {now()[11:]}"
        recent.append(src)
    save_records()
    say(f"[壞軌] {kind}：{S['last_bad']}")


# ---------- 複製（第一階段） ----------
def is_done(dst, size, mtime):
    try:
        st = os.stat(lp(dst))
    except OSError:
        return False
    return st.st_size == size and abs(st.st_mtime - mtime) <= 2


class DstError(OSError):
    """寫入 G 槽（目的地）失敗。跟來源硬碟的讀取錯誤分開，才不會把 Google Drive 的問題記成壞軌。"""


def _dst(e):
    return DstError(e.errno, e.strerror, e.filename, None, e.filename2)


def _rm_part(dst):
    try:
        os.remove(lp(dst + ".part"))
    except OSError:
        pass


def copy_plain(src, dst):
    """一段一段複製，每讀到一段就更新 last_progress，用來判斷是真的卡住還是大檔案還在讀。
    來源讀取錯誤原樣丟出（OSError）；目的地寫入錯誤包成 DstError。"""
    tmp = dst + ".part"
    try:
        with open(lp(src), "rb", buffering=0) as fi:
            try:
                fo = open(lp(tmp), "wb")
            except OSError as e:
                raise _dst(e) from e
            with fo:
                while True:
                    b = fi.read(COPY_CHUNK)
                    if not b:
                        break
                    try:
                        fo.write(b)
                    except OSError as e:
                        raise _dst(e) from e
                    S["last_progress"] = time.time()
        try:
            shutil.copystat(lp(src), lp(tmp))       # 修改時間跟來源一樣，續傳判斷才會對
            os.replace(lp(tmp), lp(dst))
        except OSError as e:
            raise _dst(e) from e
    except BaseException:
        _rm_part(dst)
        raise


def copy_file(src, dst, size, mtime):
    S["cur"] = src
    if is_pending(src):              # 已知壞檔，留給第二階段，不重讀
        S["known"] += 1
        return _rm_part(dst)         # 上次卡住被強制結束時可能留下 .part，順手清掉
    if is_done(dst, size, mtime):
        S["skipped"] += 1
        return
    if (top_of(src), os.path.basename(src)) in KNOWN_BAD:     # 舊清單的壞檔：不讀，直接留給第二階段
        S["known"] += 1
        _rm_part(dst)
        return record("檔案", src, OSError(23, "舊清單已記錄的壞檔"), size, mtime,
                      note="第一階段不讀取，避免卡住硬碟")
    if S["copied"] % 200 == 0:
        ensure_space()
    S["cur_start"] = S["last_progress"] = time.time()
    dst_fails = 0
    try:
        while True:
            try:
                copy_plain(src, dst)
                S["copied"] += 1
                S["copied_bytes"] += size
                S["consec"] = 0
                note_ok()
                return
            except DstError as e:                    # G 槽的問題：不是壞軌，等一下再試
                if not online():
                    wait_online()
                    continue
                if err_code(e) in (112, 28):         # 磁碟已滿（winerror 112 / errno 28）：等 Google Drive 上傳釋出空間
                    pause("G 槽空間不足，等 Google Drive 上傳釋出暫存空間後自動繼續", DST_RETRY_SLEEP)
                    continue
                dst_fails += 1
                if dst_fails <= DST_RETRY_LIMIT:
                    pause(f"寫入 G 槽失敗（{e.strerror}），{DST_RETRY_SLEEP} 秒後重試（第 {dst_fails}/{DST_RETRY_LIMIT} 次）",
                          DST_RETRY_SLEEP)
                    continue
                return record("檔案", src, e, size, mtime, ekind="目的地寫入失敗",
                              note="G 槽寫不進去，不是壞軌；第二階段會再試")
            except OSError as e:
                stalled = int(time.time() - S["last_progress"])
                if not online():
                    wait_online()
                    if stalled < STALL_SEC:          # 一般斷線：接回後同一個檔案重來
                        continue
                    return record("檔案", src, e, size, mtime, ekind="卡住(壞軌)",   # 卡住才斷線：不重讀
                                  note=f"讀取卡住 {stalled} 秒沒有進度後斷線")
                record("檔案", src, e, size, mtime, note=f"讀了 {int(time.time() - S['cur_start'])} 秒才回報錯誤")
                S["consec"] += 1
                if S["consec"] >= CONSEC_FAIL_LIMIT:   # 第一、二階段都適用：第二階段也先搬好的，壞區留到最後
                    S["consec"] = 0
                    raise DeferDir(e)            # 交給 walk：這個資料夾剩下的留給第二階段
                return
    finally:
        S["cur_start"] = 0.0


class DeferDir(Exception):
    """第一階段同一個資料夾連續失敗太多張：多半是一整片實體壞區，剩下的照片整批延後到第二階段。"""


def do_file(p, ent):
    S["idx_in_top"] += 1
    try:
        st = ent.stat(follow_symlinks=False)
    except OSError as e:
        return record("檔案", p, e)
    copy_file(p, dst_of(p), st.st_size, st.st_mtime)


def walk(src_dir):
    """回傳 False 代表這個資料夾本身列不出來（已記錄）。"""
    while True:
        try:
            with os.scandir(lp(src_dir)) as it:
                ents = sorted(it, key=lambda e: e.name.lower())
            break
        except OSError as e:
            if online():
                record("資料夾無法讀取", src_dir, e)
                streak["dir"] += 1
                if streak["dir"] >= HUNG_DIR_LIMIT:
                    check_hung(f"連續 {HUNG_DIR_LIMIT} 個資料夾都讀不到清單")
                return False
            wait_online()
    streak["dir"] = 0
    mkdir(dst_of(src_dir))
    S["consec"] = 0                                  # 連續失敗只算同一個資料夾裡的
    for ent in ents:
        if ent.name.lower() in SKIP_NAMES:
            continue
        try:
            if ent.stat(follow_symlinks=False).st_file_attributes & 0x400:   # 重新分析點（junction 等）：不跟進去
                continue
        except OSError:
            pass
        p = os.path.join(src_dir, ent.name)
        if ent.is_dir(follow_symlinks=False):
            if is_pending(p):
                S["known"] += 1
            else:
                walk(p)
                S["consec"] = 0
        else:
            try:
                do_file(p, ent)
            except DeferDir as d:
                record("資料夾（壞軌密集，延後）", src_dir, d.args[0], ekind="壞軌",
                       note=f"{S['phase']}連續 {CONSEC_FAIL_LIMIT} 張失敗，這個資料夾剩下的照片延後處理")
                streak["defer"] += 1
                if streak["defer"] >= HUNG_DEFER_LIMIT:
                    check_hung(f"連續 {HUNG_DEFER_LIMIT} 個資料夾都連續失敗、中間沒有任何一張成功")
                return True
    return True


def completed_tops():
    """從進度紀錄.log 找出第一階段已經完整走過的第一層資料夾。走過就代表裡面每個檔案不是已經搬好、就是已經記成待救援，
    重新啟動時可以直接略過、完全不讀 E 槽（硬碟退化後，連讀資料夾清單都可能卡好幾分鐘）。刪掉 log 就會全部重掃。"""
    done = set()
    root = r"E:\照片"                   # 舊紀錄沒有「來源：」行，那時的來源都是 E:\照片
    try:
        with open(os.path.join(WORK, "進度紀錄.log"), encoding="utf-8") as f:
            for line in f:
                if "] 來源：" in line:
                    root = line.split("] 來源：", 1)[1].strip()
                elif ("] [資料夾完成 " in line and "：新複製 " in line and "整個資料夾無法讀取" not in line
                      and root.rstrip("\\").lower() == SRC_ROOT.rstrip("\\").lower()):
                    done.add(line.split("] ", 2)[2].rsplit("：新複製 ", 1)[0])
    except OSError:
        pass
    return done


def phase1(only=None, tops_filter=None):
    say(f"來源：{SRC_ROOT}")             # completed_tops() 靠這一行分辨不同來源的紀錄
    say(f"目的地：{DST_ROOT}" + (f"，只搬：{tops_filter}" if tops_filter else ""))
    fails = 0
    while True:                         # 讀 E:\照片 清單失敗不要直接當掉（9/14 當掉過兩次），重試幾次
        try:
            with os.scandir(lp(SRC_ROOT)) as it:
                ents = sorted(it, key=lambda e: e.name.lower())
            break
        except OSError as e:
            if not online():
                wait_online()
                continue
            fails += 1
            if fails >= 5:
                raise
            pause(f"讀不到 {SRC_ROOT} 的資料夾清單（錯誤 {err_code(e)}），1 分鐘後重試（第 {fails}/5 次）", 60)
    ents = [e for e in ents if e.name.lower() not in SKIP_NAMES]
    tops = [e.name for e in ents if e.is_dir()]
    loose = [e for e in ents if not e.is_dir()]
    if only:
        if only not in tops:
            raise SystemExit(f"找不到資料夾：{SRC_ROOT}\\{only}")
        tops, loose = [only], []
    if tops_filter:                     # 只搬指定的幾個第一層資料夾，根目錄的散檔不搬
        missing = [t for t in tops_filter if t not in tops]
        if missing:
            say(f"注意：來源裡找不到這些資料夾，略過：{missing}")
        tops, loose = [t for t in tops if t in tops_filter], []
    S["top_n"] = len(tops)
    done_tops = set() if only else completed_tops()
    if done_tops & set(tops):
        say(f"之前已經完整處理過的 {len(done_tops & set(tops))} 個資料夾直接略過，不再讀取 E 槽")
    try:
        if loose:
            S.update(top="（根目錄的散檔）", idx_in_top=0)
            try:
                for e in loose:
                    do_file(os.path.join(SRC_ROOT, e.name), e)
            except DeferDir:
                pass
        for i, name in enumerate(tops, 1):
            if name in done_tops:
                continue
            p = os.path.join(SRC_ROOT, name)
            S.update(top=name, top_i=i, idx_in_top=0)
            if is_pending(p):
                S["known"] += 1
                say(f"[資料夾 {i}/{len(tops)}] {name}：已知整個資料夾無法讀取，留給第二階段")
                continue
            c0, k0, b0 = S["copied"], S["skipped"], S["bad"]
            ok = walk(p)
            say(f"[資料夾完成 {i}/{len(tops)}] {name}：新複製 {S['copied'] - c0}，先前已完成 {S['skipped'] - k0}，"
                f"壞檔 {S['bad'] - b0}" + ("" if ok else "（整個資料夾無法讀取）"))
    except HungDrive as h:
        S["state"] = f"硬碟疑似卡死（{h}），程式已自動停止。請拔掉 E 槽電源休息後，再重新啟動"
        say(S["state"])
        return False
    return True


# ---------- 救援（第二階段） ----------
class Reader:
    def __init__(self, path):
        self.path, self.f = path, None

    def open_file(self):
        for _ in range(3):
            try:
                self.f = open(lp(self.path), "rb", buffering=0)
                return True
            except OSError:
                if online():
                    time.sleep(P2_RETRY_SLEEP)
                else:
                    wait_online()
        return False

    def read_at(self, pos, n):
        while True:
            try:
                if self.f is None and not self.open_file():
                    raise OSError(1117, "重新開檔失敗")
                self.f.seek(pos)
                return self.f.read(n)
            except OSError:
                if online():
                    raise
                self.close()
                wait_online()

    def close(self):
        if self.f:
            try:
                self.f.close()
            except OSError:
                pass
            self.f = None


def read_block(read_at, pos, m):
    for k in range(1 + P2_RETRY):
        try:
            d = read_at(pos, m)
            if len(d) == m:
                return d
        except OSError:
            pass
        if k < P2_RETRY:
            time.sleep(P2_RETRY_SLEEP)
    return None


def rescue_stream(read_at, size, write):
    """逐段讀，讀不到的區塊補 0。回傳 (讀到的位元組數, 壞區塊數, 是否碰到時間上限)。"""
    t0, bad_time = time.time(), 0.0
    pos = good = badblk = 0
    capped = False
    while pos < size:
        if time.time() - t0 > P2_FILE_TIME_CAP_SEC:
            capped = True
            break
        n = min(P2_CHUNK, size - pos)
        try:
            data = read_at(pos, n)
        except OSError:
            data = None
        if data is not None and len(data) == n:
            write(data)
            good += n
            pos += n
            S["last_progress"] = time.time()
        elif bad_time > P2_BAD_TIME_CAP_SEC:       # 壞區時間用完：這段直接補 0
            capped = True
            write(bytes(n))
            badblk += -(-n // P2_BLOCK)
            pos += n
        else:                                       # 這 1 MiB 改成 4 KiB 慢慢讀
            tb, end = time.time(), pos + n
            while pos < end:
                if bad_time + time.time() - tb > P2_BAD_TIME_CAP_SEC:
                    capped = True
                    write(bytes(end - pos))
                    badblk += -(-(end - pos) // P2_BLOCK)
                    pos = end
                    break
                m = min(P2_BLOCK, end - pos)
                d = read_block(read_at, pos, m)
                if d is None:
                    write(bytes(m))
                    badblk += 1
                else:
                    write(d)
                    good += m
                    S["last_progress"] = time.time()
                pos += m
                S["p2_pos"], S["p2_badblk"] = pos * 100 / size, badblk
            bad_time += time.time() - tb
        S["p2_pos"], S["p2_badblk"] = pos * 100 / size, badblk
    while pos < size:                               # 總時間用完：剩下全部補 0
        n = min(P2_CHUNK, size - pos)
        write(bytes(n))
        pos += n
    return good, badblk, capped


def put(local, dst):
    while True:
        try:
            shutil.copy2(local, lp(dst + ".part"))
            os.replace(lp(dst + ".part"), lp(dst))
            os.remove(local)
            return
        except OSError:
            if online():
                raise
            wait_online()


def finish(r, status, ratio="", dst=""):
    with RL:
        r.update({"狀態": status, "救回比例%": ratio, "目的地檔案": dst, "第二階段完成時間": now()})
    key = {"已完整救回": "p2_full", "無法讀取": "p2_fail", "已可讀取（已重新搬移）": "p2_dirok",
           "資料夾結構損壞，需專業工具": "p2_dirfail"}.get(status, "p2_part")
    S[key] += 1
    S["cur_start"] = 0.0
    save_records()
    say(f"[第二階段] {status}{f' {ratio}%' if ratio != '' else ''}：{r['來源完整路徑']}")
    return status


def rescue_file(r):
    src = r["來源完整路徑"]
    dst = dst_of(src)
    S.update(cur=src, cur_start=time.time(), p2_pos=0.0, p2_badblk=0)
    st = None
    for _ in range(3):
        try:
            st = os.stat(lp(src))
            break
        except OSError:
            if online():
                time.sleep(P2_RETRY_SLEEP)
            else:
                wait_online()
    if st is None:
        return finish(r, "無法讀取")
    mkdir(os.path.dirname(dst))
    while True:                                     # 1) 先用普通方式再試一次（32、121 常常這樣就好了）
        try:
            copy_plain(src, dst)
            note_ok()
            return finish(r, "已完整救回", 100, dst)
        except OSError:
            if online():
                break
            wait_online()
    rd = Reader(src)                                # 2) 逐段讀，讀不到補 0
    if not rd.open_file():
        return finish(r, "無法讀取")
    tmp = os.path.join(WORK, "_rescue.tmp")         # 先寫在本機，避免 Google Drive 鎖住半成品
    try:
        with open(tmp, "wb") as out:
            good, badblk, capped = rescue_stream(rd.read_at, st.st_size, out.write)
    finally:
        rd.close()
    if good == 0 and st.st_size > 0:
        os.remove(tmp)
        return finish(r, "無法讀取")
    note_ok()
    os.utime(tmp, (st.st_atime, st.st_mtime))
    if good == st.st_size:
        put(tmp, dst)
        return finish(r, "已完整救回", 100, dst)
    pdst = partial_name(dst)                        # 不用原檔名，免得被當成好檔
    put(tmp, pdst)
    return finish(r, "部分救回（逾時）" if capped else "部分救回", round(good * 100 / st.st_size, 1), pdst)


def rescue_dir(r):
    src = r["來源完整路徑"]
    S.update(cur=src, cur_start=time.time(), p2_pos=0.0, p2_badblk=0, top=top_of(src), idx_in_top=0)
    ok = False
    for k in range(P2_DIR_RETRY):
        if k:
            time.sleep(P2_DIR_RETRY_SLEEP)
        try:
            with os.scandir(lp(src)) as it:
                list(it)
            ok = True
            break
        except OSError:
            if not online():
                wait_online()
    ok = ok and walk(src)
    with RL:
        cur = records.get(src)
        deferred = ok and cur is not None and cur is not r and cur["狀態"] == PENDING
        if not deferred:
            records[src] = r        # walk 失敗時會覆蓋這筆，放回原本那筆再更新狀態
    if deferred:                    # 這個資料夾又連續失敗好幾張：新的「延後」紀錄留著，下次第二階段從斷點接著處理
        S["cur_start"] = 0.0
        return PENDING
    status = "已可讀取（已重新搬移）" if ok else "資料夾結構損壞，需專業工具"
    finish(r, status)
    return status


def drive_alive(timeout=30):
    """不經快取讀 PROBE_FILE 開頭 4 KB，timeout 秒內讀得到就代表硬碟還有回應。"""
    k = ctypes.windll.kernel32
    k.CreateFileW.restype = ctypes.c_void_p
    k.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
                              ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
    k.ReadFile.argtypes = [ctypes.c_void_p, ctypes.c_void_p, ctypes.c_uint32, ctypes.POINTER(ctypes.c_uint32), ctypes.c_void_p]
    k.CloseHandle.argtypes = [ctypes.c_void_p]
    k.VirtualAlloc.restype = ctypes.c_void_p
    k.VirtualAlloc.argtypes = [ctypes.c_void_p, ctypes.c_size_t, ctypes.c_uint32, ctypes.c_uint32]
    k.VirtualFree.argtypes = [ctypes.c_void_p, ctypes.c_size_t, ctypes.c_uint32]
    result = []

    def probe():
        h = k.CreateFileW(lp(PROBE_FILE), 0x80000000, 1, None, 3, 0x20000000, None)   # 唯讀、FILE_FLAG_NO_BUFFERING
        if h in (None, ctypes.c_void_p(-1).value):
            return result.append(False)
        buf = k.VirtualAlloc(None, 4096, 0x3000, 0x04)          # 不經快取的讀取需要對齊 4 KB 的緩衝區
        n = ctypes.c_uint32(0)
        try:
            result.append(bool(k.ReadFile(h, buf, 4096, ctypes.byref(n), None)))
        finally:
            k.CloseHandle(h)
            k.VirtualFree(buf, 0, 0x8000)
    t = threading.Thread(target=probe, daemon=True)
    t.start()
    t.join(timeout)
    return bool(result and result[0])


def p2_order(r):
    """第二階段的處理順序：救回機會大、數量多的先做，趁硬碟還活著先搬。
    1. 資料夾：清單讀得到、只是裡面有壞檔而被跳過的（錯誤 13/22/23，例如幾個出國旅遊資料夾，裡面有幾千張好照片）
    2. 資料夾：清單讀取逾時的（121）；9/13 就讀不到的老問題資料夾排在同類最後
    3. 資料夾：結構損壞（1392/1117）
    4. 最後才一張一張慢慢救個別壞檔（每張可能要 20 分鐘以上），暫時性錯誤先、小檔先"""
    c = code_of(r)
    if r["類型"] != "檔案":
        return (0, 0 if c in (13, 22, 23) else 1 if c == 121 else 2, r["第一階段時間"].startswith("2026/09/13"), 0, 0)
    return (1, 0, False, c not in TRANSIENT, float(r["大小(MB)"] or 0))


def phase2():
    tried = set()
    while True:
        with RL:
            pend = [r for r in records.values() if r["狀態"] == PENDING and r["來源完整路徑"] not in tried]
        if not pend:
            return None
        pend.sort(key=p2_order)
        S["p2_n"] = len(tried) + len(pend)
        say(f"第二階段：待救援 {len(pend)} 項（資料夾 {sum(r['類型'] != '檔案' for r in pend)} 個、"
            f"檔案 {sum(r['類型'] == '檔案' for r in pend)} 個）")
        for r in pend:
            tried.add(r["來源完整路徑"])
            S["p2_i"] = len(tried)
            try:
                status = rescue_file(r) if r["類型"] == "檔案" else rescue_dir(r)
            except HungDrive as h:
                with RL:                # 走到一半硬碟卡死：這一項放回待救援
                    records[r["來源完整路徑"]] = r
                    r.update({"狀態": PENDING, "救回比例%": "", "目的地檔案": "", "第二階段完成時間": ""})
                save_records()
                S["state"] = f"硬碟疑似卡死（{h}），程式已自動停止，剛才那一項已還原成待救援。請拔掉 E 槽電源休息後，再重新啟動"
                say(S["state"])
                return False
            if status in ("無法讀取", "資料夾結構損壞，需專業工具") and not drive_alive():
                with RL:                # 是硬碟卡死、不是這一項真的壞了：還原成待救援
                    r.update({"狀態": PENDING, "救回比例%": "", "目的地檔案": "", "第二階段完成時間": ""})
                save_records()
                S["state"] = "硬碟疑似卡死（連確定完好的檔案都讀不到），程式已自動停止，剛才那一項已還原成待救援。請拔掉 E 槽電源休息後，再重新啟動"
                say(S["state"])
                return False
            time.sleep(P2_REST_BETWEEN_FILES)


# ---------- 回報 ----------
SHEET = "無法複製清單"
BAD_KINDS = {"壞軌", "結構損壞", "逾時(可能壞軌)", "卡住(壞軌)"}
FAILED = {"無法讀取", "部分救回", "部分救回（逾時）", "資料夾結構損壞，需專業工具"}


def in_excel(r):
    """xlsx 只放真的失敗、或因壞軌無法複製的項目；暫時性錯誤和已救回的只留在 CSV。"""
    return r["狀態"] in FAILED or (r["狀態"] == PENDING and r["錯誤類型"] in BAD_KINDS)


def build_xlsx():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    with RL:
        rows = [dict(r) for r in records.values() if in_excel(r)]
        dirty[0] = False
    try:
        if os.path.exists(XLSX):
            wb = load_workbook(XLSX)
            if SHEET in wb.sheetnames:
                del wb[SHEET]
        else:
            wb = Workbook()
            wb.remove(wb.active)
        ws = wb.create_sheet(SHEET, 0)
        ws.append(FIELDS)
        for i, r in enumerate(rows, 1):
            vals = [i] + [r.get(f, "") for f in FIELDS[1:]]
            for j in (5, 12):                       # 大小(MB)、救回比例% 存成數字
                try:
                    vals[j] = float(vals[j])
                except (TypeError, ValueError):
                    pass
            ws.append(vals)
            fill = PatternFill("solid", fgColor=FILLS.get(r["狀態"], "FCE4D6"))
            for c in ws[ws.max_row]:
                c.fill = fill
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, w in zip("ABCDEFGHIJKLMNO", (6, 14, 28, 60, 26, 9, 19, 14, 16, 30, 19, 24, 10, 60, 19)):
            ws.column_dimensions[col].width = w
        for other in wb.worksheets:
            other.sheet_view.tabSelected = False
        ws.sheet_view.tabSelected = True
        wb.active = 0
        wb.save(XLSX)
    except BaseException:
        dirty[0] = True
        raise


def report():
    t = time.time()
    s = dict(S)
    st = s["state"]
    stall = t - max(s["last_progress"], s["cur_start"]) if s["cur_start"] else 0   # 這個檔案開始前的進度不算
    if st == "執行中" and stall > STALL_SEC:
        st = (f"卡在壞軌（已經 {int(stall // 60)} 分 {int(stall % 60)} 秒沒有讀到任何資料，硬碟正在重試）。"
              f"可以等 Windows 放棄，或把 E 槽拔掉重插：這張會直接記成壞軌跳過")
    if s["offline_since"] and t - s["offline_since"] > 1800:
        st += "\n      ★★★ 已斷線超過 30 分鐘，請檢查 E 槽的 USB 線和電源 ★★★"
    secs = f"，已讀 {int(t - s['cur_start'])} 秒" if s["cur_start"] else ""
    L = [f"【照片搬移回報】{now()}   {s['phase']}", f"狀態：{st}"]
    if s["phase"] == "第一階段":
        L += [f"目前資料夾：{s['top']}（第 {s['top_i']} / {s['top_n']} 個資料夾）",
              f"目前檔案：{s['cur'] or '—'}（這個資料夾的第 {s['idx_in_top']} 個檔{secs}）"]
    else:
        L += [f"目前救援：第 {s['p2_i']} / {s['p2_n']} 項  {s['cur'] or '—'}",
              f"          已讀 {s['p2_pos']:.0f}%，這個檔的壞區塊 {s['p2_badblk']} 個{secs}",
              f"第二階段結果：完整救回 {s['p2_full']}，部分救回 {s['p2_part']}，無法讀取 {s['p2_fail']}，"
              f"資料夾恢復可讀 {s['p2_dirok']}，資料夾仍無法讀取 {s['p2_dirfail']}"]
    dt = max(t - PREV["t"], 1)
    dc, db = s["copied"] - PREV["copied"], s["copied_bytes"] - PREV["bytes"]
    L += [f"本次累計：新複製 {s['copied']:,} 檔 / {s['copied_bytes'] / 1e9:.2f} GB，先前已完成跳過 {s['skipped']:,} 檔，"
          f"新壞檔 {s['bad']} 個，已知待救援略過 {s['known']} 個"
          + (f"（卡住強制跳過後自動重啟的第 {s['relaunched']} 次接續）" if s["relaunched"] else ""),
          f"過去 {dt / 60:.0f} 分鐘：新複製 {dc:,} 檔 / {db / 1e6:.0f} MB（{db / 1e6 / dt:.1f} MB/s），"
          f"新壞檔 {s['bad'] - PREV['bad']} 個",
          f"最近一次壞軌：{s['last_bad'] or '無'}"]
    try:
        L.append(f"G 槽剩餘空間：{free_gb():.0f} GB")
    except OSError:
        L.append("G 槽剩餘空間：讀不到")
    L += [f"備註：{n}" for n in notes]
    PREV.update(t=t, copied=s["copied"], bytes=s["copied_bytes"], bad=s["bad"])
    return "\n".join(L)


def emit():
    notes.clear()
    if dirty[0]:
        try:
            build_xlsx()
        except PermissionError:
            notes.append("xlsx 正被開啟（Excel？），這次沒更新，下次再寫")
        except Exception as e:
            notes.append(f"xlsx 更新失敗：{e}")
    text = report()
    print("\n" + "=" * 70 + "\n" + text + "\n" + "=" * 70 + "\n", flush=True)
    log(text)
    try:
        with open(REPORT_TXT, "w", encoding="utf-8-sig") as f:     # 加 BOM：PowerShell 5.1 / 記事本才不會亂碼
            f.write(text + "\n")
    except OSError as e:
        print(f"（進度回報.txt 寫入失敗，下次再寫：{e}）")


def set_title():
    t = time.time()
    secs = f" {int(t - S['cur_start'])}秒" if S["cur_start"] else ""
    where = f"{S['top_i']}/{S['top_n']} {S['top']}" if S["phase"] == "第一階段" else f"{S['p2_i']}/{S['p2_n']}"
    try:
        ctypes.windll.kernel32.SetConsoleTitleW(
            f"[{S['phase']}] {S['state'][:14]} | {where} | {os.path.basename(S['cur'])}{secs} | "
            f"新複製{S['copied']} 壞{S['bad']}")
    except Exception:
        pass


def reporter(stop):
    last = 0.0
    while not stop.is_set():
        if time.time() - last >= REPORT_EVERY_SEC:
            last = time.time()
            try:
                emit()
            except Exception as e:
                print(f"（回報失敗：{e}）")
        try:
            stall_watchdog()
        except Exception as e:
            print(f"（卡住偵測失敗：{e}）")
        set_title()
        stop.wait(2)


# ---------- 主程式 ----------
def pid_alive(pid):
    out = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"],
                         capture_output=True, text=True, errors="ignore").stdout
    return f'"{pid}"' in out and "python" in out.lower()


def proc_start_time(pid):
    """程式啟動時間（epoch 秒），讀不到回傳 None。"""
    k = ctypes.windll.kernel32
    h = k.OpenProcess(0x1000, False, pid)            # PROCESS_QUERY_LIMITED_INFORMATION
    if not h:
        return None
    try:
        ft = (ctypes.c_ulonglong * 4)()
        if not k.GetProcessTimes(h, ctypes.byref(ft, 0), ctypes.byref(ft, 8), ctypes.byref(ft, 16), ctypes.byref(ft, 24)):
            return None
        return ft[0] / 1e7 - 11644473600
    finally:
        k.CloseHandle(h)


def lock_pid():
    try:
        with open(os.path.join(WORK, "rescue.lock")) as f:
            return int(f.read().strip())
    except (OSError, ValueError):
        return 0


def running_phase(pid):
    """從進度紀錄.log 找這個 PID 跑的是哪個階段（「===== 第X階段開始（PID n）=====」）。"""
    phase = ""
    try:
        with open(os.path.join(WORK, "進度紀錄.log"), encoding="utf-8") as f:
            for line in f:
                if f"開始（PID {pid}）" in line:
                    phase = "第二階段" if "第二階段" in line else "第一階段"
    except OSError:
        pass
    return phase


def restart(mode="phase1", max_wait=300):
    """給執行 AI 用：結束「跑舊版程式碼」的搬移程式，等它真的結束後，用檔案總管啟動新版（等於使用者雙擊）。
    可以重複執行：新版已經在跑就什麼都不做。mode=phase2 時絕不會結束正在跑的程式，只在沒有程式執行時啟動第二階段。"""
    code_mtime = os.path.getmtime(os.path.abspath(__file__))
    old = lock_pid()
    if old and pid_alive(old) and mode == "phase2" and running_phase(old) != "第二階段":
        return print(f"目前還有搬移程式在執行（PID {old}），第二階段要等它結束才能啟動。沒有啟動任何東西。")
    if old and pid_alive(old):
        st = proc_start_time(old)
        if st and st > code_mtime:
            return print(f"新版已經在執行中（PID {old}），不需要重啟。")
        subprocess.run(["taskkill", "/F", "/PID", str(old)], capture_output=True)
        t0 = time.time()
        while pid_alive(old) and time.time() - t0 < max_wait:
            time.sleep(5)
        if pid_alive(old):
            return print(f"舊版（PID {old}）正在結束中：它卡在讀取壞檔，要等 Windows 放棄那次讀取才會真正結束"
                         f"（可能要 30 分鐘）。請直接再執行一次 --restart。")
        print(f"舊版（PID {old}）已結束。")
    bat = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       {"phase1": "開始第一階段.bat", "phase2": "開始第二階段.bat", "others": "開始其他資料夾.bat"}[mode])
    subprocess.Popen(["explorer.exe", bat])
    time.sleep(30)
    new = lock_pid()
    if new and new != old and pid_alive(new):
        print(f"新版已啟動（PID {new}）。")
    else:
        print("啟動失敗：30 秒內沒有看到新的程式在執行。請照 docs/方法一_Windows逐檔複製.md「啟動」一節改用排程工作啟動。")


def acquire_lock():
    p = os.path.join(WORK, "rescue.lock")
    try:
        with open(p) as f:
            pid = int(f.read().strip())
    except (OSError, ValueError):
        pid = 0
    if pid and pid_alive(pid):
        raise SystemExit(f"已經有一份 rescue.py 在執行（PID {pid}），同時只能跑一份。本次結束。")
    with open(p, "w") as f:
        f.write(str(os.getpid()))
    return p


def wait_report(max_wait):
    """給執行 AI 用：等到下一次 5 分鐘回報產生（最多等 max_wait 秒）就印出來。只讀回報檔，不碰 E 槽。"""
    def mtime():
        try:
            return os.path.getmtime(REPORT_TXT)
        except OSError:
            return 0.0
    t0, m0 = time.time(), mtime()
    while mtime() == m0 and time.time() - t0 < max_wait:
        time.sleep(2)
    try:
        with open(REPORT_TXT, encoding="utf-8-sig") as f:
            text = f.read().rstrip()
    except OSError:
        text = "（還沒有進度回報.txt）"
    age = time.time() - mtime()
    running = os.path.exists(os.path.join(WORK, "rescue.lock"))
    if running and age > REPORT_EVERY_SEC * 2:
        text = f"！！回報已經 {int(age // 60)} 分鐘沒有更新，程式可能當住或視窗被關閉，請提醒使用者！！\n" + text
    print(text + f"\n（程式是否執行中：{'是' if running else '否（rescue.lock 不存在，程式已結束或沒有啟動）'}）", flush=True)


def main():
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="壞軌外接硬碟照片搬移")
    ap.add_argument("mode", nargs="?", choices=["phase1", "phase2"])
    ap.add_argument("--only", help=r"只跑來源底下這一個資料夾（試跑用）")
    ap.add_argument("--src", help=rf"來源根目錄（預設 {SRC_ROOT}）")
    ap.add_argument("--dst", help=rf"目的地根目錄（預設 {DST_ROOT}）")
    ap.add_argument("--tops", help=f"只搬來源底下這幾個第一層資料夾，用逗號分隔；或用預設組名 {list(PRESETS)}")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--wait-report", nargs="?", type=int, const=REPORT_EVERY_SEC + 60, metavar="秒",
                    help="等下一次回報產生後印出來（預設最多等 360 秒），給執行 AI 在對話中回報用")
    ap.add_argument("--restart", nargs="?", const="phase1", choices=["phase1", "phase2", "others"],
                    help="結束舊版、啟動新版（phase1=照片第一階段；others=其他資料夾；phase2 只在沒有程式執行時啟動第二階段）。"
                         "給執行 AI 用，可重複執行")
    ap.add_argument("--relaunched", type=int, default=0, help=argparse.SUPPRESS)        # 程式自己用：第幾次強制跳過後的接續
    ap.add_argument("--relaunch-after", type=int, metavar="PID", help=argparse.SUPPRESS)  # 程式自己用：小幫手模式
    a = ap.parse_args()
    if a.relaunch_after:
        return relaunch_helper(a.relaunch_after, relaunch_args(sys.argv[1:], a.relaunched))
    S["relaunched"] = a.relaunched
    if a.tops == "others" and not (a.src or a.dst):     # 開始其他資料夾.bat 只給 --tops others，來源與目的地用下面的預設
        a.src, a.dst = OTHERS_SRC, OTHERS_DST
    if a.src or a.dst:
        gl = globals()
        src = os.path.abspath(a.src) if a.src else SRC_ROOT             # "E:\" 保留反斜線，其他去掉尾巴
        gl["SRC_ROOT"] = src if src.endswith(":\\") else src.rstrip("\\")
        gl["DST_ROOT"] = os.path.abspath(a.dst).rstrip("\\") if a.dst else DST_ROOT
    tops_filter = (PRESETS.get(a.tops) or [t.strip() for t in a.tops.split(",") if t.strip()]) if a.tops else None
    if a.selftest:
        return selftest()
    if a.wait_report is not None:
        return wait_report(a.wait_report)
    if a.restart:
        return restart(a.restart)
    if not a.mode:
        return ap.print_help()
    os.makedirs(WORK, exist_ok=True)
    lock = acquire_lock()
    try:    # 執行期間不讓電腦睡眠（螢幕可以關）；程式結束後自動解除
        ctypes.windll.kernel32.SetThreadExecutionState(0x80000001)   # ES_CONTINUOUS | ES_SYSTEM_REQUIRED
    except Exception:
        pass
    load_records()
    bak = os.path.join(WORK, "原始清單備份.xlsx")
    if os.path.exists(XLSX) and not os.path.exists(bak):
        shutil.copy2(XLSX, bak)
    S["phase"] = "第一階段" if a.mode == "phase1" else "第二階段"
    S["state"] = "執行中"
    say(f"===== {S['phase']}開始（PID {os.getpid()}）=====")
    stop = threading.Event()
    th = threading.Thread(target=reporter, args=(stop,), daemon=True)
    th.start()
    try:
        if not online():
            wait_online()
        mkdir(DST_ROOT)
        finished = phase1(a.only, tops_filter) if a.mode == "phase1" else phase2() is None
        if finished:
            S["state"] = "已完成"
    except KeyboardInterrupt:
        S["state"] = "已手動中斷（重新執行同一個指令會從中斷處接續）"
    except BaseException:
        S["state"] = "程式發生錯誤而停止（請把畫面上的錯誤訊息交給 AI）"
        raise
    finally:
        S["cur_start"] = 0.0
        stop.set()
        th.join(10)
        save_records()
        emit()
        try:
            os.remove(lock)
        except OSError:
            pass
    if S["state"] != "已完成":
        return
    if a.mode == "phase1":
        say("===== 第一階段完成。請把結果回報使用者，等使用者同意後才可以執行 phase2 =====")
    else:
        say("===== 第二階段完成 =====")


# ---------- 自我測試（不碰 E 槽） ----------
def selftest():
    import tempfile
    from openpyxl import Workbook, load_workbook
    g = globals()
    base = tempfile.mkdtemp(prefix="rescue_selftest_")
    g.update(SRC_ROOT=os.path.join(base, "src"), DST_ROOT=os.path.join(base, "dst"), WORK=os.path.join(base, "work"),
             XLSX=os.path.join(base, "清單.xlsx"), REPORT_TXT=os.path.join(base, "回報.txt"),
             P2_RETRY_SLEEP=0, OFFLINE_POLL_SEC=0, COPY_CHUNK=4096, P2_REST_BETWEEN_FILES=0, P2_DIR_RETRY_SLEEP=0,
             DST_RETRY_SLEEP=0)
    for d in (SRC_ROOT, DST_ROOT, WORK):
        os.makedirs(d)
    real_online, real_wait, real_read, real_plain = g["online"], g["wait_online"], Reader.read_at, g["copy_plain"]
    real_alive, real_relaunch = g["drive_alive"], g["relaunch_self"]
    relaunches = []
    g["relaunch_self"] = lambda: relaunches.append(1)   # 測試裡絕對不能真的結束自己

    def mk(rel, data):
        p = os.path.join(SRC_ROOT, rel)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        with open(p, "wb") as f:
            f.write(data)
        return p

    try:
        # 1. 已完成判斷：大小+時間一樣就跳過；robocopy 失敗殘檔（1980 年）要重抄
        a = mk(r"A\a.jpg", os.urandom(10000))       # 10000 bytes ÷ 4096 = 分 3 段複製
        sa = os.stat(a)
        mkdir(os.path.dirname(dst_of(a)))
        copy_plain(a, dst_of(a))
        assert is_done(dst_of(a), sa.st_size, sa.st_mtime) and not os.path.exists(dst_of(a) + ".part")
        with open(a, "rb") as f1, open(dst_of(a), "rb") as f2:
            assert f1.read() == f2.read()
        os.utime(dst_of(a), (315532800, 315532800))
        assert not is_done(dst_of(a), sa.st_size, sa.st_mtime)

        # 2. 逐段救援：第二個 1 MiB 裡的第 4 個 4 KiB 區塊壞掉
        data = os.urandom(3 * P2_CHUNK + 5000)
        bad = P2_CHUNK + 3 * P2_BLOCK

        def ra(pos, n):
            if pos <= bad < pos + n:
                raise OSError(23, "CRC")
            return data[pos:pos + n]
        buf = bytearray()
        good, bb, capped = rescue_stream(ra, len(data), buf.extend)
        assert (len(buf), good, bb, capped) == (len(data), len(data) - P2_BLOCK, 1, False)
        assert buf[bad:bad + P2_BLOCK] == bytes(P2_BLOCK)
        assert buf[:bad] == data[:bad] and buf[bad + P2_BLOCK:] == data[bad + P2_BLOCK:]

        # 2b. 壞檔：第一階段記錄 → 重跑不重讀 → 第二階段存成 _部分損壞
        b = mk(r"A\b.jpg", data)
        sb = os.stat(b)

        def bad_plain(s, d):
            if s.endswith("b.jpg"):
                raise OSError(23, "CRC")
            return real_plain(s, d)
        g["copy_plain"] = bad_plain
        copy_file(b, dst_of(b), sb.st_size, sb.st_mtime)
        assert records[b]["狀態"] == PENDING and records[b]["錯誤類型"] == "壞軌"
        k0 = S["known"]
        copy_file(b, dst_of(b), sb.st_size, sb.st_mtime)
        assert S["known"] == k0 + 1
        Reader.read_at = lambda self, pos, n: ra(pos, n)
        rescue_file(records[b])
        pd = partial_name(dst_of(b))
        assert pd.endswith("b_部分損壞.jpg")
        assert records[b]["狀態"] == "部分救回" and os.path.getsize(pd) == len(data)
        assert not os.path.exists(dst_of(b)) and abs(os.path.getmtime(pd) - sb.st_mtime) <= 2
        assert records[b]["救回比例%"] == round((len(data) - P2_BLOCK) * 100 / len(data), 1)
        Reader.read_at = real_read

        # 3. 斷線不能記成壞軌，接回後同一個檔案重來
        c = mk(r"A\c.jpg", b"x" * 100)
        sc = os.stat(c)
        waited = []

        def flaky(s, d):
            if s.endswith("c.jpg") and not waited:
                raise OSError(1167, "裝置未連接")
            return real_plain(s, d)
        g["copy_plain"] = flaky
        g["online"], g["wait_online"] = (lambda: False), (lambda: waited.append(1))
        copy_file(c, dst_of(c), sc.st_size, sc.st_mtime)
        g["online"], g["wait_online"] = real_online, real_wait
        g["copy_plain"] = real_plain
        assert c not in records and waited and is_done(dst_of(c), sc.st_size, sc.st_mtime)

        # 4. CSV → xlsx：只放真的失敗 / 壞軌待救援；暫時性錯誤、已救回的不進 xlsx；舊工作表保留
        wb = Workbook()
        wb.active.title = "壞軌與異常清單"
        wb.active.append(["舊資料"])
        wb.save(XLSX)
        save_records()
        load_records()
        base_row = dict(records[b])                 # b = 部分救回 → 要進 xlsx
        for k, kind, status in (("x1", "暫時性(檔案被佔用)", PENDING), ("x2", "壞軌", "已完整救回"),
                                ("x3", "壞軌", PENDING), ("x4", "其他", PENDING)):
            records[k] = dict(base_row, **{"來源完整路徑": k, "錯誤類型": kind, "狀態": status})
        build_xlsx()
        wb = load_workbook(XLSX)
        assert wb.sheetnames == [SHEET, "壞軌與異常清單"] and wb["壞軌與異常清單"]["A1"].value == "舊資料"
        in_sheet = sorted(row[3] for row in wb[SHEET].iter_rows(min_row=2, values_only=True))
        assert in_sheet == sorted([b, "x3"]), in_sheet
        for k in ("x1", "x2", "x3", "x4"):
            del records[k]
        # 4b. 舊版歸成「其他」的錯誤 13 / 22，重新載入後要更正成「壞軌」並進 xlsx
        records["y"] = dict(base_row, **{"來源完整路徑": "y", "錯誤代碼": "13 (0xD)", "錯誤類型": "其他", "狀態": PENDING})
        save_records()
        load_records()
        assert records["y"]["錯誤類型"] == "壞軌" and in_excel(records["y"])
        del records["y"]

        # 5. 第一階段整趟：a 是殘檔要重抄、b 已部分救回但現在讀得到 → 抄完整版、c 已完成跳過
        S.update(copied=0, skipped=0)
        phase1()
        assert S["copied"] == 2 and S["skipped"] == 1 and is_done(dst_of(a), sa.st_size, sa.st_mtime)

        # 5b. 卡住超過 STALL_SEC 才斷線：接回後不重讀，直接記成「卡住(壞軌)」並進 xlsx
        s5 = mk(r"A\s.jpg", b"y" * 100)
        ss = os.stat(s5)
        calls = []

        def stuck(s, d):
            calls.append(1)
            S["last_progress"] = time.time() - STALL_SEC - 5
            raise OSError(1167, "裝置未連接")
        g["copy_plain"] = stuck
        g["online"], g["wait_online"] = (lambda: False), (lambda: None)
        copy_file(s5, dst_of(s5), ss.st_size, ss.st_mtime)
        g["online"], g["wait_online"], g["copy_plain"] = real_online, real_wait, real_plain
        assert len(calls) == 1 and records[s5]["錯誤類型"] == "卡住(壞軌)" and in_excel(records[s5])

        # 5c. 舊清單的已知壞檔：第一階段完全不讀，直接記成待救援
        kb = mk(r"A\k.jpg", b"z" * 100)
        sk = os.stat(kb)
        KNOWN_BAD.add(("A", "k.jpg"))

        def must_not_read(s, d):
            raise AssertionError("不應該讀取已知壞檔")
        g["copy_plain"] = must_not_read
        copy_file(kb, dst_of(kb), sk.st_size, sk.st_mtime)
        g["copy_plain"] = real_plain
        KNOWN_BAD.discard(("A", "k.jpg"))
        assert records[kb]["狀態"] == PENDING and records[kb]["錯誤類型"] == "壞軌" and in_excel(records[kb])

        # 5d. --restart 用的「程式啟動時間」要讀得到（用自己這個程式測）
        assert abs(proc_start_time(os.getpid()) - time.time()) < 600

        # 5e. 第 5 步已經完整走過資料夾 A → 重新啟動時直接略過，完全不讀 E 槽
        assert "A" in completed_tops()
        # 不同來源根目錄的紀錄不能互相影響：換了來源就不算已完成，換回來又算
        logp = os.path.join(WORK, "進度紀錄.log")
        with open(logp, "a", encoding="utf-8") as f:
            f.write("[00:00:00] 來源：X:\\別的地方\n[00:00:00] [資料夾完成 1/1] Q：新複製 0，先前已完成 1，壞檔 0\n")
        assert "Q" not in completed_tops() and "A" in completed_tops()
        # 磁碟根目錄當來源時，dst_of / top_of 的路徑要接得對
        sr, dr = SRC_ROOT, DST_ROOT
        g["SRC_ROOT"], g["DST_ROOT"] = "E:\\", r"G:\雲端\!家裡硬碟"
        assert dst_of(r"E:\家人A\a\b.jpg") == r"G:\雲端\!家裡硬碟\家人A\a\b.jpg" and top_of(r"E:\家人A\a\b.jpg") == "家人A"
        g["SRC_ROOT"], g["DST_ROOT"] = sr, dr
        real_walk = g["walk"]

        def must_not_walk(p):
            raise AssertionError("已完成的資料夾不應該再讀")
        g["walk"] = must_not_walk
        try:
            phase1()
        finally:
            g["walk"] = real_walk

        # 5f. 同一個資料夾連續失敗 CONSEC_FAIL_LIMIT 張 → 剩下的整批延後、不再讀（第二階段的行為見 5j）
        dd = os.path.join(SRC_ROOT, "D")
        for i in range(CONSEC_FAIL_LIMIT + 3):
            mk(rf"D\f{i}.jpg", b"q" * 10)
        tried = []

        def all_fail(s, d):
            tried.append(s)
            raise OSError(22, "裝置無法辨識此命令")
        g["copy_plain"] = all_fail
        S["phase"] = "第一階段"
        try:
            assert walk(dd) is True
            assert len(tried) == CONSEC_FAIL_LIMIT, tried
            assert records[dd]["狀態"] == PENDING and records[dd]["類型"] != "檔案" and in_excel(records[dd])
        finally:
            g["copy_plain"] = real_plain
            for k in [k for k in records if k.startswith(dd)]:
                del records[k]

        # 5g. 連續 HUNG_DIR_LIMIT 個資料夾（不限層級）都讀不到清單 → 探測硬碟：
        #     卡死 → 撤銷這段期間的紀錄、自動停止；還活著 → 紀錄保留、繼續。讀不到的資料夾都不能算「已完成」
        zs = [os.path.join(SRC_ROOT, "Zt", f"Z{i}") for i in range(HUNG_DIR_LIMIT)]   # 巢狀在 Zt 底下，測非第一層
        for z in zs:
            mk(os.path.join("Zt", os.path.basename(z), "p.jpg"), b"w")
        real_scandir = os.scandir

        def bad_scandir(p, *x):
            if "\\Z" in str(p) and str(p).rstrip("\\").endswith(("Z0", "Z1", "Z2")):
                raise OSError(1117, "I/O 裝置錯誤")
            return real_scandir(p, *x)
        os.scandir = bad_scandir
        g["PROBE_FILE"] = a                        # 探測檔要在來源那顆碟上，機制才會啟用
        S["phase"] = "第一階段"
        try:
            g["drive_alive"] = lambda timeout=30: False
            assert phase1() is False
            assert not any(z in records for z in zs) and "Zt" not in completed_tops() and "D" in completed_tops()

            # 5g3. 第一階段單一檔案太久沒進度 → 回報執行緒把它記成「卡住(壞軌)」（進 xlsx）並啟動強制結束+重啟；
            #      同一個檔案只做一次、暫停中不做；自動重啟後一個檔都沒複製成功就又卡住 → 不再跳過（疑似硬碟卡死）
            relaunches.clear()
            xs = os.path.join(SRC_ROOT, "x.jpg")
            S.update(phase="第一階段", state="執行中", cur=xs, relaunched=0, copied=0,
                     cur_start=time.time() - FILE_STALL_KILL_SEC - 1)
            S["last_progress"] = S["cur_start"]
            assert stall_watchdog() is True and len(relaunches) == 1
            assert is_pending(xs) and records[xs]["錯誤類型"] == "卡住(壞軌)" and in_excel(records[xs])
            assert stall_watchdog() is False and len(relaunches) == 1
            S.update(state="G 槽空間不足", cur_start=time.time() - FILE_STALL_KILL_SEC - 1)
            assert stall_watchdog() is False and len(relaunches) == 1
            del records[xs]
            S.update(state="執行中", cur=xs + "2", relaunched=1, copied=0, cur_start=time.time() - FILE_STALL_KILL_SEC - 1)
            S["last_progress"] = S["cur_start"]
            assert stall_watchdog() is False and len(relaunches) == 1 and (xs + "2") not in records
            S.update(copied=3, cur_start=time.time() - FILE_STALL_KILL_SEC - 1)
            S["last_progress"] = S["cur_start"]
            assert stall_watchdog() is True and len(relaunches) == 2 and is_pending(xs + "2")
            del records[xs + "2"]
            S.update(state="執行中", cur_start=0.0, relaunched=0, copied=0)
            # 重啟參數：拿掉舊的 --relaunched / --relaunch-after，補上新的次數
            assert relaunch_args(["phase1", "--tops", "others", "--relaunched", "1", "--relaunch-after", "99"], 2) == \
                ["phase1", "--tops", "others", "--relaunched", "2"]
        finally:
            os.scandir = real_scandir
            g["drive_alive"] = real_alive
            for z in zs:
                records.pop(z, None)
            shutil.rmtree(os.path.join(SRC_ROOT, "Zt"))

        # 5g2. 寫入 G 槽失敗不是壞軌：磁碟滿 → 等一下重試到成功；其他寫入錯誤重試 DST_RETRY_LIMIT 次後記成「目的地寫入失敗」，不進 xlsx
        w = mk(r"A\w.jpg", b"v" * 50)
        sw = os.stat(w)
        calls = []

        def full_then_ok(s, d):
            calls.append(1)
            if len(calls) < 3:
                raise DstError(28, "No space left on device")
            return real_plain(s, d)
        g["copy_plain"] = full_then_ok
        b0 = S["bad"]
        copy_file(w, dst_of(w), sw.st_size, sw.st_mtime)
        assert len(calls) == 3 and w not in records and S["bad"] == b0 and is_done(dst_of(w), sw.st_size, sw.st_mtime)
        w2 = mk(r"A\w2.jpg", b"v" * 50)
        sw2 = os.stat(w2)
        calls.clear()

        def always_denied(s, d):
            calls.append(1)
            raise DstError(5, "Access is denied")
        g["copy_plain"] = always_denied
        c0 = S["consec"]
        copy_file(w2, dst_of(w2), sw2.st_size, sw2.st_mtime)
        g["copy_plain"] = real_plain
        assert len(calls) == DST_RETRY_LIMIT + 1 and records[w2]["錯誤類型"] == "目的地寫入失敗"
        assert not in_excel(records[w2]) and S["consec"] == c0
        del records[w2]
        for f_ in (w, w2):
            os.remove(f_)
        os.remove(dst_of(w))

        # 5h. 硬碟探測：讀得到的檔案 → 活著；不存在 → 不算活著
        g["PROBE_FILE"] = a
        assert drive_alive() is True
        g["PROBE_FILE"] = os.path.join(SRC_ROOT, "不存在.jpg")
        assert drive_alive() is False

        # 5i. 第二階段：資料夾先處理、檔案最後；資料夾失敗時硬碟探測不到 → 還原成待救援並自動停止
        saved = dict(records)
        records.clear()
        mkp = lambda k, kind: dict(base_row, **{"來源完整路徑": k, "類型": kind, "錯誤代碼": "22 (0x16)",
                                                "錯誤類型": "壞軌", "狀態": PENDING, "大小(MB)": ""})
        records.update({"f1": mkp("f1", "檔案"), "d1": mkp("d1", "資料夾（壞軌密集，延後）"), "d2": mkp("d2", "資料夾無法讀取")})
        order = []
        real_rf, real_rd = g["rescue_file"], g["rescue_dir"]
        g["rescue_file"] = lambda r: (order.append(r["來源完整路徑"]), finish(r, "已完整救回", 100))[1]
        g["rescue_dir"] = lambda r: (order.append(r["來源完整路徑"]), finish(r, "已可讀取（已重新搬移）"))[1]
        try:
            assert phase2() is None and order[-1] == "f1" and set(order[:2]) == {"d1", "d2"}, order
            for r in records.values():
                r["狀態"] = PENDING
            g["rescue_dir"] = lambda r: finish(r, "資料夾結構損壞，需專業工具")
            g["drive_alive"] = lambda timeout=30: False
            assert phase2() is False
            assert sum(r["狀態"] == PENDING for r in records.values()) == 3     # 失敗的那一項已還原
        finally:
            g["rescue_file"], g["rescue_dir"], g["drive_alive"] = real_rf, real_rd, real_alive
            records.clear()
            records.update(saved)

        # 5j. 第二階段處理資料夾時，同一個資料夾又連續失敗 → 留在待救援（不能標成「已可讀取」），下次從斷點接著處理
        ee = os.path.join(SRC_ROOT, "E")
        for i in range(CONSEC_FAIL_LIMIT + 3):
            mk(rf"E\g{i}.jpg", b"e" * 10)
        er = mkp(ee, "資料夾（壞軌密集，延後）")
        records[ee] = er
        g["copy_plain"] = all_fail
        S["phase"] = "第二階段"
        try:
            assert rescue_dir(er) == PENDING and records[ee]["狀態"] == PENDING and records[ee] is not er
            assert sum(k.startswith(ee + "\\") for k in records) == CONSEC_FAIL_LIMIT
        finally:
            g["copy_plain"] = real_plain
            for k in [k for k in records if k.startswith(ee)]:
                del records[k]

        # 6. 回報可以產生；沒進度超過 STALL_SEC 要顯示卡住
        S.update(state="執行中", cur_start=time.time() - STALL_SEC - 1, last_progress=0.0)   # 從沒讀到過資料
        assert "沒有讀到任何資料" in report() and "分" in report() and "2982" not in report()
        S.update(cur_start=time.time(), last_progress=0.0)                   # 剛開始處理，還沒卡住
        assert "沒有讀到任何資料" not in report()
        with open(os.path.join(WORK, "進度紀錄.log"), "a", encoding="utf-8") as f:
            f.write("[03:10:15] ===== 第二階段開始（PID 424242）=====\n")
        assert running_phase(424242) == "第二階段" and running_phase(1) == ""
        S.update(cur_start=0.0)
        for ph in ("第一階段", "第二階段"):
            S["phase"] = ph
            assert "目前" in report()

        # 7. --wait-report：回報一更新就印出；程式執行中但回報太久沒更新要警告
        import contextlib, io
        with open(REPORT_TXT, "w", encoding="utf-8-sig") as f:
            f.write("第一份回報")

        def later():
            time.sleep(1.5)
            with open(REPORT_TXT, "w", encoding="utf-8-sig") as f:
                f.write("第二份回報")
        threading.Thread(target=later).start()
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            wait_report(20)
        assert "第二份回報" in out.getvalue() and "沒有更新" not in out.getvalue(), out.getvalue()
        open(os.path.join(WORK, "rescue.lock"), "w").close()
        os.utime(REPORT_TXT, (time.time() - 3600, time.time() - 3600))
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            wait_report(0)
        assert "沒有更新" in out.getvalue() and "執行中：是" in out.getvalue(), out.getvalue()
    finally:
        g["online"], g["wait_online"], g["copy_plain"], Reader.read_at = real_online, real_wait, real_plain, real_read
        g["relaunch_self"] = real_relaunch
    shutil.rmtree(base, ignore_errors=True)
    print("SELFTEST OK")


if __name__ == "__main__":
    main()
